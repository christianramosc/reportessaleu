# ============================================================
#  CUADRE SALE-U · MG COLIMA
#  Autoexpress Inbursa
#
#  Compara por TELÉFONO los créditos capturados en el archivo de
#  seguimiento (hojas "MG Colima Rechazos" y "MG Colima Créditos
#  Finalizados") contra la exportación de leads de Sale-U.
#
#  Lógica del cuadre (sin interfaz). La usa paginas/cuadre_saleu.py
# ============================================================

import io
import re
import unicodedata
import warnings
from datetime import datetime, date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

AGENCIA = "COLIMA"          # solo se toman hojas cuyo nombre contenga esto
EN_SALEU, FALTA = "EN SALE-U", "FALTA EN SALE-U"


# ------------------------------------------------------------
# UTILIDADES
# ------------------------------------------------------------
def quitar_acentos(s):
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c))


def norm_nombre(s):
    if s is None or (isinstance(s, float) and pd.isna(s)):
        return ""
    s = quitar_acentos(s).upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z ]", " ", s)).strip()


def norm_tel(t):
    """Deja solo dígitos y quita lada de país (52 / 521)."""
    if t is None or (isinstance(t, float) and pd.isna(t)):
        return ""
    txt = str(int(t)) if isinstance(t, float) else str(t)
    d = re.sub(r"\D", "", txt)
    if len(d) == 12 and d.startswith("52"):
        d = d[2:]
    elif len(d) == 13 and d.startswith("521"):
        d = d[3:]
    return d


def a_fecha(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        return pd.to_datetime(v, unit="D", origin="1899-12-30").date()
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return pd.to_datetime(v, dayfirst=True).date()
    except Exception:
        return None


def arreglar_encabezado(c):
    c = str(c).strip()
    try:
        c = c.encode("latin-1").decode("utf-8")   # "TelÃ©fono" -> "Teléfono"
    except Exception:
        pass
    return c


def buscar_col(df, *claves, exacta=None, requerida=True):
    cols = {c: norm_nombre(c) for c in df.columns}
    if exacta:
        for c, n in cols.items():
            if n == exacta:
                return c
    for c, n in cols.items():
        if all(k in n for k in claves):
            return c
    if requerida:
        raise KeyError(f"No encontré la columna {claves or exacta}. Columnas: {list(df.columns)}")
    return None


def valor(x, col):
    if not col:
        return None
    v = x[col]
    return None if (v is None or (isinstance(v, float) and pd.isna(v))) else v


# ------------------------------------------------------------
# CARGA
# ------------------------------------------------------------
def cargar_saleu(archivo):
    df = pd.read_excel(archivo)
    df.columns = [arreglar_encabezado(c) for c in df.columns]
    c_tel = buscar_col(df, "TELEFONO")
    c_nom = buscar_col(df, exacta="NOMBRE") or buscar_col(df, "NOMBRE")
    c_est = buscar_col(df, exacta="ESTATUS", requerida=False)
    c_ase = buscar_col(df, exacta="ASESOR", requerida=False)
    c_fec = buscar_col(df, "FECHA", "ALTA", requerida=False)
    out = pd.DataFrame({
        "tel": df[c_tel].map(norm_tel),
        "nombre": df[c_nom],
        "nombre_n": df[c_nom].map(norm_nombre),
        "estatus": df[c_est] if c_est else None,
        "asesor": df[c_ase] if c_ase else None,
        "fecha_alta": df[c_fec].map(a_fecha) if c_fec else None,
    })
    return out


def hojas_colima(archivo):
    xls = pd.ExcelFile(archivo)
    return [h for h in xls.sheet_names
            if AGENCIA in norm_nombre(h) and ("RECHAZ" in norm_nombre(h) or "FINALIZ" in norm_nombre(h))]


def cargar_hoja(archivo, hoja):
    df = pd.read_excel(archivo, sheet_name=hoja)
    df.columns = [arreglar_encabezado(c) for c in df.columns]
    origen = "Finalizado" if "FINALIZ" in norm_nombre(hoja) else "Rechazo"

    c_nom = buscar_col(df, "NOMBRE")
    c_tel = buscar_col(df, "CELULAR", requerida=False) or buscar_col(df, "TELEFONO")
    c_ase = buscar_col(df, "ASESOR", "VENTAS", requerida=False) or buscar_col(df, exacta="ASESOR", requerida=False)
    c_fec = buscar_col(df, "FECHA", "INGRESO", requerida=False)
    c_fin = buscar_col(df, exacta="FINANCIERA", requerida=False)
    c_uni = buscar_col(df, exacta="UNIDAD", requerida=False)
    c_seg = buscar_col(df, "SALESU", requerida=False)
    if c_seg is None:
        # en Finalizados el seguimiento viene en la columna sin nombre junto a ESTATUS
        c_est = buscar_col(df, exacta="ESTATUS", requerida=False)
        if c_est is not None:
            i = list(df.columns).index(c_est)
            if i + 1 < len(df.columns) and str(df.columns[i + 1]).startswith("Unnamed"):
                c_seg = df.columns[i + 1]

    df = df[df[c_nom].notna() & (df[c_nom].astype(str).str.strip() != "")]
    return df, origen, dict(nom=c_nom, tel=c_tel, ase=c_ase, fec=c_fec, fin=c_fin, uni=c_uni, seg=c_seg)


# ------------------------------------------------------------
# CRUCE
# ------------------------------------------------------------
def posible_por_nombre(nombre_n, saleu_nombres):
    """Todas las palabras (>=3 letras, mínimo 2) del nombre en Sale-U están en el nombre capturado."""
    tb = set(nombre_n.split())
    for n, orig, tel in saleu_nombres:
        ts = {w for w in n.split() if len(w) >= 3}
        if len(ts) >= 2 and ts <= tb:
            return f"{' '.join(str(orig).split())} (tel. Sale-U {tel or 'sin tel.'})"
    return None


def cuadrar(saleu, hojas):
    saleu = saleu.sort_values("fecha_alta", ascending=False, na_position="last")
    idx = {}
    for r in saleu.itertuples(index=False):
        if len(r.tel) >= 8:
            idx.setdefault(r.tel, r)
    nombres = list(saleu[["nombre_n", "nombre", "tel"]].itertuples(index=False, name=None))

    filas = []
    for df, origen, c in hojas:
        for i, x in df.iterrows():
            t = norm_tel(x[c["tel"]])
            m = idx.get(t) if t else None
            seg = valor(x, c["seg"])
            seg_n = norm_nombre(seg) if seg else ""
            fecha = a_fecha(valor(x, c["fec"]))

            obs = []
            if not t:
                obs.append("Sin teléfono")
            elif len(t) != 10:
                obs.append(f"Teléfono con {len(t)} dígitos ({t})")
            posible = None
            if m is None:
                posible = posible_por_nombre(norm_nombre(x[c["nom"]]), nombres)
                dice_registrado = ("REGISTRADO" in seg_n and "NO " not in seg_n and "RESGISTRADO" not in seg_n) \
                    or ("EN SALE U" in seg_n and not seg_n.startswith("NO "))
                if dice_registrado:
                    obs.append("El Excel dice que está en Sale-U, pero su teléfono no aparece")
                if posible:
                    obs.append("Hay un nombre parecido en Sale-U con otro teléfono")
            elif "VERIFICAR" in seg_n:
                obs.append("Marcado 'verificar número', pero el teléfono sí está en Sale-U")
            if fecha and fecha.year < 2020:
                obs.append(f"Fecha de ingreso sospechosa ({fecha:%d/%m/%Y})")

            filas.append({
                "Origen": origen,
                "Fila Excel": i + 2,
                "Fecha ingreso": fecha,
                "Nombre de Cliente": x[c["nom"]],
                "Teléfono": t or None,
                "Financiera": valor(x, c["fin"]),
                "Unidad": valor(x, c["uni"]),
                "Asesor": valor(x, c["ase"]),
                "Seguimiento SalesU (Excel)": seg,
                "Estatus Cuadre": EN_SALEU if m is not None else FALTA,
                "Nombre en Sale-U": m.nombre if m is not None else None,
                "Estatus Sale-U": m.estatus if m is not None else None,
                "Asesor Sale-U": m.asesor if m is not None else None,
                "Fecha Alta Sale-U": m.fecha_alta if m is not None else None,
                "Posible coincidencia por nombre": posible,
                "Observaciones": "; ".join(obs) or None,
            })
    return pd.DataFrame(filas)


def generar_observaciones(cuadre, saleu):
    obs = []
    falt = cuadre[cuadre["Estatus Cuadre"] == FALTA]
    clave = falt["Teléfono"].fillna(falt["Nombre de Cliente"])
    if len(falt):
        obs.append(("alta", f"Faltan {len(falt)} registros en Sale-U, que corresponden a "
                            f"{clave.nunique()} clientes distintos (un cliente puede venir varias veces con distintas financieras)."))
    else:
        obs.append(("ok", "Todos los registros capturados están en Sale-U."))

    for origen in cuadre["Origen"].unique():
        d = cuadre[cuadre["Origen"] == origen]
        f = (d["Estatus Cuadre"] == FALTA).sum()
        obs.append(("info", f"{origen}s: {len(d)} capturados, {len(d) - f} en Sale-U, {f} faltantes "
                            f"({(len(d) - f) / len(d) * 100:.1f}% cuadrado)."))

    txt = cuadre["Observaciones"].fillna("")
    n = (txt.str.contains("El Excel dice") & (cuadre["Estatus Cuadre"] == FALTA)).sum()
    if n:
        obs.append(("alta", f"{n} registros dicen en el Excel que ya están en Sale-U, pero su teléfono no aparece. "
                            f"Probablemente se registraron con otro número o la exportación no cubre ese periodo."))
    pos = falt[falt["Posible coincidencia por nombre"].notna()]
    for _, r in pos.drop_duplicates("Nombre de Cliente").iterrows():
        obs.append(("media", f"{str(r['Nombre de Cliente']).title()} ({'tel. ' + r['Teléfono'] if pd.notna(r['Teléfono']) else 'sin teléfono'}) podría estar en Sale-U como "
                             f"{r['Posible coincidencia por nombre']}. Revisar cuál teléfono es el correcto."))
    ver = cuadre[txt.str.contains("verificar número")]
    if len(ver):
        obs.append(("ok", f"{len(ver)} registros marcados como 'verificar número' sí aparecen en Sale-U con ese teléfono: "
                          f"{', '.join(ver['Nombre de Cliente'].astype(str).str.title())}."))
    sin_tel = cuadre[txt.str.contains("Sin teléfono")]
    if len(sin_tel):
        obs.append(("media", f"{len(sin_tel)} registros no tienen teléfono capturado: "
                             f"{', '.join(sin_tel['Nombre de Cliente'].astype(str).str.title())}."))
    raros = cuadre[txt.str.contains("dígitos")]
    if len(raros):
        obs.append(("media", (f"{len(raros)} teléfonos no tienen" if len(raros) > 1 else "1 teléfono no tiene") + " 10 dígitos: "
                             + ", ".join(f"{a.title()} ({b})" for a, b in zip(raros['Nombre de Cliente'].astype(str), raros['Teléfono']))))
    fechas = cuadre[txt.str.contains("Fecha de ingreso sospechosa")]
    if len(fechas):
        obs.append(("media", f"{len(fechas)} registros tienen fecha de ingreso anterior a 2020 "
                             f"(parecen fechas de nacimiento capturadas en la columna equivocada)."))
    rep = falt[falt.duplicated("Teléfono", keep=False) & falt["Teléfono"].notna()]
    if len(rep):
        g = rep.groupby("Teléfono")["Nombre de Cliente"].agg(["first", "size"])
        obs.append(("info", "Clientes faltantes que aparecen varias veces: "
                            + ", ".join(f"{a.title()} (x{b})" for a, b in g.values) + "."))
    if len(falt) and falt["Asesor"].notna().any():
        top = falt["Asesor"].value_counts()
        obs.append(("info", f"El asesor con más faltantes es {str(top.index[0]).title()} con {top.iloc[0]}."))
    fa = saleu["fecha_alta"].dropna()
    if len(fa):
        malos = (saleu["tel"].str.len() != 10).sum()
        obs.append(("info", f"La exportación de Sale-U tiene {len(saleu):,} leads del {min(fa):%d/%m/%Y} al "
                            f"{max(fa):%d/%m/%Y}; {malos} de ellos tienen teléfono vacío o con formato inválido."))
    return obs


# ------------------------------------------------------------
# EXCEL
# ------------------------------------------------------------
AZUL = "1F4E78"
F_HDR = Font(name="Arial", bold=True, color="FFFFFF", size=10)
F_TXT = Font(name="Arial", size=10)
F_B = Font(name="Arial", size=10, bold=True)
RELLENO = {EN_SALEU: "E2EFDA", FALTA: "F8CBAD"}
COLOR_OBS = {"alta": "F8CBAD", "media": "FFF2CC", "ok": "E2EFDA", "info": "FFFFFF"}
ANCHOS = [12, 8, 12, 34, 13, 12, 24, 30, 26, 17, 30, 12, 26, 13, 40, 48]


def hoja_tabla(wb, nombre, df):
    ws = wb.create_sheet(nombre)
    ws.append(list(df.columns))
    for fila in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in fila])
    for c in ws[1]:
        c.font, c.fill = F_HDR, PatternFill("solid", start_color=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    col_est = list(df.columns).index("Estatus Cuadre")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = F_TXT, Border(bottom=Side(style="thin", color="D9D9D9"))
            if isinstance(c.value, (date, datetime)):
                c.number_format = "DD/MM/YYYY"
        e = row[col_est]
        e.fill, e.font = PatternFill("solid", start_color=RELLENO.get(e.value, "FFFFFF")), F_B
    for i, w in enumerate(ANCHOS[:len(df.columns)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions


def generar_excel(cuadre, obs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    falt = cuadre[cuadre["Estatus Cuadre"] == FALTA].sort_values(["Origen", "Asesor", "Fecha ingreso"], na_position="last")
    origenes = [o for o in ["Rechazo", "Finalizado"] if o in cuadre["Origen"].unique()]
    hojas = {o: f"{o}s MG Colima" for o in origenes}

    hoja_tabla(wb, "Faltan en Sale-U", falt)
    for o in origenes:
        hoja_tabla(wb, hojas[o], cuadre[cuadre["Origen"] == o])

    # --- Resumen con fórmulas ---
    ws["A1"] = "CUADRE SALE-U · MG COLIMA"
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color=AZUL)
    ws["A2"] = f"Generado el {datetime.now():%d/%m/%Y %H:%M} · Cruce por teléfono"
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="808080")
    encabezados = ["Concepto"] + [hojas[o].replace(" MG Colima", "") for o in origenes] + ["Total"]
    for j, h in enumerate(encabezados, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill = F_HDR, PatternFill("solid", start_color=AZUL)
        c.alignment = Alignment(horizontal="center")
    conceptos = ["Registros capturados", "Encontrados en Sale-U", "Faltan en Sale-U", "% cuadrado"]
    col_tot = len(origenes) + 2
    for k, txt in enumerate(conceptos, 5):
        ws.cell(row=k, column=1, value=txt).font = F_TXT
        for j in range(2, col_tot + 1):
            L = get_column_letter(j)
            if k == 8:
                f = f"=IF({L}5=0,0,{L}6/{L}5)"
            elif j == col_tot:
                f = f"=SUM(B{k}:{get_column_letter(col_tot - 1)}{k})"
            else:
                h = hojas[origenes[j - 2]]
                f = {5: f"=COUNTA('{h}'!D:D)-1",
                     6: f"=COUNTIF('{h}'!J:J,\"{EN_SALEU}\")",
                     7: f"=COUNTIF('{h}'!J:J,\"{FALTA}\")"}[k]
            c = ws.cell(row=k, column=j, value=f)
            c.font = F_B if (k == 7 or j == col_tot) else F_TXT
            c.alignment = Alignment(horizontal="center")
            if k == 8:
                c.number_format = "0.0%"
    for j in range(1, col_tot + 1):
        ws.cell(row=7, column=j).fill = PatternFill("solid", start_color="F8CBAD")

    ws["A10"] = "Faltantes por asesor"
    ws["A10"].font = Font(name="Arial", bold=True, size=11, color=AZUL)
    for j, h in enumerate(["Asesor"] + origenes + ["Total"], 1):
        c = ws.cell(row=11, column=j, value=h)
        c.font, c.fill = F_HDR, PatternFill("solid", start_color=AZUL)
    for k, a in enumerate(sorted(falt["Asesor"].dropna().astype(str).unique()), 12):
        ws.cell(row=k, column=1, value=a).font = F_TXT
        for j, o in enumerate(origenes, 2):
            ws.cell(row=k, column=j,
                    value=f"=COUNTIFS('Faltan en Sale-U'!H:H,A{k},'Faltan en Sale-U'!A:A,\"{o}\")").font = F_TXT
        ws.cell(row=k, column=len(origenes) + 2,
                value=f"=SUM(B{k}:{get_column_letter(len(origenes) + 1)}{k})").font = F_B
    ws.column_dimensions["A"].width = 42
    for j in range(2, 6):
        ws.column_dimensions[get_column_letter(j)].width = 14

    # --- Observaciones ---
    wo = wb.create_sheet("Observaciones", 1)
    wo["A1"], wo["B1"] = "Prioridad", "Observación"
    for c in wo[1]:
        c.font, c.fill = F_HDR, PatternFill("solid", start_color=AZUL)
    etiqueta = {"alta": "Alta", "media": "Revisar", "ok": "OK", "info": "Info"}
    for i, (nivel, texto) in enumerate(obs, 2):
        a = wo.cell(row=i, column=1, value=etiqueta[nivel])
        b = wo.cell(row=i, column=2, value=texto)
        a.font, b.font = F_B, F_TXT
        a.fill = PatternFill("solid", start_color=COLOR_OBS[nivel])
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.alignment = Alignment(vertical="top", horizontal="center")
    wo.column_dimensions["A"].width = 11
    wo.column_dimensions["B"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


