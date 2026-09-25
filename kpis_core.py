# -*- coding: utf-8 -*-
"""
kpis_core.py
Genera un PDF y PNGs descargables con los KPIs de capacitación (Ventas y Posventa) de MG Colima
a partir de las hojas "KPIS-VENTAS" y "KPIS POSVENTA" del archivo PLAN DE ACCIÓN.

Los totales y avances se RECALCULAN desde las filas de cada colaborador
(no se usan las celdas TOTAL del Excel), para evitar errores de rangos.
Avance de cada módulo = promedio del avance individual (concluidas / total).
"""
import os
import re
import unicodedata
from datetime import datetime

import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyBboxPatch, Circle, Wedge, Rectangle, Polygon

# ---------------------------------------------------------------- CONFIG
ARCHIVO = "PLAN_DE_ACCIÓN.xlsx"
SALIDA = "KPIs_Capacitacion_MG_Colima.pdf"
HOJAS = {"VENTAS": "KPIS-VENTAS", "POSVENTA": "KPIS POSVENTA"}

MODULOS = {  # clave de búsqueda en el título de la tabla -> (nombre, color, color claro, icono)
    "DOT":        ("DOT",        "#1E6FD9", "#EAF3FD", "grupo"),
    "CAPSULAS":   ("Cápsulas",   "#138A4B", "#E9F6EE", "pantalla"),
    "LOL":        ("LOL",        "#5B3FA8", "#F1EDFA", "grupo"),
    "PRESENCIAL": ("Presencial", "#E8820C", "#FDF3E7", "presentador"),
}
ORDEN = ["DOT", "CAPSULAS", "LOL", "PRESENCIAL"]
AZUL_OSCURO = "#12305C"
GRIS = "#D5DBE3"
TEXTO = "#1B2A41"
TEXTO_2 = "#5A6B82"

plt.rcParams["font.family"] = "DejaVu Sans"


# ---------------------------------------------------------------- LECTURA
def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return s.upper()


def _num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return 0.0  # 'N/A', vacío, etc.


def leer_hoja(ws):
    """Busca las 4 tablas por su título y devuelve {modulo: [filas]}."""
    tablas = {}
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str):
                continue
            t = _norm(c.value)
            m = re.search(r"\((CAPSULAS|DOT|LOL|PRESENCIAL)\)", t)
            if not m or m.group(1) in tablas:
                continue
            clave, col, fila = m.group(1), c.column, c.row + 2  # título, encabezado, datos
            filas = []
            while fila <= ws.max_row:
                nombre = ws.cell(fila, col).value
                if nombre is None or "TOTAL" in _norm(nombre):
                    break
                total = _num(ws.cell(fila, col + 2).value)
                concl = _num(ws.cell(fila, col + 3).value)
                filas.append({
                    "nombre": str(nombre).strip(),
                    "letra": str(ws.cell(fila, col + 1).value or "").strip().upper(),
                    "total": total,
                    "concluidas": concl,
                    "pendientes": _num(ws.cell(fila, col + 4).value),
                    "noshow": _num(ws.cell(fila, col + 5).value),
                    "avance": (concl / total) if total else 0.0,
                })
                fila += 1
            tablas[clave] = filas
    return tablas


def resumir(tablas):
    res = {}
    for k in ORDEN:
        f = tablas.get(k, [])
        n = len(f)
        res[k] = {
            "total": sum(x["total"] for x in f),
            "concluidas": sum(x["concluidas"] for x in f),
            "pendientes": sum(x["pendientes"] for x in f),
            "noshow": sum(x["noshow"] for x in f),
            "avance": (sum(x["avance"] for x in f) / n) if n else 0.0,
        }
    # por letra: promedio de todos los avances individuales de esa letra
    letras = {}
    for k in ORDEN:
        for x in tablas.get(k, []):
            d = letras.setdefault(x["letra"], {"av": [], "total": 0})
            d["av"].append(x["avance"])
            d["total"] += x["total"]
    por_letra = {l: {"avance": sum(d["av"]) / len(d["av"]), "total": d["total"]}
                 for l, d in letras.items() if l}
    general = sum(res[k]["avance"] for k in ORDEN) / len(ORDEN)
    return res, por_letra, general


# ---------------------------------------------------------------- DIBUJO
def caja(ax, x, y, w, h, fc, ec, r=0.012, lw=1.0):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0,rounding_size={r}",
                                fc=fc, ec=ec, lw=lw, transform=ax.transAxes, zorder=1))


def dona(fig, rect, pct, color, texto_centro, sub="Avance", grosor=0.22, fs=15):
    a = fig.add_axes(rect)
    a.set_aspect("equal")
    a.axis("off")
    a.set_xlim(-1.05, 1.05)
    a.set_ylim(-1.05, 1.05)
    a.add_patch(Wedge((0, 0), 1, 0, 360, width=grosor, fc=GRIS, ec="none"))
    p = max(0.0, min(pct, 1.0))
    if p > 0:
        a.add_patch(Wedge((0, 0), 1, 90 - 360 * p, 90, width=grosor, fc=color, ec="none"))
    a.text(0, 0.1, texto_centro, ha="center", va="center", fontsize=fs, weight="bold", color=TEXTO)
    a.text(0, -0.28, sub, ha="center", va="center", fontsize=fs * 0.45, color=TEXTO_2)
    return a


def icono(fig, rect, color, tipo):
    a = fig.add_axes(rect)
    a.set_aspect("equal")
    a.axis("off")
    a.set_xlim(-1, 1)
    a.set_ylim(-1, 1)
    a.add_patch(Circle((0, 0), 1, fc=color, ec="none"))
    w = "white"
    if tipo == "grupo":
        for dx, s in [(-0.38, 0.8), (0.38, 0.8), (0, 1.0)]:
            a.add_patch(Circle((dx, 0.18 * s + 0.02), 0.17 * s, fc=w, ec=color, lw=1.2))
            a.add_patch(Wedge((dx, -0.42), 0.34 * s, 0, 180, fc=w, ec=color, lw=1.2))
    elif tipo == "pantalla":
        a.add_patch(Rectangle((-0.55, -0.3), 1.1, 0.75, fc="none", ec=w, lw=2.2))
        a.add_patch(Rectangle((-0.3, -0.5), 0.6, 0.07, fc=w, ec="none"))
        a.add_patch(Polygon([(-0.12, -0.1), (-0.12, 0.25), (0.2, 0.07)], fc=w, ec="none"))
    else:  # presentador
        a.add_patch(Rectangle((-0.1, -0.05), 0.65, 0.5, fc="none", ec=w, lw=2))
        a.add_patch(Circle((-0.38, 0.18), 0.17, fc=w, ec="none"))
        a.add_patch(Wedge((-0.38, -0.5), 0.33, 0, 180, fc=w, ec="none"))
        a.plot([0.05, 0.4], [0.05, 0.3], color=w, lw=2)


def pct(v):
    return f"{v * 100:.0f}%"


def entero(v):
    return f"{v:,.0f}"


def pagina_resumen(area, res, por_letra, general, n_colab):
    fig = plt.figure(figsize=(11, 8.5))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.set_autoscale_on(False)

    # encabezado
    ax.add_patch(Rectangle((0, 0.915), 1, 0.085, fc=AZUL_OSCURO, ec="none", transform=ax.transAxes))
    ax.text(0.035, 0.957, f"AVANCE DE CAPACITACIÓN {area} · MG COLIMA", color="white",
            fontsize=17, weight="bold", va="center")
    ax.text(0.965, 0.957, f"{n_colab} colaboradores · {datetime.now():%d/%m/%Y}",
            color="#C9D6EA", fontsize=9.5, va="center", ha="right")

    # ---------- tarjetas de módulo
    mx, gap, y0, h = 0.03, 0.015, 0.545, 0.345
    w = (1 - 2 * mx - 3 * gap) / 4
    for i, k in enumerate(ORDEN):
        nombre, col, claro, tipo = MODULOS[k]
        r = res[k]
        x = mx + i * (w + gap)
        caja(ax, x, y0, w, h, claro, col + "55")
        ax.add_patch(Rectangle((x + 0.002, y0 + 0.002), w - 0.004, 0.05, fc=col + "22",
                               ec="none", transform=ax.transAxes))
        icono(fig, [x + 0.01, y0 + h - 0.085, 0.07, 0.075], col, tipo)
        ax.text(x + 0.085, y0 + h - 0.0475, nombre, fontsize=15, weight="bold", color=TEXTO, va="center")
        dona(fig, [x + w / 2 - 0.07, y0 + 0.118, 0.14, 0.132], r["avance"], col, pct(r["avance"]), fs=15)
        stats = [(r["concluidas"], "Concluidas"), (r["pendientes"], "Pendientes"), (r["noshow"], "No show")]
        for j, (v, lbl) in enumerate(stats):
            cx = x + w * (j + 0.5) / 3
            ax.text(cx, y0 + 0.083, entero(v), fontsize=14, weight="bold", color=TEXTO, ha="center")
            ax.text(cx, y0 + 0.062, lbl, fontsize=7.5, color=TEXTO_2, ha="center")
            if j < 2:
                dx = x + w * (j + 1) / 3
                ax.plot([dx, dx], [y0 + 0.06, y0 + 0.105], color=col + "55", lw=1)
        ax.text(x + w / 2, y0 + 0.026, f"Total de {nombre}: {entero(r['total'])}",
                fontsize=10.5, weight="bold", color=TEXTO, ha="center", va="center")

    # ---------- avance por letra
    lx, ly, lw_, lh = 0.03, 0.06, 0.455, 0.46
    caja(ax, lx, ly, lw_, lh, "white", GRIS)
    ax.add_patch(Rectangle((lx, ly + lh - 0.065), lw_, 0.065, fc=AZUL_OSCURO, ec="none", transform=ax.transAxes))
    ax.text(lx + 0.02, ly + lh - 0.0325, "AVANCE GENERAL POR LETRA", color="white", fontsize=12,
            weight="bold", va="center")
    orden_l = sorted(por_letra, key=lambda l: (l != "P", l))
    n = max(len(orden_l), 1)
    cw = (lw_ - 0.04) / n
    for i, l in enumerate(orden_l):
        d = por_letra[l]
        cx0 = lx + 0.02 + i * cw
        caja(ax, cx0 + 0.005, ly + 0.03, cw - 0.01, lh - 0.12, "#F6F8FB", "#E3E8EF")
        mid = cx0 + cw / 2
        ax.text(mid, ly + 0.3, l, fontsize=16, weight="bold", color=TEXTO, ha="center")
        ax.text(mid, ly + 0.21, pct(d["avance"]), fontsize=22, weight="bold", color=TEXTO, ha="center")
        bw, bx, by = cw - 0.05, cx0 + 0.025, ly + 0.155
        caja(ax, bx, by, bw, 0.026, GRIS, "none", r=0.01)
        if d["avance"] > 0:
            caja(ax, bx, by, max(bw * min(d["avance"], 1), 0.012), 0.026, "#3DBE55", "none", r=0.01)
        ax.text(mid, ly + 0.075, f"Asignaciones: {entero(d['total'])}", fontsize=9.5,
                color=TEXTO_2, ha="center")

    # ---------- avance promedio general
    gx, gy, gw, gh = 0.515, 0.06, 0.455, 0.46
    caja(ax, gx, gy, gw, gh, "white", GRIS)
    ax.add_patch(Rectangle((gx, gy + gh - 0.065), gw, 0.065, fc=AZUL_OSCURO, ec="none", transform=ax.transAxes))
    ax.text(gx + 0.02, gy + gh - 0.0325, "AVANCE PROMEDIO GENERAL", color="white", fontsize=12,
            weight="bold", va="center")
    dona(fig, [gx + 0.01, gy + 0.05, 0.2, 0.28], general, "#138A4B", pct(general),
         sub="Avance total", grosor=0.2, fs=24)
    lx2, ly2 = gx + 0.235, gy + 0.05
    caja(ax, lx2, ly2, 0.2, 0.28, "#F6F8FB", "#E3E8EF")
    for i, k in enumerate(ORDEN):
        nombre, col, _, _ = MODULOS[k]
        yy = ly2 + 0.245 - i * 0.065
        ax.add_patch(Circle((lx2 + 0.022, yy), 0.008, fc=col, ec="none", transform=ax.transAxes))
        ax.text(lx2 + 0.04, yy, nombre, fontsize=11, color=TEXTO, va="center")
        ax.text(lx2 + 0.185, yy, pct(res[k]["avance"]), fontsize=11, weight="bold", color=TEXTO,
                va="center", ha="right")
        if i < 3:
            ax.plot([lx2 + 0.012, lx2 + 0.188], [yy - 0.0325] * 2, color="#E3E8EF", lw=1,
                    transform=ax.transAxes)

    ax.text(0.03, 0.025, "Avance por módulo = promedio del avance individual (concluidas / total). "
            "Avance total = promedio de los 4 módulos.", fontsize=7.5, color=TEXTO_2)
    return fig


def pagina_detalle(area, tablas):
    nombres = [x["nombre"] for x in tablas[ORDEN[0]]]
    datos = {k: {x["nombre"]: x["avance"] for x in tablas.get(k, [])} for k in ORDEN}
    fig = plt.figure(figsize=(11, 8.5))
    ax0 = fig.add_axes([0, 0, 1, 1])
    ax0.axis("off")
    ax0.set_xlim(0, 1)
    ax0.set_ylim(0, 1)
    ax0.set_autoscale_on(False)
    ax0.add_patch(Rectangle((0, 0.915), 1, 0.085, fc=AZUL_OSCURO, ec="none", transform=ax0.transAxes))
    ax0.text(0.035, 0.957, f"DETALLE DE AVANCE POR COLABORADOR · {area}", color="white",
             fontsize=17, weight="bold", va="center")

    ax = fig.add_axes([0.30, 0.1, 0.66, 0.77])
    n, bh = len(nombres), 0.19
    for j, k in enumerate(ORDEN):
        nombre, col, _, _ = MODULOS[k]
        ys = [i + (j - 1.5) * bh for i in range(n)]
        vals = [datos[k].get(nm, 0) for nm in nombres]
        ax.barh(ys, vals, height=bh * 0.9, color=col, label=nombre)
    ax.set_yticks(range(n))
    ax.set_yticklabels([nm.title() for nm in nombres], fontsize=9, color=TEXTO)
    ax.invert_yaxis()
    ax.set_xlim(0, 1.0)
    ax.xaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    ax.grid(axis="x", color="#E3E8EF")
    ax.set_axisbelow(True)
    for s in ["top", "right", "left"]:
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.legend(ncol=4, loc="lower center", bbox_to_anchor=(0.5, 1.0), frameon=False, fontsize=10)
    return fig


# ---------------------------------------------------------------- EXPORTACIÓN
from io import BytesIO
import zipfile
from matplotlib.transforms import Bbox

FIG_W, FIG_H = 11, 8.5
# recortes de la página de resumen, en fracción de la figura (x0, y0, x1, y1)
SECCIONES = {
    "1_tarjetas_modulos": (0.02, 0.535, 0.98, 0.90),
    "2_avance_por_letra": (0.02, 0.05, 0.495, 0.53),
    "3_avance_promedio_general": (0.505, 0.05, 0.98, 0.53),
}


def _png(fig, recorte=None, dpi=200):
    buf = BytesIO()
    kw = {}
    if recorte:
        x0, y0, x1, y1 = recorte
        kw["bbox_inches"] = Bbox([[x0 * FIG_W, y0 * FIG_H], [x1 * FIG_W, y1 * FIG_H]])
    fig.savefig(buf, format="png", dpi=dpi, facecolor="white", **kw)
    return buf.getvalue()


def generar(archivo):
    """archivo: ruta o bytes/BytesIO del Excel.
    Devuelve dict con: pdf (bytes), imagenes {nombre: bytes}, zip (bytes), resumen {area: datos}."""
    if isinstance(archivo, (bytes, bytearray)):
        archivo = BytesIO(archivo)
    wb = openpyxl.load_workbook(archivo, data_only=True)
    imagenes, resumen, figs = {}, {}, []
    for area, hoja in HOJAS.items():
        if hoja not in wb.sheetnames:
            raise ValueError(f"No se encontró la hoja '{hoja}' en el archivo.")
        tablas = leer_hoja(wb[hoja])
        faltan = [k for k in ORDEN if k not in tablas]
        if faltan:
            raise ValueError(f"En la hoja '{hoja}' no se encontraron las tablas: {faltan}")
        res, por_letra, general = resumir(tablas)
        n_colab = len(tablas[ORDEN[0]])
        resumen[area] = {"modulos": res, "por_letra": por_letra, "general": general,
                         "colaboradores": n_colab, "tablas": tablas}

        f1 = pagina_resumen(area, res, por_letra, general, n_colab)
        f2 = pagina_detalle(area, tablas)
        pref = area.lower()
        imagenes[f"{pref}_0_resumen_completo.png"] = _png(f1)
        for nombre, rec in SECCIONES.items():
            imagenes[f"{pref}_{nombre}.png"] = _png(f1, rec)
        imagenes[f"{pref}_4_detalle_por_colaborador.png"] = _png(f2)
        figs += [f1, f2]

    buf_pdf = BytesIO()
    with PdfPages(buf_pdf) as pdf:
        for f in figs:
            pdf.savefig(f)
            plt.close(f)

    buf_zip = BytesIO()
    with zipfile.ZipFile(buf_zip, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, data in imagenes.items():
            z.writestr(nombre, data)
        z.writestr(SALIDA, buf_pdf.getvalue())
    return {"pdf": buf_pdf.getvalue(), "imagenes": imagenes, "zip": buf_zip.getvalue(),
            "resumen": resumen}


def main(archivo=ARCHIVO, carpeta="."):
    out = generar(archivo)
    os.makedirs(carpeta, exist_ok=True)
    ruta_pdf = os.path.join(carpeta, SALIDA)
    ruta_zip = os.path.join(carpeta, "KPIs_Capacitacion_imagenes.zip")
    with open(ruta_pdf, "wb") as f:
        f.write(out["pdf"])
    with open(ruta_zip, "wb") as f:
        f.write(out["zip"])
    for area, d in out["resumen"].items():
        print(f"\n{area}: avance total {pct(d['general'])}")
        for k in ORDEN:
            r = d["modulos"][k]
            print(f"  {MODULOS[k][0]:<11} {pct(r['avance']):>5}  concl {r['concluidas']:.0f}  "
                  f"pend {r['pendientes']:.0f}  no show {r['noshow']:.0f}  total {r['total']:.0f}")
    print(f"\nGenerados: {ruta_pdf} y {ruta_zip}")
    return ruta_pdf, ruta_zip


if __name__ == "__main__":
    try:  # Google Colab: sube el archivo si no está y descarga PDF + ZIP de imágenes
        from google.colab import files
        if not os.path.exists(ARCHIVO):
            subido = files.upload()
            ARCHIVO = next(iter(subido))
        for r in main(ARCHIVO):
            files.download(r)
    except ImportError:
        main()
