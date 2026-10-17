# -*- coding: utf-8 -*-
"""
Formato común de los archivos que generan los cuadres (Excel y PDF).

  · Nombres de columna legibles (ETIQUETAS) en el Excel, el PDF y las tablas de la app.
  · Excel con el estilo del Sheet (Calibri 9, encabezado gris), teléfonos como texto, fechas dd/mm/aaaa.
  · PDF con cifras clave al inicio, secciones seguidas (las vacías no ocupan página), texto escapado
    (un «&» o «<» en un comentario ya no rompe el PDF) y número de página.
"""
import html
import re
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (CondPageBreak, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

PDV = "MG Colima"
AZUL = "#1F3864"
GRIS_TEXTO = "#595959"
TONOS = {"bien": "#1E8449", "mal": "#C00000", "aviso": "#BF8F00", "neutro": AZUL}

# Nombre legible de cada columna interna. Lo que no esté aquí se muestra cambiando "_" por espacio.
ETIQUETAS = {
    "Cliente_bitacora": "Cliente (bitácora)", "Tel_bitacora": "Teléfono (bitácora)",
    "Cuadre_pestana": "¿Pestaña correcta?", "Pestanas_encontrado": "Pestaña donde está",
    "Cliente_seguimiento": "Cliente (seguimiento)", "Tel_seguimiento": "Teléfono (seguimiento)",
    "Folio_seguimiento": "Folio (seguimiento)", "Asesor_seguimiento": "Asesor (seguimiento)",
    "Fecha_ingreso_seg": "Fecha ingreso (seguimiento)", "Similitud_nombre": "Parecido del nombre",
    "Estatus_final": "Último estatus", "Folio_norm": "Folio", "Fila_excel": "Fila", "Fecha_ingreso": "Fecha ingreso",
    "Clasificacion": "Clasificación", "Estatus_seg": "Estatus (seguimiento)", "Veces_capturado": "Veces capturado",
    "Telefono": "Teléfono", "Cuadre_estatus": "¿Estatus cuadra?", "Cliente_SaleU": "Cliente (Sale-U)",
    "Tel_SaleU": "Teléfono (Sale-U)", "Asesor_SaleU": "Asesor (Sale-U)", "Estatus_SaleU": "Estatus (Sale-U)",
    "Financiera_SaleU": "Financiera (Sale-U)", "Fuente_SaleU": "Fuente (Sale-U)",
    "Fecha_alta_SaleU": "Alta en Sale-U", "Fecha_Alta_Credito": "Alta en Sale-U", "Monto_Financiar": "Monto a financiar",
    "Vehiculo": "Vehículo", "VIN_BlueService": "VIN (BlueService)", "Modelo_inventario": "Modelo (inventario)",
    "Version": "Versión", "Ubicacion": "Ubicación", "Entrada_inventario": "Entrada a inventario",
    "Ultimo_PDI_inventario": "Último PDI (inventario)", "Proximo_PDI_inventario": "Próximo PDI (inventario)",
    "Etapa_BS": "Etapa (BS)", "Status_BS": "Estatus (BS)", "Modelo_BS": "Modelo (BS)",
    "Dealer_origen": "Dealer origen", "Dealer_actual": "Dealer actual", "Fecha_venta_BS": "Fecha venta (BS)",
    "Fecha_inventario_BS": "Fecha inventario (BS)", "Aparece_en_inventario": "Aparece en inventario",
    "Situacion": "Situación", "Estatus_Lead": "Estatus del lead", "Auto_Interes": "Auto de interés",
}
COLS_TELEFONO = {"Tel_bitacora", "Tel_seguimiento", "Tel_SaleU", "Telefono", "Celular"}
COLS_MONEDA = {"Mensualidad", "Monto_Financiar"}
COLS_PORCENTAJE = {"Similitud_nombre"}
COLS_LARGAS = {"Observaciones", "Historial", "Aparece_en_inventario", "Situacion", "Clasificacion"}

MESES_CORTOS = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
_MES_TXT = {m.upper(): i for i, m in enumerate(MESES_CORTOS, 1)}


def etiqueta(col):
    return ETIQUETAS.get(col, str(col).replace("_", " "))


def vacio(v):
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() in ("", "nan", "NaT", "None")
    if isinstance(v, (list, dict, set, tuple)):
        return False
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def telefono(v):
    """3123801493.0 / '3123801493.0' → '3123801493'."""
    if vacio(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\.0$", "", str(v).strip())


def texto(v, col=""):
    """Valor tal como se lee en el PDF."""
    if vacio(v):
        return ""
    if col in COLS_TELEFONO:
        return telefono(v)
    if isinstance(v, (datetime, date, pd.Timestamp)):
        return v.strftime("%d/%m/%Y")
    if col in COLS_MONEDA and isinstance(v, (int, float)):
        return f"${v:,.2f}"
    if col in COLS_PORCENTAJE and isinstance(v, (int, float)):
        return f"{v:.0%}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def esc(v):
    """Texto seguro para un Paragraph de reportlab («F&I», «<5000»)."""
    s = html.escape(str(v), quote=False).replace("\n", "<br/>")
    sangria = len(s) - len(s.lstrip(" "))
    return "&nbsp;" * sangria * 2 + s.lstrip(" ")       # conserva la sangría de «   · por folio»


# =============================================================================
# Nombres de archivo
# =============================================================================
def _mes_anio(etiqueta_mes):
    """'JUN 2026' / 'JUNIO 2026' → (2026, 6)."""
    m = re.match(r"\s*([A-Za-zÁÉÍÓÚáéíóú]+)\s+(\d{4})", str(etiqueta_mes))
    if not m:
        return None
    mes = _MES_TXT.get(m.group(1)[:3].upper().replace("SET", "SEP"))
    return (int(m.group(2)), mes) if mes else None


def periodo_de_meses(meses):
    """['JUN 2026', …, 'OCT 2026'] → 'Jun-Oct2026'; si cruza de año → 'Dic2025-Feb2026'."""
    fechas = sorted(f for f in (_mes_anio(m) for m in meses) if f)
    if not fechas:
        return ""
    (a0, m0), (a1, m1) = fechas[0], fechas[-1]
    if (a0, m0) == (a1, m1):
        return f"{MESES_CORTOS[m0 - 1]}{a0}"
    if a0 == a1:
        return f"{MESES_CORTOS[m0 - 1]}-{MESES_CORTOS[m1 - 1]}{a1}"
    return f"{MESES_CORTOS[m0 - 1]}{a0}-{MESES_CORTOS[m1 - 1]}{a1}"


def periodo_de_fechas(ini, fin):
    """01/09/2026 – 15/10/2026 → '01Sep-15Oct2026'."""
    f = lambda d: f"{d.day:02d}{MESES_CORTOS[d.month - 1]}"
    return f"{f(ini)}-{f(fin)}{fin.year}" if ini.year == fin.year else f"{f(ini)}{ini.year}-{f(fin)}{fin.year}"


def nombre_archivo(tipo, periodo=""):
    """Cuadre_Seguimiento_Jun-Oct2026_20261016-1743 (sin extensión)."""
    partes = ["Cuadre", tipo] + ([periodo] if periodo else []) + [datetime.now().strftime("%Y%m%d-%H%M")]
    return "_".join(partes)


# =============================================================================
# Excel
# =============================================================================
FUENTE_XL = "Calibri"
ENC_FONDO = "666666"          # mismo encabezado gris del Sheet


def _valor_excel(v, col):
    if vacio(v):
        return None
    if col in COLS_TELEFONO:
        return telefono(v)
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, date) and not isinstance(v, datetime):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (set, list, tuple)):
        return ", ".join(map(str, v))
    if isinstance(v, str) and re.fullmatch(r"\d{2}/\d{2}/\d{4}", v.strip()):     # «08/01/2026» → fecha real
        try:
            return datetime.strptime(v.strip(), "%d/%m/%Y")
        except ValueError:
            return v
    if hasattr(v, "item"):            # numpy → python
        return v.item()
    return v


def excel(ruta, titulo, subtitulo, hojas, colores=None):
    """hojas: lista de (nombre, DataFrame, columnas). La primera (Resumen) lleva título y subtítulo."""
    colores = colores or {}
    borde = Side(style="thin", color="D9D9D9")
    caja = Border(top=borde, bottom=borde, left=borde, right=borde)
    wb = Workbook()
    wb.remove(wb.active)
    for n, (nombre, df, cols) in enumerate(hojas):
        ws = wb.create_sheet(nombre[:31])
        cols = [c for c in cols if c in df.columns]
        f0 = 1
        if n == 0:
            ws.cell(1, 1, titulo).font = Font(name=FUENTE_XL, size=13, bold=True, color=AZUL[1:])
            ws.cell(2, 1, subtitulo).font = Font(name=FUENTE_XL, size=9, color=GRIS_TEXTO[1:])
            f0 = 4
        for j, c in enumerate(cols, 1):
            celda = ws.cell(f0, j, etiqueta(c))
            celda.font = Font(name=FUENTE_XL, size=9, bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=ENC_FONDO)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.border = caja
        if df.empty:
            ws.cell(f0 + 1, 1, "Sin registros").font = Font(name=FUENTE_XL, size=9, italic=True, color="808080")
        for i, (_, r) in enumerate(df[cols].iterrows(), start=f0 + 1):
            for j, c in enumerate(cols, 1):
                v = _valor_excel(r[c], c)
                celda = ws.cell(i, j, v)
                celda.font = Font(name=FUENTE_XL, size=9)
                celda.border = caja
                celda.alignment = Alignment(vertical="top", wrap_text=c in COLS_LARGAS,
                                            horizontal="left" if n == 0 else None)
                if isinstance(v, datetime):
                    celda.number_format = "dd/mm/yyyy"
                elif c in COLS_MONEDA and isinstance(v, (int, float)):
                    celda.number_format = '"$"#,##0.00'
                elif c in COLS_PORCENTAJE and isinstance(v, (int, float)):
                    celda.number_format = "0%"
                elif c in COLS_TELEFONO or c in ("Folio", "Folios", "Folio_norm", "Folio_seguimiento", "VIN", "VIN_BlueService"):
                    celda.number_format = "@"
                if isinstance(v, str) and v in colores:
                    celda.fill = PatternFill("solid", fgColor=colores[v])
        for j, c in enumerate(cols, 1):
            valores = [len(texto(v, c)) for v in df[c].head(500)] if not df.empty else []
            largo = max([len(etiqueta(c)) * 0.9] + valores)
            ancho = 55 if c in COLS_LARGAS and largo > 55 else min(max(largo + 2, 9), 45)
            ws.column_dimensions[get_column_letter(j)].width = ancho
        if n == 0:
            ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 62)
        ws.freeze_panes = ws.cell(f0 + 1, 1)
        if not df.empty and cols:
            ws.auto_filter.ref = f"A{f0}:{get_column_letter(len(cols))}{f0 + len(df)}"
    wb.save(ruta)


# =============================================================================
# PDF
# =============================================================================
class PDF:
    """Reporte horizontal: título, cifras clave, secciones seguidas y el detalle al final.

        pdf = PDF(ruta, "Cuadre …", "Bitácora: … · Generado: …")
        pdf.cifras([("Rechazados", 51, "", "neutro"), …])
        pdf.seccion("Faltan en el seguimiento", df, cols, anchos, color, "Qué hacer con ellos.")
        pdf.guardar(resumen, nota)
    """
    ANCHO = landscape(letter)[0] - 2.4 * cm

    def __init__(self, ruta, titulo, subtitulo):
        self.ruta, self.titulo, self.subtitulo = ruta, titulo, subtitulo
        base = getSampleStyleSheet()
        self.st = dict(
            h1=ParagraphStyle("h1", parent=base["Title"], fontSize=17, leading=21, alignment=0,
                              textColor=colors.HexColor(AZUL), spaceAfter=2),
            sub=ParagraphStyle("sub", parent=base["Normal"], fontSize=8.5, leading=11,
                               textColor=colors.HexColor(GRIS_TEXTO)),
            h2=ParagraphStyle("h2", parent=base["Heading2"], fontSize=12.5, leading=15, spaceBefore=12,
                              spaceAfter=3, textColor=colors.HexColor(AZUL)),
            txt=ParagraphStyle("txt", parent=base["Normal"], fontSize=8.8, leading=11.5),
            celda=ParagraphStyle("celda", parent=base["Normal"], fontSize=7.5, leading=9),
            enc=ParagraphStyle("enc", parent=base["Normal"], fontSize=7.5, leading=9, textColor=colors.white,
                               fontName="Helvetica-Bold"),
            k_n=ParagraphStyle("k_n", parent=base["Normal"], fontSize=8, leading=10, alignment=TA_CENTER,
                               textColor=colors.HexColor(GRIS_TEXTO)),
            k_v=ParagraphStyle("k_v", parent=base["Normal"], fontSize=20, leading=24, alignment=TA_CENTER,
                               fontName="Helvetica-Bold"),
            k_s=ParagraphStyle("k_s", parent=base["Normal"], fontSize=7.5, leading=9, alignment=TA_CENTER,
                               textColor=colors.HexColor(GRIS_TEXTO)),
        )
        self._cifras, self._secciones, self._vacias = [], [], []

    # ------------------------------------------------------------------ piezas
    def cifras(self, items):
        """items: (nombre, valor, detalle, tono) con tono en bien | mal | aviso | neutro."""
        self._cifras = items

    def tabla(self, df, cols, anchos, color=AZUL):
        escala = self.ANCHO / (sum(anchos) * cm)
        datos = [[Paragraph(esc(etiqueta(c)), self.st["enc"]) for c in cols]]
        for _, r in df.iterrows():
            datos.append([Paragraph(esc(texto(r[c], c)), self.st["celda"]) for c in cols])
        t = Table(datos, colWidths=[a * cm * escala for a in anchos], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(color)),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BFBFBF")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
        return t

    def seccion(self, titulo, df, cols, anchos, color=AZUL, descripcion="", conteos=None):
        """Las secciones vacías no se dibujan: se listan juntas como «sin pendientes»."""
        if df is None or df.empty:
            self._vacias.append(titulo)
            return
        self._secciones.append((titulo, df, cols, anchos, color, descripcion, conteos or []))

    def _bloque_cifras(self):
        if not self._cifras:
            return []
        celdas, estilo = [], [("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                              ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F7FB")),
                              ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]
        for j, (nombre, valor, detalle, tono) in enumerate(self._cifras):
            color = TONOS.get(tono, AZUL)
            v = ParagraphStyle("v", parent=self.st["k_v"], textColor=colors.HexColor(color))
            celdas.append([Paragraph(esc(nombre), self.st["k_n"]), Paragraph(esc(texto(valor)), v),
                           Paragraph(esc(detalle), self.st["k_s"])])
            estilo.append(("LINEABOVE", (j, 0), (j, 0), 3, colors.HexColor(color)))
        n = len(celdas)
        sep = 0.25 * cm
        ancho = (self.ANCHO - sep * (n - 1)) / n
        filas = [[c for c in celdas]]
        t = Table(filas, colWidths=[ancho] * n)
        estilo += [("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                   ("LINEAFTER", (0, 0), (-2, 0), sep, colors.white)]
        t.setStyle(TableStyle(estilo))
        return [t, Spacer(1, 8)]

    def _pie(self, canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor(GRIS_TEXTO))
        canvas.drawString(1.2 * cm, 0.7 * cm, self.titulo)
        canvas.drawRightString(landscape(letter)[0] - 1.2 * cm, 0.7 * cm, f"Página {doc.page}")
        canvas.restoreState()

    # ------------------------------------------------------------------ armado
    def guardar(self, resumen=None, nota=""):
        st = self.st
        el = [Paragraph(esc(self.titulo), st["h1"]), Paragraph(esc(self.subtitulo), st["sub"]), Spacer(1, 10)]
        el += self._bloque_cifras()
        if self._vacias:
            el.append(Paragraph("<font color='#1E8449'><b>Sin pendientes:</b></font> "
                                + esc(" · ".join(self._vacias)), st["txt"]))
        if not self._secciones and not self._vacias:
            el.append(Paragraph("Sin pendientes.", st["txt"]))
        for i, (titulo, df, cols, anchos, color, desc, conteos) in enumerate(self._secciones, 1):
            cabeza = [Paragraph(f"{i}. {esc(titulo)} ({len(df)})", st["h2"])]
            if desc:
                cabeza.append(Paragraph(esc(desc), st["txt"]))
            for etiqueta_c, serie in conteos:
                cabeza.append(Paragraph(f"<b>{esc(etiqueta_c)}:</b> "
                                        + " · ".join(f"{esc(k)} {v}" for k, v in serie.items()), st["txt"]))
            cabeza.append(Spacer(1, 4))
            tabla = self.tabla(df, cols, anchos, color)
            el += [CondPageBreak(4 * cm), KeepTogether(cabeza + [tabla]) if len(df) <= 8 else
                   KeepTogether(cabeza)] + ([] if len(df) <= 8 else [tabla])
        if resumen is not None and not resumen.empty:
            filas = resumen.assign(Valor=resumen["Valor"].map(texto))
            el += [CondPageBreak(6 * cm), Paragraph("Detalle del cuadre", st["h2"]),
                   self.tabla(filas, ["Concepto", "Valor"], [17, 8.5])]
        if nota:
            el += [Spacer(1, 8), Paragraph(esc(nota), st["sub"])]
        doc = SimpleDocTemplate(self.ruta, pagesize=landscape(letter), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                                topMargin=1.1 * cm, bottomMargin=1.3 * cm, title=self.titulo, author=PDV)
        doc.build(el, onFirstPage=self._pie, onLaterPages=self._pie)
