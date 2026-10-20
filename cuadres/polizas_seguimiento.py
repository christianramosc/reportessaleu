# -*- coding: utf-8 -*-
"""
Seguimiento de pólizas de seguro para la app (página «Pólizas»).

Trabaja sobre la hoja «Contacto» del Excel que genera cuadres/polizas.py (una fila por crédito) y calcula en
Python lo mismo que las fórmulas del Excel, para cualquier fecha:

  cargar(pdfs_y_zips, excel_anterior)  → (base, excel_bytes, errores)   lee PDFs o reusa el Excel
  calcular(base, hoy, reglas)          → base + próxima oportunidad, fechas y estatus de hoy
  proyeccion(base, hoy, meses, reglas) → todas las oportunidades futuras (aniversarios, fin de crédito)
  aplicar_seguimiento(excel, cambios)  → Excel con Contactado / Fecha / Resultado / Notas actualizados
  pdf(...)                             → reporte de contactos del periodo

Reglas (días naturales antes de cada aniversario de la póliza):
  · póliza externa: a más tardar 45 días antes (después se cierra)
  · pago de contado: desde 30 días antes hasta el día anterior
  · llamar al menos 7 días antes del límite de externa
  · al terminar el crédito el auto es del cliente: puede contratar el seguro que quiera.
"""
import datetime as dt
import io
import os
import tempfile

import pandas as pd

from cuadres import polizas as pz

REGLAS = dict(externa=pz.DIAS_POLIZA_EXTERNA, contado=pz.DIAS_PAGO_CONTADO, aviso=pz.DIAS_AVISO_PREVIO)

# Estatus en el orden en que van pasando (los mismos textos que el Excel)
PROXIMO = "PRÓXIMO"
LLAMAR_EXT = "LLAMAR YA · PÓLIZA EXTERNA"
EXT_CERRADA = "EXTERNA CERRADA · ESPERAR CONTADO"
CONTADO = "VENTANA DE CONTADO ABIERTA"
FIN_PROX = "PRÓXIMO · TERMINA SU CRÉDITO"
FIN_LLAMAR = "LLAMAR YA · TERMINA SU CRÉDITO"
LIBRE = "CRÉDITO TERMINADO · CLIENTE LIBRE"
CONTACTADO = "CONTACTADO"
ESTATUS = {e: (q, fondo, letra) for e, q, fondo, letra in pz.ESTATUS_LISTA}
URGENTES = [LLAMAR_EXT, FIN_LLAMAR, CONTADO, LIBRE]
RESULTADOS = pz.RESULTADOS
CIERRES = ["Vendida (póliza externa)", "Pagará de contado"]

COLS_SEG = ["CONTACTADO", "FECHA CONTACTO", "RESULTADO", "NOTAS"]
HDR_CONTACTO = 5          # fila de encabezados en la hoja Contacto


# =============================================================================
# Carga
# =============================================================================
def _fecha(v):
    try:
        if v is None or v == "" or pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return pd.to_datetime(v, dayfirst=True).date()
    except Exception:
        return None


def leer_contacto(excel_bytes):
    """Hoja Contacto → DataFrame con las columnas fijas (las fórmulas se recalculan en Python)."""
    xl = pd.ExcelFile(io.BytesIO(excel_bytes))
    if "Contacto" not in xl.sheet_names:
        raise ValueError("El Excel no tiene la hoja «Contacto». Sube el Polizas_extraidas.xlsx nuevo (2 hojas).")
    df = pd.read_excel(xl, "Contacto", header=HDR_CONTACTO - 1, dtype={"TELÉFONO": str, "SERIE": str,
                                                                      "FOLIO CRÉDITO": str, "PÓLIZA ACTUAL": str})
    df = df[df["CLIENTE"].notna() & (df["CLIENTE"].astype(str).str.strip() != "")].copy()
    for c in ("PRIMERA OPORTUNIDAD", "FIN DEL CRÉDITO", "FECHA CONTACTO"):
        df[c] = df[c].map(_fecha)
    for c in ("TELÉFONO", "CORREO", "FOLIO CRÉDITO", "SERIE", "PÓLIZA ACTUAL", "RESULTADO", "NOTAS", "CONTACTADO",
              "MUNICIPIO", "DIRECCIÓN", "VEHÍCULO", "ASEGURADORA"):
        if c in df.columns:
            df[c] = df[c].map(lambda v: "" if pd.isna(v) else str(v).strip().removesuffix(".0"))
    df["PRIMA ANUAL ACTUAL"] = pd.to_numeric(df.get("PRIMA ANUAL ACTUAL"), errors="coerce")
    keep = ["CLIENTE", "TELÉFONO", "CORREO", "ASEGURADORA", "PÓLIZA ACTUAL", "PRIMA ANUAL ACTUAL", "VEHÍCULO",
            "SERIE", "FOLIO CRÉDITO", "MUNICIPIO", "DIRECCIÓN", "PRIMERA OPORTUNIDAD", "FIN DEL CRÉDITO"] + COLS_SEG
    return df[[c for c in keep if c in df.columns]].reset_index(drop=True)


def filas_de_excel(excel_bytes):
    """Reconstruye las pólizas (2 por crédito) desde la hoja «Base de datos» de un Excel ya generado, para
    poder juntarlas con PDFs nuevos sin volver a subir los PDFs viejos."""
    xl = pd.ExcelFile(io.BytesIO(excel_bytes))
    if "Base de datos" not in xl.sheet_names:
        return []
    texto = ["FOLIO CRÉDITO", "RFC", "TELÉFONO", "CP", "AÑO MODELO", "SERIE", "MOTOR", "PÓLIZA 1 (CONTADO)",
             "PÓLIZA 2 (RESTO DEL PLAZO)"]
    df = pd.read_excel(xl, "Base de datos", dtype={c: str for c in texto})
    limpio = lambda v: "" if v is None or (not isinstance(v, str) and pd.isna(v)) else str(v).strip().removesuffix(".0")
    num = lambda v: None if v is None or pd.isna(v) or v == "" else float(v)
    filas = []
    for _, r in df.iterrows():
        if not limpio(r.get("CLIENTE")):
            continue
        rfc = limpio(r.get("RFC"))
        comun = {
            "FOLIO CRÉDITO": limpio(r.get("FOLIO CRÉDITO")), "NOMBRE": limpio(r.get("CLIENTE")),
            "TIPO PERSONA": limpio(r.get("TIPO PERSONA")), "RFC": rfc,
            "FECHA NAC. (RFC)": pz.nacimiento_rfc(rfc) if len(rfc) == 13 else None, "SEXO": limpio(r.get("SEXO")),
            "TELÉFONO": limpio(r.get("TELÉFONO")), "CORREO": limpio(r.get("CORREO")),
            "DIRECCIÓN": limpio(r.get("DIRECCIÓN")), "MUNICIPIO": limpio(r.get("MUNICIPIO")),
            "ESTADO": limpio(r.get("ESTADO")), "CP": limpio(r.get("CP")),
            "DATOS DE VEHÍCULO": limpio(r.get("VEHÍCULO")), "MODELO MG": limpio(r.get("MODELO MG")),
            "AÑO MODELO": limpio(r.get("AÑO MODELO")), "SERIE": limpio(r.get("SERIE")), "MOTOR": limpio(r.get("MOTOR")),
            "ASEGURADORA": limpio(r.get("ASEGURADORA")), "PAQUETE": limpio(r.get("PAQUETE")),
            "SUMA ASEGURADA DM": num(r.get("SUMA ASEGURADA DM")) if not isinstance(r.get("SUMA ASEGURADA DM"), str)
            else r.get("SUMA ASEGURADA DM"),
            "DEDUCIBLE DM": num(r.get("DEDUCIBLE DM")), "DEDUCIBLE RT": num(r.get("DEDUCIBLE RT")),
            "RC TERCEROS": num(r.get("RC TERCEROS")),
            "BENEFICIARIO PREFERENTE": limpio(r.get("BENEFICIARIO PREFERENTE")), "AGENCIA": limpio(r.get("AGENCIA")),
            "ARCHIVO": limpio(r.get("ARCHIVO PDF")) or "(de Excel anterior)", "OBSERVACIONES": "",
        }
        for n, col in ((1, "PÓLIZA 1 (CONTADO)"), (2, "PÓLIZA 2 (RESTO DEL PLAZO)")):
            desde, hasta = _fecha(r.get(f"P{n} DESDE")), _fecha(r.get(f"P{n} HASTA"))
            if not (limpio(r.get(col)) and desde and hasta):
                continue
            anios = max(1, round((hasta - desde).days / 365.25))
            prima = num(r.get("PRIMA P1" if n == 1 else "PRIMA P2 (TOTAL)"))
            filas.append({**comun, "NO. PÓLIZA": limpio(r.get(col)), "VIGENCIA DESDE": desde, "VIGENCIA HASTA": hasta,
                          "PRIMA TOTAL": prima, "PRIMA ANUAL EST.": round(prima / anios, 2) if prima else None,
                          "AÑOS DE VIGENCIA": anios,
                          "TIPO DE PÓLIZA": "ANUAL" if anios == 1 else f"MULTIANUAL ({anios} AÑOS)"})
    return filas


def cargar(archivos, excel_anterior=None):
    """archivos: [(nombre, bytes)] de PDFs y/o ZIPs. excel_anterior: bytes del Polizas_extraidas.xlsx (opcional).
    Con PDFs: los lee, los junta con los créditos del Excel anterior (los PDFs nuevos mandan si se repite una
    póliza) y arma el Excel conservando el seguimiento. Sin PDFs: usa el Excel tal cual.
    Devuelve (base, excel_bytes, errores)."""
    if not archivos:
        if not excel_anterior:
            raise ValueError("Sube los PDFs / ZIP de las pólizas o el Excel Polizas_extraidas.xlsx.")
        return leer_contacto(excel_anterior), excel_anterior, []
    with tempfile.TemporaryDirectory() as tmp:
        carpeta = os.path.join(tmp, "pdfs")
        os.makedirs(carpeta)
        for nombre, datos in archivos:
            with open(os.path.join(carpeta, os.path.basename(nombre)), "wb") as f:
                f.write(datos)
        salida = os.path.join(tmp, "Polizas_extraidas.xlsx")
        if excel_anterior:                       # guardar_excel recupera de aquí el seguimiento capturado
            with open(salida, "wb") as f:
                f.write(excel_anterior)
        filas, errores = pz.extraer([carpeta], hilos=4)
        if not filas:
            raise ValueError("No se encontró ninguna póliza legible en los archivos.")
        if excel_anterior:                       # créditos que ya estaban en el Excel y no vienen en los PDFs
            nuevas = {(f.get("ASEGURADORA"), f.get("NO. PÓLIZA")) for f in filas}
            series = {f.get("SERIE") for f in filas}
            filas += [f for f in filas_de_excel(excel_anterior)
                      if (f.get("ASEGURADORA"), f.get("NO. PÓLIZA")) not in nuevas and f.get("SERIE") not in series]
        pz.guardar_excel(filas, errores, salida)
        with open(salida, "rb") as f:
            excel = f.read()
    return leer_contacto(excel), excel, errores


# =============================================================================
# Cálculo (mismo criterio que las fórmulas del Excel)
# =============================================================================
def _anios(d, n):
    return pz.sumar_anios(d, n)


def _siguiente_aniversario(base, hoy):
    """Primer aniversario de «base» (base, base+1 año…) posterior a hoy."""
    if base > hoy:
        return base
    n = hoy.year - base.year
    while _anios(base, n) <= hoy:
        n += 1
    while n > 0 and _anios(base, n - 1) > hoy:
        n -= 1
    return _anios(base, n)


def _total_oport(base, fin):
    if not fin:
        return None
    n = 0
    while _anios(base, n) < fin:
        n += 1
    return n


def estado(fila, hoy, reglas=REGLAS):
    """Próxima oportunidad de un crédito vista desde «hoy»."""
    E, C, A = reglas["externa"], reglas["contado"], reglas["aviso"]
    base, fin = fila["PRIMERA OPORTUNIDAD"], fila["FIN DEL CRÉDITO"]
    if base is None:
        return {}
    sig = _siguiente_aniversario(base, hoy)
    if fin and hoy >= fin:
        tipo, fecha = "CRÉDITO TERMINADO", sig
    elif fin and sig >= fin:
        tipo, fecha = "FIN DE CRÉDITO", fin
    else:
        tipo, fecha = "ANIVERSARIO", sig
    d = lambda n: fecha - dt.timedelta(days=n)
    llamar = d(E + A)
    r = {"TIPO": tipo, "PRÓXIMA FECHA": fecha, "LLAMAR DESDE": llamar,
         "ÚLTIMO DÍA EXTERNA": d(E) if tipo == "ANIVERSARIO" else None,
         "ABRE CONTADO": d(C) if tipo == "ANIVERSARIO" else None,
         "ÚLTIMO DÍA CONTADO": d(1) if tipo == "ANIVERSARIO" else None}
    tot = _total_oport(base, fin)
    if tipo == "ANIVERSARIO":
        n = (fecha.year - base.year) + 1
        r["OPORTUNIDAD"] = f"{n} de {tot if tot else '?'}"
    else:
        r["OPORTUNIDAD"] = "Fin del crédito" if tipo == "FIN DE CRÉDITO" else "Auto ya liberado"
    contactado = str(fila.get("CONTACTADO", "")).strip().lower() in ("sí", "si")
    fc = fila.get("FECHA CONTACTO")
    if contactado and (fc is None or fc > _anios(fecha, -1)):
        est = CONTACTADO
    elif tipo == "CRÉDITO TERMINADO":
        est = LIBRE
    elif tipo == "FIN DE CRÉDITO":
        est = FIN_LLAMAR if hoy >= llamar else FIN_PROX
    elif hoy >= d(C):
        est = CONTADO
    elif hoy > d(E):
        est = EXT_CERRADA
    elif hoy >= llamar:
        est = LLAMAR_EXT
    else:
        est = PROXIMO
    r["ESTATUS"] = est
    r["QUÉ HACER"] = ESTATUS[est][0]
    if tipo == "CRÉDITO TERMINADO":
        r["DÍAS QUE QUEDAN"] = None
        r["CIERRA EN"] = ""
    elif tipo == "FIN DE CRÉDITO":
        r["DÍAS QUE QUEDAN"] = (fecha - hoy).days
        r["CIERRA EN"] = f"{(fecha - hoy).days} días para el fin del crédito"
    elif est in (EXT_CERRADA, CONTADO) or (est == CONTACTADO and hoy > d(E)):
        r["DÍAS QUE QUEDAN"] = (d(1) - hoy).days
        r["CIERRA EN"] = f"{(d(1) - hoy).days} días para el último día de contado"
    else:
        r["DÍAS QUE QUEDAN"] = (d(E) - hoy).days
        r["CIERRA EN"] = f"{(d(E) - hoy).days} días para el límite de externa"
    return r


def calcular(base, hoy=None, reglas=REGLAS):
    hoy = hoy or dt.date.today()
    calc = pd.DataFrame([estado(f, hoy, reglas) for _, f in base.iterrows()], index=base.index)
    df = pd.concat([base, calc], axis=1)
    orden = {e: i for i, e in enumerate([LLAMAR_EXT, FIN_LLAMAR, CONTADO, LIBRE, EXT_CERRADA, CONTACTADO,
                                         FIN_PROX, PROXIMO])}
    df["_o"] = df["ESTATUS"].map(orden).fillna(99)
    return df.sort_values(["_o", "LLAMAR DESDE", "CLIENTE"]).drop(columns="_o").reset_index(drop=True)


def proyeccion(base, hoy=None, meses=12, reglas=REGLAS):
    """Todas las oportunidades cuyo «llamar desde» cae de hoy a +meses: aniversarios antes del fin del crédito y
    el propio fin del crédito. Una fila por oportunidad."""
    hoy = hoy or dt.date.today()
    hasta = (pd.Timestamp(hoy) + pd.DateOffset(months=meses)).date()
    E, A = reglas["externa"], reglas["aviso"]
    filas = []
    for _, f in base.iterrows():
        b, fin = f["PRIMERA OPORTUNIDAD"], f["FIN DEL CRÉDITO"]
        if b is None:
            continue
        tot = _total_oport(b, fin)
        eventos, n = [], 0
        while True:
            a = _anios(b, n)
            if fin and a >= fin:
                break
            eventos.append(("Aniversario", a, f"{n + 1} de {tot if tot else '?'}"))
            n += 1
            if a - dt.timedelta(days=E + A) > hasta:
                break
        if fin:
            eventos.append(("Fin de crédito", fin, "Fin del crédito"))
        for tipo, fecha, op in eventos:
            llamar = fecha - dt.timedelta(days=E + A)
            if fecha > hoy and llamar <= hasta:
                filas.append({"MES": dt.date(llamar.year, llamar.month, 1) if llamar > hoy
                              else dt.date(hoy.year, hoy.month, 1),
                              "LLAMAR DESDE": llamar, "FECHA": fecha, "TIPO": tipo, "OPORTUNIDAD": op,
                              "CLIENTE": f["CLIENTE"], "TELÉFONO": f["TELÉFONO"], "CORREO": f["CORREO"],
                              "ASEGURADORA": f["ASEGURADORA"], "VEHÍCULO": f["VEHÍCULO"], "SERIE": f["SERIE"],
                              "PRIMA ANUAL ACTUAL": f["PRIMA ANUAL ACTUAL"]})
    return pd.DataFrame(filas).sort_values(["LLAMAR DESDE", "CLIENTE"]).reset_index(drop=True) if filas else \
        pd.DataFrame(columns=["MES", "LLAMAR DESDE", "FECHA", "TIPO", "OPORTUNIDAD", "CLIENTE"])


# =============================================================================
# Guardar el seguimiento en el Excel (sin tocar sus fórmulas)
# =============================================================================
def aplicar_seguimiento(excel_bytes, seg):
    """seg: DataFrame con SERIE + CONTACTADO, FECHA CONTACTO, RESULTADO, NOTAS. Devuelve los bytes del Excel."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb["Contacto"]
    hdr = {c.value: c.column for c in ws[HDR_CONTACTO]}
    por_serie = {str(s): r for s, r in zip(seg["SERIE"], seg.to_dict("records"))}
    for fila in range(HDR_CONTACTO + 1, ws.max_row + 1):
        serie = str(ws.cell(fila, hdr["SERIE"]).value or "")
        if serie in por_serie:
            r = por_serie[serie]
            for c in COLS_SEG:
                v = r.get(c)
                if v is None or (isinstance(v, float) and pd.isna(v)) or v == "":
                    v = None
                elif c == "FECHA CONTACTO":
                    v = _fecha(v)
                ws.cell(fila, hdr[c]).value = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# PDF
# =============================================================================
def pdf(df_hoy, df_periodo, titulo_periodo, hoy=None):
    """Reporte: a quién llamar hoy (urgentes) y las oportunidades del periodo elegido."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Table, TableStyle)
    from cuadres.reporte_saleu import NAVY, TINTA_3, LINEA, _e, _estilos, _fuentes
    from cuadres.salidas import telefono

    _fuentes()
    st = _estilos()
    hoy = hoy or dt.date.today()
    ancho, alto = landscape(letter)
    m = 34
    util = ancho - 2 * m
    buf = io.BytesIO()

    def fondo(c, doc):
        c.saveState()
        c.setFillColor(colors.HexColor(NAVY))
        c.rect(0, alto - 54, ancho, 54, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Archivo-ExtraBold", 16)
        c.drawString(m, alto - 30, "Pólizas de seguro · clientes a contactar")
        c.setFont("Archivo", 8.5)
        c.setFillColor(colors.HexColor("#C9D6EA"))
        c.drawString(m, alto - 44, f"Póliza externa hasta 45 días antes del aniversario · contado de 30 días antes al día "
                                   f"anterior · corte {hoy:%d/%m/%Y}")
        c.setFont("Archivo", 7)
        c.setFillColor(colors.HexColor(TINTA_3))
        c.drawString(m, 16, "MG Colima · Fuente: pólizas de Crédito Click")
        c.drawRightString(ancho - m, 16, f"Página {c.getPageNumber()}")
        c.restoreState()

    doc = BaseDocTemplate(buf, pagesize=(ancho, alto), leftMargin=m, rightMargin=m, topMargin=66, bottomMargin=30,
                          title="Pólizas a contactar", author="MG Colima")
    doc.addPageTemplates([PageTemplate("p", [Frame(m, 30, util, alto - 96, 0, 0, 0, 0)], onPage=fondo)])
    f = lambda d: d.strftime("%d/%m/%Y") if isinstance(d, dt.date) else ""
    P = lambda t, s="celda": Paragraph(_e(t), st[s])

    def tabla(df, cols, anchos, color_col=None):
        filas = [[P(c, "enc") for c in cols]]
        estilos = [("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor(NAVY)),
                   ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor(LINEA)),
                   ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                   ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        for i, (_, r) in enumerate(df.iterrows(), start=1):
            fila = []
            for c in cols:
                v = r.get(c, "")
                if isinstance(v, dt.date):
                    v = f(v)
                elif c == "TELÉFONO":
                    v = telefono(v) or ""
                elif isinstance(v, float):
                    v = "" if pd.isna(v) else (f"${v:,.0f}" if "PRIMA" in c else f"{v:.0f}")
                fila.append(P(v, "nombre" if c == "CLIENTE" else "celda"))
            filas.append(fila)
            if color_col and r.get(color_col) in ESTATUS:
                estilos.append(("BACKGROUND", (0, i), (0, i), colors.HexColor("#" + ESTATUS[r[color_col]][1])))
        t = Table(filas, colWidths=[util * a for a in anchos], repeatRows=1)
        t.setStyle(TableStyle(estilos))
        return t

    el = [Paragraph("Para contactar ya", st["h2"])]
    urg = df_hoy[df_hoy["ESTATUS"].isin(URGENTES)]
    if urg.empty:
        el.append(P("Hoy no hay clientes con ventana abierta.", "txt"))
    else:
        el.append(tabla(urg, ["ESTATUS", "CLIENTE", "TELÉFONO", "CIERRA EN", "PRÓXIMA FECHA", "OPORTUNIDAD",
                              "ASEGURADORA", "PRIMA ANUAL ACTUAL", "VEHÍCULO"],
                        [.15, .17, .095, .14, .08, .065, .085, .07, .145], "ESTATUS"))
    el += [Spacer(1, 14), Paragraph(f"Oportunidades de {titulo_periodo}", st["h2"])]
    if df_periodo.empty:
        el.append(P("Sin oportunidades en el periodo.", "txt"))
    else:
        el.append(P(f"{len(df_periodo)} oportunidades · prima anual en juego "
                    f"${df_periodo['PRIMA ANUAL ACTUAL'].sum():,.0f}", "sub"))
        el.append(tabla(df_periodo, ["LLAMAR DESDE", "FECHA", "TIPO", "OPORTUNIDAD", "CLIENTE", "TELÉFONO",
                                     "CORREO", "ASEGURADORA", "PRIMA ANUAL ACTUAL"],
                        [.08, .08, .09, .07, .2, .095, .19, .085, .09]))
    doc.build(el)
    return buf.getvalue()
