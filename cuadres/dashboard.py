# -*- coding: utf-8 -*-
"""
Dashboard de la pestaña «Resumen» del Sheet de seguimiento.

construir(base) reescribe «Resumen» con:
  · indicadores (KPI) del año elegido en B2,
  · «Lectura del periodo»: frases que interpretan la situación (fórmulas, se actualizan solas),
  · gráficas (rechazos por mes y financiera, seguimiento, financieras, asesores, Sale-U, bitácora, pendientes),
  · debajo, las tablas de datos de las que salen (todo con fórmulas sobre «MG Colima Rechazos», «Bitacora»
    y «Revision»).
Funciona igual en Google Sheets (gráficas nativas) y en el modo Excel (gráficas de openpyxl).
El año de B2 se conserva al reconstruir.
"""
import datetime as dt

from cuadres.base import TAB_RECHAZOS, TAB_BITACORA, TAB_CAPTURAS, BaseGoogle, letra, norm
from cuadres.actualizar import C, FUENTE_SHEET, _catalogo_asesores, _asesor, es_valor_catalogo
from cuadres.fuentes import BITACORA_COLS, ETIQUETA_MES, CAP_COLS, CAP_EXTRA

TAB_RESUMEN = "Resumen"
TAB_REVISION = "Revision"
NCOLS = 12                                        # A:L
ALTO_FILAS = 16                                   # filas que ocupa cada renglón de gráficas
FILAS = 2000                                      # alcance de las fórmulas

# Colores (paleta validada: categóricos en orden fijo; estados reservados para estados)
SERIE = ["#2a78d6", "#eb6834", "#1baf7a"]        # azul, naranja, aqua
SECUENCIA = ["#86b6ef", "#2a78d6", "#104281"]    # rampa azul ordinal (embudo)
ESTADO = dict(bien="#0ca30c", aviso="#fab219", serio="#ec835a", critico="#d03b3b", neutro="#8a8984")
TINTA, TINTA_2, TINTA_3 = "#0b0b0b", "#52514e", "#8a8984"
AZUL_TITULO, FONDO_KPI, FONDO_ENC, LINEA = "#0B5394", "#F3F6FA", "#666666", "#D9D9D9"
FUENTE = "Calibri"

MESES = [ETIQUETA_MES[m].capitalize() for m in range(1, 13)]
LEYENDA = [("Pendiente: falta un dato (lee la nota de la celda)", "#D9D2E9"),
           ("Agregado desde la bitácora", "#CFE2F3"), ("Folio repetido", "#F9CB9C"),
           ("Seguimiento vencido", "#F4CCCC"), ("Vence en ≤15 días / falta correo, Sale-U o si contestó", "#FCE5CD"),
           ("Seguimiento hecho / Registrado", "#D9EAD3"), ("Reingreso: el celular está en otra solicitud", "#FFF2CC")]


# =============================================================================
# Hoja lógica: celdas con estilo, combinaciones, gráficas
# =============================================================================
class _Hoja:
    def __init__(self):
        self.celdas = {}          # (fila, col) -> (valor, rol, formato)
        self.merges = []          # (f0, c0, f1, c1) 1-based inclusive
        self.graficas = []
        self.alertas = []         # (fila, col, condición): rojo si se cumple

    def pon(self, f, c, valor, rol="td", formato=None):
        self.celdas[(f, c)] = (valor, rol, formato)

    def une(self, f0, c0, f1, c1):
        self.merges.append((f0, c0, f1, c1))

    def tabla(self, fila, encabezados, filas, formatos=None, total=False):
        """Escribe una tabla (encabezado + filas). Devuelve (fila del encabezado, primera, última fila de datos)."""
        for j, h in enumerate(encabezados, start=1):
            self.pon(fila, j, h, "th")
        for i, valores in enumerate(filas, start=1):
            es_total = total and i == len(filas)
            for j, v in enumerate(valores, start=1):
                fmt = (formatos or {}).get(j)
                self.pon(fila + i, j, v, "total" if es_total else ("td_nombre" if j == 1 else "td"), fmt)
        return fila, fila + 1, fila + len(filas)


def _rng(hoja, col, fila0=2, fila1=FILAS):
    return f"'{hoja}'!${col}${fila0}:${col}${fila1}"


def _entre(fecha_rng, mes0, mes1=None):
    """Condición COUNTIFS: fecha dentro de [mes0, mes1) del año de $B$2 (mes1=None → solo ese mes)."""
    mes1 = mes1 or mes0 + 1
    return f'{fecha_rng},">="&DATE($B$2,{mes0},1),{fecha_rng},"<"&DATE($B$2,{mes1},1)'


def _conteo(nombre_col, columnas, df, anio):
    """{valor: rechazos en el año} de una columna de MG Colima Rechazos (para ordenar las tablas)."""
    col, fcol = columnas.get(nombre_col), columnas.get("fecha")
    out = {}
    if not col or df.empty:
        return out
    from cuadres.fuentes import to_date
    for _, r in df.iterrows():
        v = str(r[col]).strip().upper() if r[col] not in (None, "") else ""
        f = to_date(r[fcol]) if fcol else None
        if v and f and f.year == anio:
            out[v] = out.get(v, 0) + 1
        elif v:
            out.setdefault(v, 0)
    return out


# =============================================================================
# Construcción del contenido
# =============================================================================
def contenido(base):
    """Arma la hoja lógica del dashboard a partir de las columnas reales del Sheet."""
    pestanas = set(base.pestanas())
    df = base.leer(TAB_RECHAZOS) if TAB_RECHAZOS in pestanas else None
    enc = {norm(h): h for h in (df.columns if df is not None else [])}
    letra_de = lambda h: letra(base.col(TAB_RECHAZOS, h)) if base.col(TAB_RECHAZOS, h) else None
    L = dict(fecha=letra_de(C["fecha"]), fin=letra_de(C["fin"]), asesor=letra_de(C["asesor"]),
             saleu=letra_de(C["saleu"]), fseg=letra_de(C["fseg"]), correo=letra_de(C["correo"]),
             folio=letra_de(C["folio"]), nombre=letra_de(C["nombre"]),
             hecho=letra_de("FECHA REALIZADO 1"), contesta=letra_de("CONTESTA 1"))
    if not (L["fecha"] and L["nombre"]):
        raise ValueError(f"'{TAB_RECHAZOS}' no tiene las columnas {C['fecha']} y {C['nombre']}.")
    R = lambda k: _rng(TAB_RECHAZOS, L[k])
    hay_bit = TAB_BITACORA in pestanas
    B = {k: _rng(TAB_BITACORA, letra(BITACORA_COLS.index(h) + 1)) for k, h in
         dict(fecha="Fecha", estatus="Estatus solicitud", banco="Banco solicitud", folio="Folio",
              fuente="Fuente en la bitácora", final="Estatus final", saleu="Registro en Sale-U",
              id_rech="ID en 'MG Colima Rechazos'").items()}

    anio = _anio_actual(base)
    h = _Hoja()

    # ------------------------------------------------------------------ encabezado
    h.pon(1, 1, "Dashboard de rechazos y seguimiento — MG Colima", "titulo"); h.une(1, 1, 1, NCOLS)
    h.pon(2, 1, "Año:", "etiqueta"); h.pon(2, 2, anio, "anio")
    h.pon(2, 3, "← cambia el año para ver otro periodo. Todo se calcula solo desde 'MG Colima Rechazos', "
                "'Bitacora' y 'Revision'.", "nota"); h.une(2, 3, 2, NCOLS)

    # ------------------------------------------------------------------ tablas de datos (abajo)
    # filas: 1-6 encabezado y KPI · 8.. lectura · luego gráficas (16 filas por renglón de 2) · luego tablas
    n_lectura = 6 + (2 if hay_bit else 0)
    fila_graficas = 10 + n_lectura + 1
    n_graficas = 8 if hay_bit else 6
    f = fila_graficas + 1 + ((n_graficas + 1) // 2) * ALTO_FILAS + 4
    h.pon(f - 2, 1, "Datos del dashboard", "seccion"); h.une(f - 2, 1, f - 2, NCOLS)
    h.pon(f - 1, 1, "Rechazos por mes", "subseccion"); h.une(f - 1, 1, f - 1, NCOLS)
    filas_mes = []
    for m in range(1, 13):
        fila = f + m
        cond = _entre(R("fecha"), m)
        filas_mes.append([
            MESES[m - 1], f"=COUNTIFS({cond})",
            f'=COUNTIFS({cond},{R("fin")},"CETELEM")' if L["fin"] else 0,
            f'=COUNTIFS({cond},{R("fin")},"BBVA")' if L["fin"] else 0,
            f"=B{fila}-C{fila}-D{fila}",
            f'=COUNTIFS({cond},{R("hecho")},"<>")' if L["hecho"] else 0,
            f'=COUNTIFS({cond},{R("contesta")},"Si Responde")' if L["contesta"] else 0,
            f'=IF(F{fila}=0,"",G{fila}/F{fila})',
            f'=COUNTIFS({cond},{R("saleu")},"Registrado")' if L["saleu"] else 0,
            f"=B{fila}-I{fila}"])
    t = f + 13
    filas_mes.append(["Total"] + [f"=SUM({c}{f + 1}:{c}{f + 12})" for c in "BCDEFG"]
                     + [f'=IF(F{t}=0,"",G{t}/F{t})', f"=SUM(I{f + 1}:I{f + 12})", f"=SUM(J{f + 1}:J{f + 12})"])
    h.tabla(f, ["Mes", "Rechazos", "CETELEM", "BBVA", "Otras financieras", "Con 1er seguimiento",
                "Contestaron (1er)", "% contacto efectivo", "Registrados en Sale-U", "Sin registro en Sale-U"],
            filas_mes, {8: "pct"}, total=True)
    T1 = dict(enc=f, ini=f + 1, fin=f + 12, tot=t)

    # Financieras y asesores (año), lado a lado
    f = t + 3
    columnas = {"fecha": enc.get(norm(C["fecha"])), "fin": enc.get(norm(C["fin"])),
                "asesor": enc.get(norm(C["asesor"]))}
    cnt_fin = _conteo("fin", columnas, df, anio) if df is not None else {}
    for v in _catalogo(base, "Financiera"):
        cnt_fin.setdefault(v.upper(), 0)
    fins = sorted(cnt_fin, key=lambda k: (-cnt_fin[k], k))
    catalogo = _catalogo_asesores(base)
    cnt_as = {}
    for k, v in (_conteo("asesor", columnas, df, anio) if df is not None else {}).items():
        nombre = (_asesor(k, catalogo) or k).upper() if catalogo else k
        cnt_as[nombre] = cnt_as.get(nombre, 0) + v
    for v in catalogo:
        cnt_as.setdefault(v.upper(), 0)
    asesores = sorted(cnt_as, key=lambda k: (-cnt_as[k], k))
    anio_cond = _entre(R("fecha"), 1, 13)
    h.pon(f - 1, 1, "Rechazos por financiera (año)", "subseccion"); h.une(f - 1, 1, f - 1, 4)
    h.pon(f - 1, 6, "Rechazos por asesor (año)", "subseccion"); h.une(f - 1, 6, f - 1, NCOLS)
    for j, txt in enumerate(["Financiera", "Rechazos", "% del año"], start=1):
        h.pon(f, j, txt, "th")
    for i, nombre in enumerate(fins, start=1):
        h.pon(f + i, 1, nombre, "td_nombre")
        h.pon(f + i, 2, f'=COUNTIFS({anio_cond},{R("fin")},A{f + i})' if L["fin"] else 0)
        h.pon(f + i, 3, f'=IF($B${T1["tot"]}=0,"",B{f + i}/$B${T1["tot"]})', "td", "pct")
    for j, txt in enumerate(["Asesor", "", "Nombre completo", "", "", "Rechazos", "% del año"], start=6):
        h.pon(f, j, txt, "th")
    h.une(f, 6, f, 7); h.une(f, 8, f, 10)
    for i, nombre in enumerate(asesores, start=1):
        h.pon(f + i, 6, nombre_corto(nombre), "td_nombre"); h.une(f + i, 6, f + i, 7)
        h.pon(f + i, 8, nombre, "td_nombre"); h.une(f + i, 8, f + i, 10)
        h.pon(f + i, 11, f'=COUNTIFS({anio_cond},{R("asesor")},H{f + i})' if L["asesor"] else 0)
        h.pon(f + i, 12, f'=IF($B${T1["tot"]}=0,"",K{f + i}/$B${T1["tot"]})', "td", "pct")
    T2 = dict(enc=f, ini=f + 1, fin=f + len(fins))
    T3 = dict(enc=f, ini=f + 1, fin=f + len(asesores))

    # Pendientes + leyenda
    f = max(T2["fin"], T3["fin"]) + 3
    h.pon(f - 1, 1, "Pendientes (toda la hoja, sin importar el año)", "subseccion"); h.une(f - 1, 1, f - 1, 4)
    folio = R("folio")
    pend = [("Seguimiento vencido y sin hacer",
             f'=COUNTIFS({R("fseg")},"<="&TODAY(),{R("hecho")},"")' if L["fseg"] and L["hecho"] else 0),
            ("Vence en los próximos 15 días",
             f'=COUNTIFS({R("fseg")},">"&TODAY(),{R("fseg")},"<="&(TODAY()+15),{R("hecho")},"")'
             if L["fseg"] and L["hecho"] else 0),
            ("Filas con folio repetido",
             f'=SUMPRODUCT(({folio}<>"")*(COUNTIF({folio},{folio})>1))' if L["folio"] else 0),
            ("Registros sin correo", f'=COUNTIFS({R("nombre")},"<>",{R("correo")},"")' if L["correo"] else 0),
            ("CETELEM sin folio", f'=COUNTIFS({R("fin")},"CETELEM",{folio},"",{R("nombre")},"<>")'
             if L["fin"] and L["folio"] else 0),
            ("Pendientes en la pestaña Revision",
             f"=COUNTIFS('{TAB_REVISION}'!$B$2:$B$1000,\"Pendiente\")" if TAB_REVISION in pestanas else 0)]
    h.pon(f, 1, "Pendiente", "th"); h.une(f, 1, f, 3); h.pon(f, 4, "Cantidad", "th")
    for i, (txt, formula) in enumerate(pend, start=1):
        h.pon(f + i, 1, txt, "td_nombre"); h.une(f + i, 1, f + i, 3); h.pon(f + i, 4, formula)
        if i in (1, 2):
            h.alertas.append((f + i, 4, ">0"))
    T4 = dict(enc=f, ini=f + 1, fin=f + len(pend))
    h.pon(f - 1, 6, "Leyenda de colores de 'MG Colima Rechazos'", "subseccion"); h.une(f - 1, 6, f - 1, NCOLS)
    for i, (txt, color) in enumerate(LEYENDA):
        h.pon(f + i, 6, txt, "leyenda:" + color); h.une(f + i, 6, f + i, NCOLS)

    # Bitácora por mes y cuadre contra Cetelem
    T5 = T6 = None
    f = T4["fin"] + 4
    if hay_bit:
        h.pon(f - 1, 1, "Bitácora por mes: capturas de Cetelem y Bitácora de F&I (todas las financieras y "
                        "estatus)", "subseccion")
        h.une(f - 1, 1, f - 1, NCOLS)
        filas_bit = []
        de_bitacoras = f'{B["fuente"]},"<>{FUENTE_SHEET[0]}*"'      # sin los rechazos copiados del Sheet
        for m in range(1, 13):
            fila, cond = f + m, _entre(B["fecha"], m) + "," + de_bitacoras
            est = lambda e: f'COUNTIFS({cond},{B["estatus"]},"{e}")'
            filas_bit.append([MESES[m - 1], f"=COUNTIFS({cond})", "=" + est("Rechazado"), "=" + est("Aprobado"),
                              "=" + est("Contrapropuesta"), "=" + est("Análisis"), "=" + est("Formalizado"),
                              f'=IF(B{fila}=0,"",C{fila}/B{fila})'])
        t = f + 13
        filas_bit.append(["Total"] + [f"=SUM({c}{f + 1}:{c}{f + 12})" for c in "BCDEFG"]
                         + [f'=IF(B{t}=0,"",C{t}/B{t})'])
        h.tabla(f, ["Mes", "Solicitudes", "Rechazadas", "Aprobadas", "Contrapropuesta", "En análisis", "Formalizadas",
                    "% rechazo"], filas_bit, {8: "pct"}, total=True)
        T5 = dict(enc=f, ini=f + 1, fin=f + 12, tot=t)

        f = t + 4
        h.pon(f - 1, 1, "Cuadre contra las capturas de Cetelem (rechazos con folio)", "subseccion")
        h.une(f - 1, 1, f - 1, NCOLS)
        filas_cc = []
        if TAB_CAPTURAS in pestanas:
            # un mes cuenta como «con capturas» si se subió la captura de Cetelem de ese mes (columna Hoja), no
            # porque haya solicitudes con esa fecha (la captura de junio trae solicitudes de mayo)
            hoja_cap = _rng(TAB_CAPTURAS, letra((CAP_COLS + CAP_EXTRA).index("Hoja") + 1))
            sin_captura = lambda m: f'COUNTIF({hoja_cap},"{ETIQUETA_MES[m]} "&$B$2)=0'
        else:
            sin_captura = lambda m: f'COUNTIFS({_entre(B["fecha"], m)},{B["fuente"]},"*CETELEM*")=0'
        for m in range(1, 13):
            fila = f + m
            cond_b = _entre(B["fecha"], m)
            cond_r = _entre(R("fecha"), m)
            filas_cc.append([
                MESES[m - 1],
                f'=COUNTIFS({cond_b},{B["estatus"]},"Rechazado",{B["banco"]},"CETELEM",{B["folio"]},"<>",'
                f'{B["fuente"]},"*CETELEM*")',
                f'=COUNTIFS({cond_r},{R("fin")},"CETELEM",{folio},"<>")' if L["fin"] and L["folio"] else 0,
                f'=IF({sin_captura(m)},"sin capturas",C{fila}-B{fila})'])
            h.alertas.append((fila, 4, "num<>0"))      # «sin capturas» no es diferencia
        h.tabla(f, ["Mes", "Rechazados en Cetelem", "En 'MG Colima Rechazos'", "Diferencia"], filas_cc)
        T6 = dict(enc=f, ini=f + 1, fin=f + 12)
        h.pon(T6["fin"] + 1, 1, "Si la diferencia no es 0, filtra la pestaña Bitacora por 'ID en MG Colima "
                                "Rechazos' vacía.", "nota"); h.une(T6["fin"] + 1, 1, T6["fin"] + 1, NCOLS)

    # ------------------------------------------------------------------ KPI (filas 4-6)
    tot = T1["tot"]
    pend_c = lambda i: f"D{T4['ini'] + i}"
    # Seguimiento al día: de los rechazos del año que ya cumplieron 180 días, cuántos tienen 1er seguimiento
    if L["fseg"] and L["hecho"]:
        ya_toca = f'COUNTIFS({anio_cond},{R("fseg")},"<="&TODAY())'
        al_dia = f'COUNTIFS({anio_cond},{R("fseg")},"<="&TODAY(),{R("hecho")},"<>")'
    else:
        ya_toca = al_dia = "0"
    # Registro en Sale-U: Rechazados (MG Colima Rechazos) + Aprobados del año de la Bitacora. Un aprobado ligado a
    # una fila de Rechazos (columna ID) ya se cuenta ahí.
    if hay_bit:
        ap = f'{_entre(B["fecha"], 1, 13)},{B["final"]},"Aprobado",{B["id_rech"]},""'
        ap_tot, ap_reg = f"COUNTIFS({ap})", f'COUNTIFS({ap},{B["saleu"]},"Registrado")'
        ap_no = f'COUNTIFS({ap},{B["saleu"]},"No Registrado")'
    else:
        ap_tot = ap_reg = ap_no = "0"
    su_total, su_reg, su_no = f"(B{tot}+{ap_tot})", f"(I{tot}+{ap_reg})", f"(J{tot}+{ap_no})"
    kpis = [("Rechazos en el año", f"=B{tot}", None,
             f'="CETELEM "&C{tot}&" · BBVA "&D{tot}&" · otras "&E{tot}'),
            ("Seguimiento al día", f'=IF({ya_toca}=0,1,{al_dia}/{ya_toca})', "pct",
             f'={al_dia}&" de "&{ya_toca}&" que ya cumplieron 180 días"'),
            ("Contacto efectivo", f'=IF(F{tot}=0,0,G{tot}/F{tot})', "pct", f'=G{tot}&" contestaron el 1er contacto"'),
            ("Registrados en Sale-U", f"=IF({su_total}=0,0,{su_reg}/{su_total})", "pct",
             f'={su_no}&" sin registro: "&J{tot}&" rechazados · "&{ap_no}&" aprobados"'),
            ("Seguimientos vencidos", f"={pend_c(0)}", None, f'={pend_c(1)}&" vencen en 15 días"'),
            ("Pendientes en Revision", f"={pend_c(5)}", None, '="por resolver en la pestaña Revision"')]
    for i, (titulo, valor, fmt, sub) in enumerate(kpis):
        c = 1 + 2 * i
        h.pon(4, c, titulo, "kpi_titulo"); h.une(4, c, 4, c + 1)
        h.pon(5, c, valor, "kpi_valor", fmt); h.une(5, c, 5, c + 1)
        h.pon(6, c, sub, "kpi_sub"); h.une(6, c, 6, c + 1)
    h.alertas.append((5, 9, ">0"))

    # ------------------------------------------------------------------ lectura del periodo (fórmulas)
    meses, b_mes = f"$A${T1['ini']}:$A${T1['fin']}", f"$B${T1['ini']}:$B${T1['fin']}"
    fn, fv = f"$A${T2['ini']}:$A${T2['fin']}", f"$B${T2['ini']}:$B${T2['fin']}"
    an, av = f"$F${T3['ini']}:$F${T3['fin']}", f"$K${T3['ini']}:$K${T3['fin']}"
    sin = f'"Sin rechazos registrados en "&$B$2&"."'
    lectura = [
        f'=IF(B{tot}=0,{sin},"En "&$B$2&" van "&B{tot}&" rechazos. El mes con más fue "&INDEX({meses},'
        f'MATCH(MAX({b_mes}),{b_mes},0))&" con "&MAX({b_mes})&" ("&TEXT(MAX({b_mes})/B{tot},"0%")&" del año).")',
        f'=IF(B{tot}=0,"","Por financiera: CETELEM "&TEXT(C{tot}/B{tot},"0%")&" ("&C{tot}&"), BBVA "&'
        f'TEXT(D{tot}/B{tot},"0%")&" ("&D{tot}&") y otras "&TEXT(E{tot}/B{tot},"0%")&" ("&E{tot}&").")',
        f'=IF(B{tot}=0,"","El asesor con más rechazos es "&INDEX({an},MATCH(MAX({av}),{av},0))&" ("&'
        f'MAX({av})&", "&TEXT(MAX({av})/B{tot},"0%")&" del año).")',
        f'=IF(B{tot}=0,"",IF({ya_toca}=0,"Ningún rechazo de "&$B$2&" ha cumplido 180 días todavía.",'
        f'"De los "&{ya_toca}&" rechazos que ya cumplieron 180 días, "&{al_dia}&" tienen primer seguimiento ("&'
        f'TEXT({al_dia}/{ya_toca},"0%")&")"&IF({ya_toca}>{al_dia},"; faltan "&({ya_toca}-{al_dia})&'
        f'" por contactar","")&".")&IF(B{tot}>{ya_toca}," Otros "&(B{tot}-{ya_toca})&'
        f'" todavía no cumplen 180 días.","")&IF(F{tot}=0,""," Contestó el "&TEXT(G{tot}/F{tot},"0%")&'
        f'" de los contactados."))',
        f'=IF({su_total}=0,"","El "&TEXT({su_reg}/{su_total},"0%")&" de los rechazados y aprobados del año está '
        f'registrado en Sale-U"&IF({su_no}>0,"; faltan "&{su_no}&" por registrar ("&J{tot}&" rechazados y "&'
        f'{ap_no}&" aprobados).","."))',
        f'={pend_c(0)}&" seguimientos vencidos sin hacer y "&{pend_c(1)}&" vencen en los próximos 15 días. Hay "&'
        f'{pend_c(5)}&" pendientes en Revision y "&{pend_c(3)}&" registros sin correo."',
    ]
    if T5:
        bt, pc, bm = T5["tot"], f"$H${T5['ini']}:$H${T5['fin']}", f"$A${T5['ini']}:$A${T5['fin']}"
        lectura.append(
            f'=IF(B{bt}=0,"La bitácora no tiene solicitudes en "&$B$2&".","En la bitácora hay "&B{bt}&'
            f'" solicitudes: "&TEXT(C{bt}/B{bt},"0%")&" rechazadas, "&TEXT(G{bt}/B{bt},"0%")&" formalizadas y "&'
            f'TEXT((D{bt}+E{bt})/B{bt},"0%")&" aprobadas o en contrapropuesta. El mes con mayor tasa de rechazo fue "&'
            f'INDEX({bm},MATCH(MAX({pc}),{pc},0))&" ("&TEXT(MAX({pc}),"0%")&").")')
    if T6:
        dif = f"$D${T6['ini']}:$D${T6['fin']}"
        absd = f'(SUMIF({dif},">0")-SUMIF({dif},"<0"))'             # suma de diferencias (ignora "sin capturas")
        lectura.append(f'=IF(COUNT({dif})=0,"Cuadre con Cetelem: todavía no hay capturas de Cetelem en "&$B$2&".",'
                       f'IF({absd}=0,"Cuadre con Cetelem: en los "&COUNT({dif})&" meses con capturas, todos los '
                       f'rechazos con folio están en MG Colima Rechazos.","Cuadre con Cetelem: en los "&COUNT({dif})&'
                       f'" meses con capturas hay "&{absd}&" diferencia(s) con MG Colima Rechazos; revisa la tabla '
                       f'«Cuadre contra las capturas de Cetelem»."))')
    h.pon(8, 1, "Lectura del periodo", "seccion"); h.une(8, 1, 8, NCOLS)
    for i, formula in enumerate(lectura):
        h.pon(9 + i, 1, f'=IFERROR({formula[1:]},"")', "lectura"); h.une(9 + i, 1, 9 + i, NCOLS)

    # ------------------------------------------------------------------ gráficas
    f0 = fila_graficas
    h.pon(f0, 1, "Gráficas", "seccion"); h.une(f0, 1, f0, NCOLS)
    g = f0 + 1
    t1 = lambda col: (T1["enc"], T1["fin"], col)
    h.graficas += [
        dict(titulo="Rechazos por mes y financiera", tipo="columna", apilada=True, dominio=t1("A"),
             series=[(t1("C"), SERIE[0]), (t1("D"), SERIE[1]), (t1("E"), SERIE[2])], eje="Rechazos"),
        dict(titulo="Seguimiento de rechazos por mes", tipo="columna", dominio=t1("A"),
             series=[(t1("B"), SECUENCIA[0]), (t1("F"), SECUENCIA[1]), (t1("G"), SECUENCIA[2])], eje="Clientes"),
        dict(titulo="Rechazos por financiera (año)", tipo="barra", dominio=(T2["enc"], T2["fin"], "A"),
             series=[((T2["enc"], T2["fin"], "B"), SERIE[0])], eje="Rechazos"),
        dict(titulo="Rechazos por asesor (año)", tipo="barra", dominio=(T3["enc"], T3["fin"], "F"),  # nombre corto
             series=[((T3["enc"], T3["fin"], "K"), SERIE[0])], eje="Rechazos"),
        dict(titulo="Registro en Sale-U por mes", tipo="columna", apilada=True, dominio=t1("A"),
             series=[(t1("I"), ESTADO["bien"]), (t1("J"), ESTADO["critico"])], eje="Rechazos"),
        dict(titulo="Pendientes", tipo="barra", dominio=(T4["enc"], T4["fin"], "A"),
             series=[((T4["enc"], T4["fin"], "D"), ESTADO["serio"])], eje="Cantidad"),
    ]
    if T5:
        t5 = lambda col: (T5["enc"], T5["fin"], col)
        h.graficas += [
            dict(titulo="Bitácora: solicitudes por estatus y mes", tipo="columna", apilada=True, dominio=t5("A"),
                 series=[(t5("C"), ESTADO["critico"]), (t5("E"), ESTADO["aviso"]), (t5("D"), ESTADO["bien"]),
                         (t5("G"), SECUENCIA[2]), (t5("F"), ESTADO["neutro"])], eje="Solicitudes"),
            dict(titulo="Bitácora: % de solicitudes rechazadas por mes", tipo="linea", dominio=t5("A"),
                 series=[(t5("H"), SERIE[0])], eje="% rechazo", pct=True),
        ]
    for i, gr in enumerate(h.graficas):
        gr["ancla"] = (g + (i // 2) * ALTO_FILAS, 1 if i % 2 == 0 else 7)
    assert len(h.graficas) == n_graficas and len(lectura) == n_lectura
    return h


def _anio_actual(base):
    if TAB_RESUMEN in base.pestanas():
        try:
            v = base._valores(TAB_RESUMEN)
            a = v[1][1] if len(v) > 1 and len(v[1]) > 1 else None
            if isinstance(a, (int, float)) and 2000 <= a <= 2100:
                return int(a)
        except Exception:
            pass
    return dt.date.today().year


def _catalogo(base, columna):
    if "Catalogos" not in base.pestanas():
        return []
    cat = base.leer("Catalogos")
    col = next((c for c in cat.columns if norm(c) == norm(columna)), None)
    return [str(v).strip() for v in cat[col].tolist() if es_valor_catalogo(v)] if col else []


def nombre_corto(nombre):
    """'JESSICA ELIZABETH MACHUCA ROQUE' → 'JESSICA MACHUCA' (para que quepa en la gráfica)."""
    p = [x for x in str(nombre).split() if x.upper() not in ("DE", "DEL", "LA", "LAS", "LOS", "Y")]
    return " ".join([p[0], p[-2]]) if len(p) >= 4 else " ".join(p[:2])


# =============================================================================
# Escritura
# =============================================================================
ANCHOS = [150] + [92] * (NCOLS - 1)                     # px; gráficas: izquierda A:F, derecha G:L
ALTO_GRAFICA = 300


def construir(base):
    """Reescribe la pestaña Resumen con el dashboard. Devuelve cuántas gráficas dibujó."""
    h = contenido(base)
    max_f = max(f for f, _ in h.celdas)
    tabla = [["" for _ in range(NCOLS)] for _ in range(max_f)]
    for (f, c), (v, _, _) in h.celdas.items():
        tabla[f - 1][c - 1] = v
    base._reemplazar(TAB_RESUMEN, tabla, False)
    if isinstance(base, BaseGoogle):
        _google(base, h, max_f)
    else:
        _excel(base, h)
    return len(h.graficas)


def _estilo(rol):
    """(tamaño, negrita, color de letra, fondo, alineación horizontal, ajuste)"""
    if rol.startswith("leyenda:"):
        return 9, False, TINTA, rol.split(":", 1)[1], "LEFT", False
    return {
        "titulo": (16, True, AZUL_TITULO, None, "LEFT", False),
        "etiqueta": (10, True, TINTA, None, "RIGHT", False),
        "anio": (12, True, AZUL_TITULO, "#FFF2CC", "CENTER", False),
        "nota": (9, False, TINTA_3, None, "LEFT", False),
        "seccion": (12, True, AZUL_TITULO, None, "LEFT", False),
        "subseccion": (10, True, TINTA_2, None, "LEFT", False),
        "kpi_titulo": (9, True, TINTA_2, FONDO_KPI, "CENTER", False),
        "kpi_valor": (20, True, TINTA, FONDO_KPI, "CENTER", False),
        "kpi_sub": (9, False, TINTA_2, FONDO_KPI, "CENTER", True),
        "lectura": (10, False, TINTA, None, "LEFT", True),
        "th": (9, True, "#FFFFFF", FONDO_ENC, "CENTER", True),
        "td_nombre": (9, False, TINTA, None, "LEFT", False),
        "td": (9, False, TINTA, None, "RIGHT", False),
        "total": (9, True, TINTA, "#EFEFEF", "RIGHT", False),
    }[rol]


def _rgb(color):
    return {k: int(color.lstrip("#")[i:i + 2], 16) / 255 for k, i in (("red", 0), ("green", 2), ("blue", 4))}


def _google(base, h, max_f):
    ws = base._hoja(TAB_RESUMEN)
    sid = ws.id
    meta = base.libro.fetch_sheet_metadata(params={"fields": "sheets(properties(sheetId),charts(chartId),"
                                                             "conditionalFormats,merges)"})
    hoja = next((x for x in meta.get("sheets", []) if x["properties"]["sheetId"] == sid), {})
    rango = lambda f0, f1, c0, c1: {"sheetId": sid, "startRowIndex": f0, "endRowIndex": f1,
                                    "startColumnIndex": c0, "endColumnIndex": c1}
    todo = rango(0, max(ws.row_count, max_f), 0, max(ws.col_count, NCOLS))
    pedidos = [{"deleteEmbeddedObject": {"objectId": c["chartId"]}} for c in hoja.get("charts", [])]
    pedidos += [{"deleteConditionalFormatRule": {"sheetId": sid, "index": i}}
                for i in reversed(range(len(hoja.get("conditionalFormats", []))))]
    pedidos += [{"unmergeCells": {"range": todo}},
                {"repeatCell": {"range": todo, "fields": "userEnteredFormat", "cell": {"userEnteredFormat": {
                    "textFormat": {"fontFamily": FUENTE, "fontSize": 9, "foregroundColor": _rgb(TINTA)},
                    "verticalAlignment": "MIDDLE", "backgroundColor": _rgb("#FFFFFF")}}}},
                {"updateSheetProperties": {"properties": {"sheetId": sid, "gridProperties": {
                    "frozenRowCount": 0, "frozenColumnCount": 0, "hideGridlines": True}},
                    "fields": "gridProperties(frozenRowCount,frozenColumnCount,hideGridlines)"}}]
    for i, px in enumerate(ANCHOS):
        pedidos.append({"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "COLUMNS",
                        "startIndex": i, "endIndex": i + 1}, "properties": {"pixelSize": px}, "fields": "pixelSize"}})
    for fila, px in ((1, 34), (5, 40), (6, 30)):
        pedidos.append({"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "ROWS",
                        "startIndex": fila - 1, "endIndex": fila}, "properties": {"pixelSize": px},
                        "fields": "pixelSize"}})
    for (f, c), (_, rol, fmt) in h.celdas.items():
        tam, neg, color, fondo, alin, ajuste = _estilo(rol)
        formato = {"textFormat": {"fontFamily": FUENTE, "fontSize": tam, "bold": neg, "foregroundColor": _rgb(color)},
                   "horizontalAlignment": alin, "verticalAlignment": "MIDDLE",
                   "wrapStrategy": "WRAP" if ajuste else "OVERFLOW_CELL"}
        if fondo:
            formato["backgroundColor"] = _rgb(fondo)
        if rol in ("td", "td_nombre", "total", "th"):
            linea = {"style": "SOLID", "color": _rgb(LINEA)}
            formato["borders"] = {"top": linea, "bottom": linea, "left": linea, "right": linea}
        if fmt == "pct":
            formato["numberFormat"] = {"type": "PERCENT", "pattern": "0%"}
        pedidos.append({"repeatCell": {"range": rango(f - 1, f, c - 1, c), "cell": {"userEnteredFormat": formato},
                                       "fields": "userEnteredFormat"}})
    for f0, c0, f1, c1 in h.merges:
        if (f0, c0) != (f1, c1):
            pedidos.append({"mergeCells": {"range": rango(f0 - 1, f1, c0 - 1, c1), "mergeType": "MERGE_ALL"}})
    for f, c, cond in h.alertas:
        ref = f"{letra(c)}{f}"
        tipo, valor = (("NUMBER_GREATER", "0") if cond == ">0" else
                       ("CUSTOM_FORMULA", f"=AND(ISNUMBER({ref}),{ref}<>0)"))
        pedidos.append({"addConditionalFormatRule": {"index": 0, "rule": {
            "ranges": [rango(f - 1, f, c - 1, c + (1 if f == 5 else 0))],
            "booleanRule": {"condition": {"type": tipo, "values": [{"userEnteredValue": valor}]},
                            "format": {"textFormat": {"foregroundColor": _rgb(ESTADO["critico"]), "bold": True}}}}}})
    for gr in h.graficas:
        pedidos.append({"addChart": {"chart": _grafica_google(gr, sid)}})
    base.libro.batch_update({"requests": pedidos})
    base._limpiar_cache(TAB_RESUMEN)


def _fuente_google(sid, f0, f1, col):
    c = ord(col) - 65
    return {"sourceRange": {"sources": [{"sheetId": sid, "startRowIndex": f0 - 1, "endRowIndex": f1,
                                         "startColumnIndex": c, "endColumnIndex": c + 1}]}}


def _grafica_google(gr, sid):
    barra = gr["tipo"] == "barra"
    eje_valor = "BOTTOM_AXIS" if barra else "LEFT_AXIS"
    texto = lambda tam, color, neg=False: {"fontFamily": FUENTE, "fontSize": tam, "bold": neg,
                                          "foregroundColor": _rgb(color)}
    series = []
    for (f0, f1, col), color in gr["series"]:
        s = {"series": _fuente_google(sid, f0, f1, col), "targetAxis": eje_valor,
             "colorStyle": {"rgbColor": _rgb(color)}}
        if gr["tipo"] == "linea":
            s["lineStyle"] = {"width": 2}
            s["pointStyle"] = {"shape": "CIRCLE", "size": 6}
        if barra:
            s["dataLabel"] = {"type": "DATA", "placement": "OUTSIDE_END", "textFormat": texto(9, TINTA_2)}
        series.append(s)
    basico = {
        "chartType": {"columna": "COLUMN", "barra": "BAR", "linea": "LINE"}[gr["tipo"]],
        "legendPosition": "BOTTOM_LEGEND" if len(series) > 1 else "NO_LEGEND",
        "headerCount": 1,
        "axis": [{"position": "BOTTOM_AXIS", "title": gr["eje"] if barra else "",
                  "format": texto(9, TINTA_2)},
                 {"position": "LEFT_AXIS", "title": "" if barra else gr["eje"], "format": texto(9, TINTA_2)}],
        "domains": [{"domain": _fuente_google(sid, *gr["dominio"])}],
        "series": series,
    }
    if gr.get("apilada"):
        basico["stackedType"] = "STACKED"
    alto = ALTO_GRAFICA
    ancho = sum(ANCHOS[:6]) - 12 if gr["ancla"][1] == 1 else sum(ANCHOS[6:]) - 12
    return {"spec": {"title": gr["titulo"], "titleTextFormat": texto(11, TINTA, True), "fontName": FUENTE,
                     "backgroundColor": _rgb("#FFFFFF"), "basicChart": basico},
            "position": {"overlayPosition": {"anchorCell": {"sheetId": sid, "rowIndex": gr["ancla"][0] - 1,
                                                            "columnIndex": gr["ancla"][1] - 1},
                                             "offsetXPixels": 4, "offsetYPixels": 4,
                                             "widthPixels": ancho, "heightPixels": alto}}}


def _excel(base, h):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.formatting import ConditionalFormattingList
    ws = base.libro[TAB_RESUMEN]
    ws.conditional_formatting = ConditionalFormattingList()      # las reglas del Resumen anterior ya no aplican
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    for i, px in enumerate(ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = px / 7
    for fila, px in ((1, 34), (5, 40), (6, 30)):
        ws.row_dimensions[fila].height = px * 0.75
    hex_ = lambda c: "FF" + c.lstrip("#").upper()
    linea = Side(style="thin", color=hex_(LINEA))
    for (f, c), (_, rol, fmt) in h.celdas.items():
        tam, neg, color, fondo, alin, ajuste = _estilo(rol)
        celda = ws.cell(f, c)
        celda.font = Font(name=FUENTE, size=tam, bold=neg, color=hex_(color))
        celda.alignment = Alignment(horizontal=alin.lower(), vertical="center", wrap_text=ajuste)
        if fondo:
            celda.fill = PatternFill("solid", fgColor=hex_(fondo))
        if rol in ("td", "td_nombre", "total", "th"):
            celda.border = Border(top=linea, bottom=linea, left=linea, right=linea)
        if fmt == "pct":
            celda.number_format = "0%"
    for f0, c0, f1, c1 in h.merges:
        if (f0, c0) != (f1, c1):
            ws.merge_cells(start_row=f0, start_column=c0, end_row=f1, end_column=c1)
    from openpyxl.formatting.rule import FormulaRule
    for f, c, cond in h.alertas:
        ref = f"{get_column_letter(c)}{f}"
        regla = (CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=hex_(ESTADO["critico"]), bold=True))
                 if cond == ">0" else
                 FormulaRule(formula=[f"AND(ISNUMBER({ref}),{ref}<>0)"],
                             font=Font(color=hex_(ESTADO["critico"]), bold=True)))
        ws.conditional_formatting.add(ref, regla)
    for gr in h.graficas:
        if gr["tipo"] == "linea":
            ch = LineChart()
        else:
            ch = BarChart()
            ch.type = "bar" if gr["tipo"] == "barra" else "col"
            if gr.get("apilada"):
                ch.grouping, ch.overlap = "stacked", 100
            ch.gapWidth = 60
        ch.title = gr["titulo"]
        ch.style = 2
        ch.height = ALTO_GRAFICA / 37.8
        ch.width = (sum(ANCHOS[:6]) if gr["ancla"][1] == 1 else sum(ANCHOS[6:])) / 37.8 - 0.3
        ch.legend = ch.legend if len(gr["series"]) > 1 else None
        if ch.legend is not None:
            ch.legend.position = "b"
        col = lambda letra_: ord(letra_) - 64
        f0, f1, dc = gr["dominio"]
        for (s0, s1, sc), color in gr["series"]:
            ch.add_data(Reference(ws, min_col=col(sc), min_row=s0, max_row=s1), titles_from_data=True)
            serie = ch.series[-1]
            if gr["tipo"] == "linea":
                serie.graphicalProperties.line.solidFill = color.lstrip("#")
                serie.graphicalProperties.line.width = 25400
                serie.smooth = False
            else:
                serie.graphicalProperties.solidFill = color.lstrip("#")
                serie.graphicalProperties.line.noFill = True
        ch.set_categories(Reference(ws, min_col=col(dc), min_row=f0 + 1, max_row=f1))
        ch.y_axis.title = gr["eje"] if gr["tipo"] != "barra" else None
        ch.y_axis.majorGridlines.spPr = None
        if gr["tipo"] == "barra":
            ch.x_axis.scaling.orientation = "maxMin"           # primero el de más rechazos
        if gr.get("pct"):
            ch.y_axis.number_format = "0%"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ws.add_chart(ch, f"{get_column_letter(gr['ancla'][1])}{gr['ancla'][0]}")
