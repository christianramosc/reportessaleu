# -*- coding: utf-8 -*-
"""
Reconoce qué archivo se subió, sin importar en qué campo se ponga ni cómo se llame:

  credito  → reporteLeadsconCredito de Sale-U (títulos «Nombre del Cliente», «Estatus del crédito»)
  fuente   → ReporteFuentes de Sale-U (todos los leads: «Fecha de Alta», «Nombre», «Telefono»)
  cetelem  → captura mensual de Cetelem (alguna hoja con la columna «Folio CCK»)
  fi       → Bitácora de F&I (hojas DETALLE / CIERRE con «INGRESOS · VENDEDOR · … · CLIENTE»)
  base     → el propio Excel de seguimiento (pestaña «MG Colima Rechazos»)
"""
import csv
import io

from cuadres.base import TAB_RECHAZOS, norm

NOMBRES = {"credito": "Sale-U · Crédito", "fuente": "Sale-U · Fuente", "cetelem": "Captura de Cetelem",
           "fi": "Bitácora de F&I", "base": "Excel de seguimiento", None: "No reconocido"}


def _texto(datos):
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return datos.decode(enc)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", "ignore")


def _filas_csv(datos, n=15):
    return [f for f, _ in zip(csv.reader(io.StringIO(_texto(datos))), range(n))]


def _tipo_filas(filas):
    for f in filas:
        t = {norm(x) for x in f if x is not None}
        if "NOMBRE DEL CLIENTE" in t and any("ESTATUS DEL CREDITO" in x for x in t):
            return "credito"
        if "FECHA DE ALTA" in t and "NOMBRE" in t and "TELEFONO" in t:
            return "fuente"
        if "FOLIO CCK" in t:
            return "cetelem"
    return None


def tipo(nombre, datos):
    """Tipo del archivo (ver arriba) o None."""
    nombre = (nombre or "").lower()
    if nombre.endswith(".csv"):
        return _tipo_filas(_filas_csv(datos))
    if not nombre.endswith((".xlsx", ".xlsm", ".xls")):
        return None
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if TAB_RECHAZOS in wb.sheetnames:
            return "base"
        from cuadres.fuentes import _formato
        encontrado = None
        for ws in wb.worksheets:
            filas = [list(f) for f, _ in zip(ws.iter_rows(values_only=True), range(15))]
            if filas and _formato(filas[0]):
                return "fi"
            t = _tipo_filas(filas)
            if t and not encontrado:
                encontrado = t
        return encontrado
    finally:
        wb.close()


def clasificar(archivos):
    """archivos: [(nombre, bytes)] → {tipo: [(nombre, bytes)]} (los no reconocidos, en la llave None)."""
    out = {}
    for nombre, datos in archivos:
        out.setdefault(tipo(nombre, datos), []).append((nombre, datos))
    return out
