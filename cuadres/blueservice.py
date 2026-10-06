# -*- coding: utf-8 -*-
"""
blueservice.py
--------------
Cuadre por VIN: reporte de PDIs de BlueService (PDIs_MG_Colima.xlsx)
vs. Excel de inventario de MG Colima (hoja de inventario físico, p. ej. ALMACEN).

Para cada VIN dice si cuadra, si falta en BlueService, si BlueService lo tiene y
el inventario no (y en qué otra hoja del inventario aparece: VENDIDOS, TRASLADOS…),
detecta VIN mal capturados con el dígito verificador y compara el avance de PDIs.

Uso en Colab: python blueservice.py  (pide subir los dos archivos)
"""

import os, re, glob, unicodedata, subprocess, sys, warnings
from datetime import datetime
import pandas as pd

warnings.filterwarnings("ignore")
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

PDV = "MG COLIMA"
ETAPAS_PDI = ["PDI ARRIBO", "PDI 1", "PDI 2", "PDI 3"]
ETIQUETA = {"PDI ARRIBO": "Arribo", "PDI 1": "PDI 1", "PDI 2": "PDI 2", "PDI 3": "PDI 3"}
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

R_CUADRA = "Cuadra"
R_CUADRA_ERR = "Cuadra – VIN mal capturado en inventario"
R_FALTA_BS = "Falta en BlueService"
R_FALTA_INV = "En BlueService, no está en inventario"
R_VENDIDA = "Vendida en BlueService"


# =============================================================================
# Utilidades
# =============================================================================
def quitar_acentos(t):
    t = unicodedata.normalize("NFKD", str(t))
    return "".join(ch for ch in t if not unicodedata.combining(ch))


def norm_txt(t):
    return re.sub(r"\s+", " ", quitar_acentos(t).upper()).strip()


def limpiar_vin(v):
    v = re.sub(r"[^A-Z0-9]", "", str(v).upper())
    return v if 15 <= len(v) <= 18 and re.search(r"\d", v) and re.search(r"[A-Z]", v) else ""


_T = {**{str(i): i for i in range(10)}, **dict(zip("ABCDEFGH", range(1, 9))),
      **dict(zip("JKLMN", range(1, 6))), "P": 7, "R": 9, **dict(zip("STUVWXYZ", range(2, 10)))}
_W = [8, 7, 6, 5, 4, 3, 2, 10, 0, 9, 8, 7, 6, 5, 4, 3, 2]


def vin_valido(v):
    """Valida longitud y dígito verificador (posición 9)."""
    if len(v) != 17 or any(c not in _T for c in v):
        return False
    s = sum(_T[c] * w for c, w in zip(v, _W)) % 11
    return v[8] == ("X" if s == 10 else str(s))


def diferencias(a, b):
    return sum(x != y for x, y in zip(a, b)) + abs(len(a) - len(b))


def a_fecha(x):
    x = str(x).strip()
    if x in ("", "nan", "NaT", "None"):
        return pd.NaT
    if re.match(r"\d{4}-\d{2}-\d{2}", x):
        return pd.to_datetime(x[:10], errors="coerce")
    return pd.to_datetime(x, dayfirst=True, errors="coerce")


def fmt(f):
    return f.strftime("%d/%m/%Y") if pd.notna(f) else ""


# =============================================================================
# Lectura
# =============================================================================
def leer_blueservice(ruta):
    crudo = pd.read_excel(ruta, header=None, dtype=str).fillna("")
    h = next(i for i, r in crudo.iterrows() if any(norm_txt(c) == "VIN" for c in r))
    df = crudo.iloc[h + 1:].copy()
    df.columns = [norm_txt(c).rstrip(".") for c in crudo.iloc[h]]
    col = lambda *ops: next((c for c in df.columns for o in ops if c.startswith(o)), None)
    bs = pd.DataFrame({
        "VIN": df[col("VIN")].map(limpiar_vin),
        "VIN_original": df[col("VIN")].str.strip(),
        "Modelo_BS": df[col("MODELO")] if col("MODELO") else "",
        "Dealer_origen": df[col("DEALER ORIGEN")] if col("DEALER ORIGEN") else "",
        "Dealer_actual": df[col("DEALER ACTUAL")] if col("DEALER ACTUAL") else "",
        "Encargado": df[col("NOMBRE ENCARGADO", "ENCARGADO")] if col("NOMBRE ENCARGADO", "ENCARGADO") else "",
        "Fecha_inventario_BS": df[col("FECHA INVENTARIO")].map(a_fecha) if col("FECHA INVENTARIO") else pd.NaT,
        "Fecha_venta_BS": df[col("FECHA VENTA")].map(a_fecha) if col("FECHA VENTA") else pd.NaT,
        "Etapa_BS": df[col("PDI")].str.strip() if col("PDI") else "",
        "Status_BS": df[col("STATUS", "ESTATUS")].str.strip() if col("STATUS", "ESTATUS") else "",
    })
    bs = bs[bs["VIN"] != ""]
    bs["Veces_en_BS"] = bs.groupby("VIN")["VIN"].transform("size")
    return bs.drop_duplicates("VIN").reset_index(drop=True)


def hojas_inventario(ruta):
    return pd.ExcelFile(ruta).sheet_names


def hoja_por_defecto(hojas):
    for clave in ("INVENTARIO", "ALMACEN"):
        for h in hojas:
            if clave in norm_txt(h):
                return h
    return hojas[0]


def leer_hoja(ruta, hoja):
    """Lee una hoja con uno o varios bloques (encabezado con 'VIN' + filas).
    Regresa los registros con VIN, la fecha de corte y el bloque (mes) de cada fila."""
    crudo = pd.read_excel(ruta, sheet_name=hoja, header=None, dtype=str).fillna("")
    registros, encabezado, col_vin, bloque, corte = [], None, None, "", pd.NaT
    for i, fila in crudo.iterrows():
        celdas = [str(c).strip() for c in fila]
        textos = [norm_txt(c) for c in celdas]
        idx_vin = next((j for j, t in enumerate(textos) if t == "VIN" or t.endswith("/ VIN")
                        or t == "SERIE"), None)
        if idx_vin is not None:
            encabezado, col_vin = textos, idx_vin
            continue
        if encabezado is None:
            if pd.isna(corte):
                for c in celdas:
                    if re.match(r"\d{4}-\d{2}-\d{2}", c):
                        corte = a_fecha(c)
                        break
            continue
        if not any(celdas):
            continue
        mes = next((m for m in MESES if m in re.sub(r"[^A-Z]", "", " ".join(textos))), None)
        vin = limpiar_vin(celdas[col_vin]) if col_vin < len(celdas) else ""
        if not vin:   # el VIN pudo quedar en otra columna (filas desfasadas)
            vin = next((limpiar_vin(c) for c in celdas if limpiar_vin(c) and len(limpiar_vin(c)) >= 16), "")
        if not vin:
            if mes:
                bloque = mes.capitalize()
            continue
        reg = {"VIN": vin, "VIN_original": celdas[col_vin] if col_vin < len(celdas) else vin,
               "Hoja": hoja, "Bloque": bloque, "Fila_excel": i + 1}
        for j, t in enumerate(encabezado):
            if t and j < len(celdas) and t not in reg:
                reg[t] = celdas[j]
        registros.append(reg)
    return pd.DataFrame(registros), corte


def leer_inventario(ruta, hoja_principal):
    hojas = hojas_inventario(ruta)
    principal, corte = leer_hoja(ruta, hoja_principal)
    otras = []
    for h in hojas:
        if h != hoja_principal:
            d, _ = leer_hoja(ruta, h)
            if not d.empty:
                otras.append(d)
    otras = pd.concat(otras, ignore_index=True) if otras else pd.DataFrame(columns=["VIN", "Hoja"])
    if pd.isna(corte):
        corte = pd.Timestamp(datetime.now().date())
    return principal, otras, corte


# =============================================================================
# Cruce
# =============================================================================
def resumen_pdis(fila, corte):
    """Fechas de PDI en el inventario: realizados (<= corte), programados y fechas inválidas."""
    hechos, programados, malas = [], [], []
    for e in ETAPAS_PDI:
        valor = str(fila.get(e, "")).strip()
        if not valor or valor == "nan":
            continue
        f = a_fecha(valor)
        if pd.isna(f) or f.year < corte.year - 2 or (f - corte).days > 400:
            malas.append(f"{ETIQUETA[e]}: {valor[:10]}")
        elif f <= corte:
            hechos.append((e, f))
        else:
            programados.append((e, f))
    return hechos, programados, malas


def ubicaciones(vin, otras):
    sub = otras[otras["VIN"] == vin]
    partes = []
    for _, r in sub.iterrows():
        detalle = r["Hoja"].title()
        if r.get("Bloque"):
            detalle += f" ({r['Bloque']})"
        ubic = r.get("UBICACION", "")
        ubic = "" if pd.isna(ubic) else str(ubic).strip()
        if ubic and ubic.lower() != "nan" and not re.match(r"\d{4}-\d{2}-\d{2}", ubic):
            detalle += f" – {ubic}"
        partes.append(detalle)
    return "; ".join(dict.fromkeys(partes))


def cruzar(bs, inv, otras, corte, hoja):
    vins_bs = set(bs["VIN"])
    B = bs.set_index("VIN")
    filas, usados_bs = [], set()

    # --- 1. Cada unidad del inventario físico
    for _, r in inv.iterrows():
        vin, obs = r["VIN"], []
        if not vin_valido(vin):
            obs.append("VIN inválido (longitud o dígito verificador)")
        vin_bs, resultado = None, R_FALTA_BS
        if vin in vins_bs:
            vin_bs, resultado = vin, R_CUADRA
        else:
            parecidos = [v for v in vins_bs if diferencias(vin, v) <= 2 and v not in inv["VIN"].values]
            parecidos.sort(key=lambda v: (not vin_valido(v), diferencias(vin, v)))
            if parecidos and (not vin_valido(vin) or diferencias(vin, parecidos[0]) == 1):
                vin_bs, resultado = parecidos[0], R_CUADRA_ERR
                obs.append(f"En BlueService es {vin_bs}: corregir VIN en inventario")
        hechos, prog, malas = resumen_pdis(r, corte)
        if malas:
            obs.append("Fecha de PDI inválida: " + ", ".join(malas))
        nota_bp = norm_txt(r.get("OBCERVACIONES DE BEEFEATER", r.get("OBSERVACIONES DE BEEFEATER", "")))
        fila = {
            "Resultado": resultado, "VIN": vin, "VIN_BlueService": vin_bs or "",
            "Modelo_inventario": r.get("MODELO", ""), "Version": r.get("VERSION", ""),
            "Color": r.get("COLOR", ""), "Ubicacion": r.get("UBICACION", ""),
            "Entrada_inventario": fmt(a_fecha(r.get("ENTRADA", ""))),
            "Ultimo_PDI_inventario": f"{ETIQUETA[hechos[-1][0]]} ({fmt(hechos[-1][1])})" if hechos else "Ninguno",
            "Proximo_PDI_inventario": f"{ETIQUETA[prog[0][0]]} ({fmt(prog[0][1])})" if prog else "",
            "Fila_excel": r["Fila_excel"],
        }
        if vin_bs:
            usados_bs.add(vin_bs)
            b = B.loc[vin_bs]
            fila.update({"Modelo_BS": b["Modelo_BS"], "Etapa_BS": b["Etapa_BS"], "Status_BS": b["Status_BS"],
                         "Dealer_origen": b["Dealer_origen"], "Fecha_venta_BS": fmt(b["Fecha_venta_BS"])})
            if pd.notna(b["Fecha_venta_BS"]):
                obs.append(f"BlueService la tiene vendida ({fmt(b['Fecha_venta_BS'])}) pero sigue en {hoja}")
            etapa = norm_txt(b["Etapa_BS"])
            if norm_txt(b["Status_BS"]) == "PENDIENTE" and etapa in ("PDI 1", "PDI 2", "PDI 3"):
                hecho = next((f for e, f in hechos if e == etapa), None)
                if hecho is not None:
                    obs.append(f"Inventario registra {etapa} el {fmt(hecho)}; en BlueService está pendiente")
            if "NO ESTA EN BP" in nota_bp:
                obs.append("Inventario dice 'NO ESTÁ EN BP' pero sí aparece en BlueService")
        else:
            fila.update({"Modelo_BS": "", "Etapa_BS": "", "Status_BS": "", "Dealer_origen": "",
                         "Fecha_venta_BS": ""})
            if "NO ESTA EN BP" in nota_bp:
                obs.append("Ya marcado en inventario como 'NO ESTÁ EN BP'")
            otra = ubicaciones(vin, otras)
            if otra:
                obs.append(f"También aparece en: {otra}")
        if (inv["VIN"] == vin).sum() > 1:
            obs.append("VIN repetido en la hoja de inventario")
        fila["Observaciones"] = "; ".join(obs)
        filas.append(fila)
    cuadre_inv = pd.DataFrame(filas).drop_duplicates("VIN")

    # --- 2. Unidades de BlueService que no están en el inventario físico
    filas = []
    for _, b in bs[~bs["VIN"].isin(usados_bs)].iterrows():
        donde = ubicaciones(b["VIN"], otras)
        vendida = pd.notna(b["Fecha_venta_BS"])
        if vendida:
            resultado = R_VENDIDA
            situacion = ("Vendida y registrada en inventario" if "Vendidos" in donde
                         else "Vendida en BlueService; no aparece en VENDIDOS del inventario")
        else:
            resultado = R_FALTA_INV
            if "Vendidos" in donde:
                situacion = "Inventario la tiene vendida; BlueService sin fecha de venta"
            elif "Traslados" in donde:
                situacion = "Inventario la tiene en traslados; en BlueService sigue en Colima"
            elif donde:
                situacion = f"Solo aparece en {donde}"
            else:
                situacion = "No aparece en ninguna hoja del inventario"
        obs = []
        if b["Veces_en_BS"] > 1:
            obs.append(f"VIN repetido {b['Veces_en_BS']} veces en BlueService")
        filas.append({
            "Resultado": resultado, "Situacion": situacion, "VIN": b["VIN"], "Modelo_BS": b["Modelo_BS"],
            "Dealer_origen": b["Dealer_origen"], "Dealer_actual": b["Dealer_actual"],
            "Encargado": b["Encargado"], "Fecha_inventario_BS": fmt(b["Fecha_inventario_BS"]),
            "Fecha_venta_BS": fmt(b["Fecha_venta_BS"]), "Etapa_BS": b["Etapa_BS"], "Status_BS": b["Status_BS"],
            "Aparece_en_inventario": donde, "Observaciones": "; ".join(obs)})
    solo_bs = pd.DataFrame(filas)
    if solo_bs.empty:
        solo_bs = pd.DataFrame(columns=["Resultado", "Situacion", "VIN"])
    return cuadre_inv, solo_bs


# =============================================================================
# Salidas
# =============================================================================
COLORES = {R_CUADRA: "C6EFCE", R_CUADRA_ERR: "FFEB9C", R_FALTA_BS: "FFC7CE",
           R_FALTA_INV: "FFC7CE", R_VENDIDA: "DDEBF7"}
COLS_INV = ["Resultado", "VIN", "VIN_BlueService", "Modelo_inventario", "Version", "Color", "Ubicacion",
            "Entrada_inventario", "Ultimo_PDI_inventario", "Proximo_PDI_inventario", "Etapa_BS",
            "Status_BS", "Modelo_BS", "Dealer_origen", "Fecha_venta_BS", "Fila_excel", "Observaciones"]
COLS_BS = ["Resultado", "Situacion", "VIN", "Modelo_BS", "Dealer_origen", "Dealer_actual", "Encargado",
           "Fecha_inventario_BS", "Fecha_venta_BS", "Etapa_BS", "Status_BS", "Aparece_en_inventario",
           "Observaciones"]


def exportar_excel(ruta, resumen, cuadre_inv, solo_bs, revisar):
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        resumen.to_excel(w, sheet_name="Resumen", index=False)
        cuadre_inv[COLS_INV].to_excel(w, sheet_name="Inventario_vs_BS", index=False)
        cuadre_inv[cuadre_inv["Resultado"] == R_FALTA_BS][COLS_INV].to_excel(
            w, sheet_name="Falta_en_BlueService", index=False)
        solo_bs[solo_bs["Resultado"] == R_FALTA_INV][COLS_BS].to_excel(
            w, sheet_name="En_BS_no_en_inventario", index=False)
        revisar[COLS_INV].to_excel(w, sheet_name="Revisar", index=False)
        solo_bs[solo_bs["Resultado"] == R_VENDIDA][COLS_BS].to_excel(
            w, sheet_name="Vendidas_BS", index=False)
    wb = load_workbook(ruta)
    b = Side(style="thin", color="BFBFBF")
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.font = Font(name="Arial", size=10)
                c.border = Border(top=b, bottom=b, left=b, right=b)
                if isinstance(c.value, str) and c.value in COLORES:
                    c.fill = PatternFill("solid", fgColor=COLORES[c.value])
        for i, col in enumerate(ws.columns, 1):
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 10), 60)
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)


def exportar_pdf(ruta, resumen, falta_bs, falta_inv, revisar, corte, hoja):
    st = getSampleStyleSheet()
    chico = ParagraphStyle("c", parent=st["Normal"], fontSize=7.5, leading=9)
    blanco = ParagraphStyle("b", parent=chico, textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=st["Title"], fontSize=18, textColor=colors.HexColor("#1F3864"))
    h2 = ParagraphStyle("h2", parent=st["Heading2"], textColor=colors.HexColor("#1F3864"))
    txt = ParagraphStyle("t", parent=st["Normal"], fontSize=9.5, leading=13)

    def tabla(df, cols, anchos, color="#1F3864"):
        if df.empty:
            return Paragraph("<i>Sin registros.</i>", txt)
        datos = [[Paragraph(f"<b>{c.replace('_', ' ')}</b>", blanco) for c in cols]]
        for _, r in df.iterrows():
            datos.append([Paragraph("" if pd.isna(r[c]) else str(r[c]), chico) for c in cols])
        t = Table(datos, colWidths=[a * cm for a in anchos], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(color)),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BFBFBF")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    doc = SimpleDocTemplate(ruta, pagesize=landscape(letter), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm)
    el = [Paragraph(f"Cuadre BlueService vs Inventario – {PDV}", h1),
          Paragraph(f"Hoja de inventario: {hoja} · Fecha de corte: {fmt(corte)} · "
                    f"Generado: {datetime.now():%d/%m/%Y %H:%M}", txt), Spacer(1, 10),
          tabla(resumen.assign(Valor=resumen["Valor"].astype(str)), ["Concepto", "Valor"], [13, 4]),
          Spacer(1, 10),
          Paragraph("El cruce es por VIN. Si un VIN del inventario no existe en BlueService, se busca uno "
                    "que difiera en uno o dos caracteres y se usa el dígito verificador (posición 9) para "
                    "saber cuál está bien capturado.", txt)]
    el += [PageBreak(), Paragraph(f"1. Unidades en {hoja.lower()} que NO están en BlueService ({len(falta_bs)})", h2),
           tabla(falta_bs, ["VIN", "Modelo_inventario", "Version", "Color", "Ubicacion", "Entrada_inventario",
                            "Observaciones"], [3.6, 2.2, 4.2, 1.8, 2.4, 2.2, 9], "#C00000")]
    el += [Spacer(1, 14), Paragraph(f"2. Datos a revisar ({len(revisar)})", h2),
           Paragraph("VIN mal capturados, fechas de PDI inválidas y PDIs que no coinciden entre archivos.", txt),
           Spacer(1, 6),
           tabla(revisar, ["Resultado", "VIN", "Modelo_inventario", "Ultimo_PDI_inventario", "Etapa_BS",
                           "Status_BS", "Observaciones"], [3.4, 3.4, 2, 3.4, 1.6, 1.8, 9.8], "#BF8F00")]
    el += [PageBreak(), Paragraph(f"3. En BlueService sin venta, pero no en {hoja.lower()} ({len(falta_inv)})", h2),
           Paragraph("BlueService las considera en el inventario de la agencia. La columna 'Situación' "
                     "indica dónde aparecen en el Excel de inventario.", txt), Spacer(1, 6)]
    if not falta_inv.empty:
        conteo = falta_inv.groupby("Situacion").size().sort_values(ascending=False)
        el.append(Paragraph(" · ".join(f"<b>{k}</b>: {v}" for k, v in conteo.items()), txt))
        el.append(Spacer(1, 6))
    el.append(tabla(falta_inv.sort_values(["Situacion", "Fecha_inventario_BS"]),
                    ["VIN", "Modelo_BS", "Dealer_origen", "Fecha_inventario_BS", "Etapa_BS", "Situacion",
                     "Aparece_en_inventario"], [3.5, 2.6, 2.8, 2.2, 1.8, 6.2, 6.3], "#7F7F7F"))
    doc.build(el)


# =============================================================================
# Proceso
# =============================================================================
def procesar(ruta_bs, ruta_inv, carpeta=".", hoja=None):
    hoja = hoja or hoja_por_defecto(hojas_inventario(ruta_inv))
    bs = leer_blueservice(ruta_bs)
    inv, otras, corte = leer_inventario(ruta_inv, hoja)
    if inv.empty:
        raise ValueError(f"No encontré VINs en la hoja '{hoja}'.")
    cuadre_inv, solo_bs = cruzar(bs, inv, otras, corte, hoja)

    falta_bs = cuadre_inv[cuadre_inv["Resultado"] == R_FALTA_BS]
    falta_inv = solo_bs[solo_bs["Resultado"] == R_FALTA_INV]
    vendidas = solo_bs[solo_bs["Resultado"] == R_VENDIDA]
    revisar = cuadre_inv[(cuadre_inv["Resultado"] != R_FALTA_BS) & (cuadre_inv["Observaciones"] != "")]
    en_bs = cuadre_inv["Resultado"].isin([R_CUADRA, R_CUADRA_ERR])

    resumen = pd.DataFrame([
        ("Fecha de corte del inventario", fmt(corte)),
        (f"Unidades en la hoja {hoja}", len(cuadre_inv)),
        ("VINs en BlueService", len(bs)),
        (f"Unidades de {hoja} que cuadran con BlueService", int(en_bs.sum())),
        ("   · de ellas con VIN mal capturado", int((cuadre_inv["Resultado"] == R_CUADRA_ERR).sum())),
        (f"Unidades de {hoja} que faltan en BlueService", len(falta_bs)),
        ("Porcentaje de cuadre del inventario", f"{en_bs.mean():.1%}" if len(cuadre_inv) else "—"),
        (f"En BlueService sin venta y fuera de {hoja}", len(falta_inv)),
        ("   · el inventario las tiene vendidas", int(falta_inv["Situacion"].str.startswith("Inventario la tiene vendida").sum())),
        ("   · el inventario las tiene en traslados", int(falta_inv["Situacion"].str.startswith("Inventario la tiene en traslados").sum())),
        ("   · no aparecen en el inventario", int((falta_inv["Situacion"] == "No aparece en ninguna hoja del inventario").sum())),
        ("Vendidas en BlueService (informativo)", len(vendidas)),
        ("Unidades con datos a revisar", len(revisar)),
    ], columns=["Concepto", "Valor"])

    sello = datetime.now().strftime("%Y%m%d_%H%M")
    rx = os.path.join(carpeta, f"Cuadre_BlueService_{sello}.xlsx")
    rp = os.path.join(carpeta, f"Cuadre_BlueService_{sello}.pdf")
    exportar_excel(rx, resumen, cuadre_inv, solo_bs, revisar)
    exportar_pdf(rp, resumen, falta_bs, falta_inv, revisar, corte, hoja)
    return dict(resumen=resumen, cuadre_inv=cuadre_inv, solo_bs=solo_bs, falta_bs=falta_bs,
                falta_inv=falta_inv, vendidas=vendidas, revisar=revisar, corte=corte, hoja=hoja,
                xlsx=rx, pdf=rp)


def main():
    try:
        from google.colab import files
        print("Sube el reporte de PDIs de BlueService y el Excel de inventario:")
        rutas = list(files.upload().keys())
    except ImportError:
        rutas = glob.glob("*.xlsx")
    rutas = [r for r in rutas if not r.startswith("Cuadre_")]
    ruta_bs = next(r for r in rutas if "PDI" in r.upper())
    ruta_inv = next(r for r in rutas if "INVENTARIO" in r.upper())
    r = procesar(ruta_bs, ruta_inv)
    print(r["resumen"].to_string(index=False))
    print(f"\n✔ Archivos generados: {r['xlsx']} y {r['pdf']}")
    try:
        from google.colab import files
        files.download(r["xlsx"]); files.download(r["pdf"])
    except ImportError:
        pass


if __name__ == "__main__":
    main()
