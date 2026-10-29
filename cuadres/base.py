# -*- coding: utf-8 -*-
"""
Base de seguimiento: lectura y escritura del Google Sheet "Seguimiento Rechazos MG Colima".

Dos implementaciones con la misma interfaz:
  · BaseGoogle  → el Google Sheet real (cuenta de servicio, gspread).
  · BaseExcel   → un .xlsx en memoria (modo sin conexión y pruebas). Al final se descarga.

Reglas de escritura:
  · Solo se escribe en las celdas que se indican (nunca se reescribe una fila completa), así
    las columnas de seguimiento que captura el equipo no se tocan.
  · Antes de aplicar cambios se guarda una copia oculta de la pestaña (respaldo).
  · Cada cambio queda en la pestaña "Historial".
"""
import io
import re
import unicodedata
from datetime import datetime, date, timedelta

import pandas as pd

TAB_RECHAZOS = "MG Colima Rechazos"
TAB_FINALIZADOS = "MG Colima Créditos Finalizados"
TAB_BITACORA = "Bitacora"
TAB_CAPTURAS = "Capturas Cetelem"
TAB_FI = "Bitacora F&I"          # copia de los registros de la Bitácora de F&I (para no perderlos entre cargas)
TAB_HISTORIAL = "Historial"
HISTORIAL_COLS = ["Fecha", "Origen", "Hoja", "Celda", "Cliente", "Campo", "Antes", "Después", "Motivo"]
RESPALDOS_A_GUARDAR = 1          # un respaldo por pestaña (el último); los anteriores se borran

_EPOCA = datetime(1899, 12, 30)        # día 0 de Google Sheets / Excel

# Estilo uniforme de las pestañas que escribe la app (el mismo que ya usaba la Bitacora)
ESTILO = dict(fuente="Calibri", tamano=9, texto="#000000", fondo="#FFFFFF", enc_texto="#FFFFFF", enc_fondo="#666666")
FORMATO_NUMERO = {"fecha": ("DATE", "dd/mm/yyyy", "DD/MM/YYYY"), "entero": ("NUMBER", "0", "0"),
                  "moneda": ("CURRENCY", "$#,##0.00", '"$"#,##0.00'),
                  "fechahora": ("DATE_TIME", "dd/mm/yyyy hh:mm", "DD/MM/YYYY HH:MM")}
# (tipo Google, patrón Google, patrón Excel)

# Encabezados que cambiaron de nombre: la app los lee con el nombre nuevo aunque el Sheet todavía tenga el viejo.
ALIAS_ENCABEZADO = {"CLIENTE CON REGISTRO EN SALESU": "Registro en Sale-U",
                    "CLIENTE CON REGISTRO EN SALE-U": "REGISTRO EN SALE-U"}


def _rgb(color):
    return {k: int(color.lstrip("#")[i:i + 2], 16) / 255 for k, i in (("red", 0), ("green", 2), ("blue", 4))}


def norm(t):
    """Texto comparable: sin acentos, mayúsculas, espacios simples."""
    t = unicodedata.normalize("NFKD", str(t if t is not None else ""))
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", t).strip().upper()


def a_serial(v):
    """Fecha → número de serie de Sheets (así la columna conserva su formato dd/mm/aaaa).
    También deja los valores listos para la API (sin tipos de numpy ni NaN)."""
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        v = v.item()                                   # numpy.int64 / float64 → int / float
    if isinstance(v, float) and v != v:
        return ""
    if isinstance(v, datetime):
        return (v - _EPOCA).total_seconds() / 86400
    if isinstance(v, date):
        return (datetime(v.year, v.month, v.day) - _EPOCA).days
    return v


def de_serial(v):
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return _EPOCA + timedelta(days=float(v))
    return v


def canonico(encabezado):
    """Nombre actual de un encabezado (los renombrados se traducen con ALIAS_ENCABEZADO)."""
    return ALIAS_ENCABEZADO.get(norm(encabezado), encabezado)


def letra(col):
    s = ""
    while col:
        col, r = divmod(col - 1, 26)
        s = chr(65 + r) + s
    return s


def tel_texto(v):
    """Teléfono como texto de 10 dígitos, venga como venga del Sheet o de un reporte: 3121234567 (número),
    3121234567.0 (decimal), '312-123-4567', '+52 312 123 4567', '044 312…'. Si no se puede (vacío, notación
    científica 3.12E+09 que ya perdió dígitos, menos de 10 dígitos), regresa "" ."""
    if v is None or (isinstance(v, float) and v != v):
        return ""
    if isinstance(v, float):
        if not v.is_integer():
            return ""
        v = int(v)
    s = str(v).strip()
    if re.search(r"[eE][+-]?\d", s):           # notación científica: el número ya se dañó
        return ""
    s = re.sub(r"\.0+$", "", s)                # «3121234567.0» leído como texto
    d = re.sub(r"\D", "", s)
    if len(d) > 10 and d.startswith("52"):
        d = d[2:]
    if len(d) > 10 and d[:1] in ("1", "0"):    # 1 (celular viejo) / 044 / 045
        d = d.lstrip("0")
        d = d[1:] if len(d) == 11 and d.startswith(("1", "4")) else d
    return d[-10:] if len(d) >= 10 else ""


TEL_COLUMNAS = {"CELULAR", "TELEFONO", "TELEFONO MOVIL", "TEL", "TELEFONO CELULAR"}


def _es_telefono(encabezado):
    return norm(encabezado) in TEL_COLUMNAS


def _es_fecha(encabezado):
    h = norm(encabezado)
    return h.startswith("FECHA") or h in ("FECHA INGRESO",)


class _BaseComun:
    """Lógica compartida; las subclases implementan _valores, _escribir, _crear_pestana, etc."""

    origen = "App"

    def _limpiar_cache(self, pestana):
        pass

    # ---------------------------------------------------------------- lectura
    def leer(self, pestana):
        """DataFrame con los encabezados de la fila 1 y la columna '_fila' (número de fila en la hoja)."""
        valores = self._valores(pestana)
        if not valores:
            return pd.DataFrame(columns=["_fila"])
        enc = [canonico(str(h).strip()) if h not in (None, "") else f"_col{i + 1}" for i, h in enumerate(valores[0])]
        vistos = {}
        for i, h in enumerate(enc):                       # encabezados repetidos → "X", "X (2)"
            if h in vistos:
                vistos[h] += 1
                enc[i] = f"{h} ({vistos[h]})"
            else:
                vistos[h] = 1
        filas = []
        for n, fila in enumerate(valores[1:], start=2):
            fila = list(fila) + [None] * (len(enc) - len(fila))
            if not any(v not in (None, "") for v in fila[:min(len(fila), 20)]):
                continue
            filas.append(fila[:len(enc)] + [n])
        df = pd.DataFrame(filas, columns=enc + ["_fila"])
        for c in enc:
            if _es_fecha(c):
                df[c] = df[c].map(de_serial)
            elif _es_telefono(c):          # siempre texto de 10 dígitos (Sheets los da como número o decimal)
                df[c] = df[c].map(lambda v: tel_texto(v) or (None if v is None or (isinstance(v, float) and v != v)
                                                              or str(v).strip() == "" else str(v).strip()))
        df = df.astype(object)
        return df.where(df.notna(), None)

    def columnas(self, pestana):
        valores = self._valores(pestana, solo_encabezado=True)
        return [canonico(str(h).strip()) if h not in (None, "") else "" for h in (valores[0] if valores else [])]

    def col(self, pestana, nombre):
        """Número de columna (1-based) por nombre de encabezado, sin importar acentos/mayúsculas."""
        objetivo = norm(canonico(nombre))
        for i, h in enumerate(self.columnas(pestana), start=1):
            if norm(h) == objetivo:
                return i
        return None

    def ultima_fila_con_datos(self, pestana, columnas_clave=(1, 3)):
        valores = self._valores(pestana)
        ultima = 1
        for n, fila in enumerate(valores[1:], start=2):
            if any(len(fila) >= c and fila[c - 1] not in (None, "") for c in columnas_clave):
                ultima = n
        return ultima

    # ---------------------------------------------------------------- escritura
    def escribir(self, pestana, celdas):
        """celdas: lista de (fila, columna, valor). Las fórmulas van como texto que empieza con '='."""
        if celdas:
            self._escribir(pestana, [(f, c, a_serial(v)) for f, c, v in celdas])

    def reemplazar(self, pestana, encabezados, filas, oculta=False):
        """Borra la pestaña y escribe la tabla completa (se usa para Bitacora y Capturas Cetelem)."""
        self._reemplazar(pestana, [list(encabezados)] + [[a_serial(v) for v in f] for f in filas], oculta)

    def marcar(self, pestana, celdas, color):
        """Pinta el fondo de las celdas [(fila, columna)] con color '#RRGGBB'."""
        if celdas:
            self._marcar(pestana, celdas, color)

    def quitar_marca(self, pestana, celdas):
        """Deja las celdas [(fila, columna)] sin relleno y sin nota."""
        if celdas:
            self._marcar(pestana, celdas, "#FFFFFF")
            self._quitar_notas(pestana, celdas)

    def reemplazar_textos(self, pestana, cambios):
        """Cambia texto dentro de las celdas (también dentro de fórmulas). cambios: [(viejo, nuevo)].
        Devuelve cuántas celdas cambiaron."""
        if pestana not in self.pestanas():
            return 0
        celdas = []
        for n, fila in enumerate(self._formulas(pestana), start=1):
            for m, v in enumerate(fila, start=1):
                if isinstance(v, str):
                    nv = v
                    for a, b in cambios:
                        nv = nv.replace(a, b)
                    if nv != v:
                        celdas.append((n, m, nv))
        if celdas:
            self._escribir(pestana, celdas)
        return len(celdas)

    def reemplazar_en_formatos(self, pestana, cambios, filas_min=0):
        """Cambia texto dentro de las reglas de formato condicional de la pestaña (p. ej. ="Contrato").
        cambios: [(viejo, nuevo)]; una regla cuyo valor es exactamente `viejo` sin comillas también cambia.
        filas_min: alarga hasta esa fila el rango de las reglas de la pestaña (en Google).
        Devuelve cuántas reglas cambiaron."""
        if pestana not in self.pestanas():
            return 0
        return self._reemplazar_en_formatos(pestana, cambios, filas_min)

    def dar_formato(self, pestana, numeros=None):
        """Un solo estilo para toda la pestaña (ESTILO): misma letra, tamaño y color, fondo blanco y encabezado
        gris con letra blanca. numeros: {encabezado: "fecha" | "entero" | "moneda"} para el formato de esas columnas.
        No toca el formato condicional (los colores por estatus)."""
        enc = self.columnas(pestana)
        cols = {i: tipo for i, h in enumerate(enc, start=1) for nombre, tipo in (numeros or {}).items()
                if norm(h) == norm(nombre)}
        self._dar_formato(pestana, len(enc), cols)

    # ---------------------------------------------------------------- estructura y formato (sin tocar datos)
    def agregar_columna(self, pestana, nombre):
        """Columna con ese encabezado; si no existe se agrega al final (después del último encabezado), con el
        formato de la columna anterior. Devuelve su número (1-based)."""
        c = self.col(pestana, nombre)
        if c:
            return c
        enc = self.columnas(pestana)
        c = max([i for i, h in enumerate(enc, start=1) if h] or [0]) + 1
        self.escribir(pestana, [(1, c, nombre)])
        if c > 1:
            self._copiar_formato_columna(pestana, c - 1, c)
        self._limpiar_cache(pestana)
        return c

    def renombrar_encabezado(self, pestana, viejo, nuevo):
        """Cambia el texto del encabezado (fila 1). Devuelve True si lo cambió."""
        valores = self._valores(pestana, solo_encabezado=True)
        for i, h in enumerate(valores[0] if valores else [], start=1):
            if norm(h) == norm(viejo) and str(h).strip() != nuevo:
                self.escribir(pestana, [(1, i, nuevo)])
                return True
        return False

    def formato_columnas(self, pestana, numeros, fila_fin=None):
        """Solo el formato de número de esas columnas {columna: "fecha" | "fechahora" | "entero" | "moneda"};
        no cambia letra, colores ni nada más."""
        if numeros:
            self._formato_columnas(pestana, numeros, fila_fin)

    def regla_color(self, pestana, columnas, formula, color, fila_fin=1000, letra_color=None):
        """Formato condicional: pinta el fondo de `columnas` (filas 2..fila_fin) cuando `formula` (escrita para la
        fila 2, p. ej. '=AND($V2<>"",$X2="")') es verdadera; letra_color, el color de la letra (opcional).
        No la duplica si ya existe. Devuelve True si la agregó."""
        return self._regla_color(pestana, columnas, formula, color, fila_fin, letra_color)

    def ocultar(self, pestana):
        if pestana in self.pestanas():
            self._ocultar(pestana)

    def borrar(self, pestana):
        """Elimina la pestaña (si existe)."""
        if pestana in self.pestanas():
            self._borrar(pestana)

    def ordenar(self, primeras):
        """Pone esas pestañas al principio, en ese orden (las demás quedan después, en su orden actual)."""
        orden = [p for p in primeras if p in self.pestanas()]
        orden += [p for p in self.pestanas() if p not in orden]
        self._ordenar(orden)

    def borrar_columna(self, pestana, col):
        """Elimina la columna (1-based) con sus datos y formato; las de la derecha se recorren."""
        if pestana in self.pestanas() and col:
            self._borrar_columna(pestana, col)
            self._limpiar_cache(pestana)

    def mover_columna(self, pestana, de, a):
        """Mueve la columna `de` para que quede en la posición `a` (1-based), con datos, formato, validaciones y
        colores; las fórmulas que la usan se ajustan solas."""
        if pestana in self.pestanas() and de and a and de != a:
            self._mover_columna(pestana, de, a)
            self._limpiar_cache(pestana)

    def congelar(self, pestana, filas=1, columnas=0):
        if pestana in self.pestanas():
            self._congelar(pestana, filas, columnas)

    def proteger(self, pestana, columnas=None, descripcion=""):
        """Protección «solo advertencia»: si alguien edita esas columnas (o toda la pestaña si columnas es None),
        Google pide confirmar. No duplica protecciones. En Excel no hace nada."""
        if pestana in self.pestanas():
            return self._proteger(pestana, sorted(set(columnas)) if columnas else None, descripcion)
        return False

    def filtro_completo(self, pestana, solo_si_hay=False):
        """Filtro de la pestaña sobre TODA la hoja (todas las filas y columnas, también las que se agreguen
        después). Un filtro con rango fijo (p. ej. A1:W121) deja fuera filas y columnas nuevas y la hoja se ve
        «seccionada». solo_si_hay: solo lo corrige si la pestaña ya tiene filtro. Devuelve True si lo puso."""
        if pestana in self.pestanas():
            return self._filtro_completo(pestana, solo_si_hay)
        return False

    def vista_filtrada(self, pestana, titulo, columna, valor):
        """Vista filtrada (Google) con columna = valor. No la duplica. En Excel no hace nada."""
        if pestana in self.pestanas():
            return self._vista_filtrada(pestana, titulo, columna, valor)
        return False

    def lista_valores(self, pestana, columnas, valores, fila_fin=1000):
        """Validación de lista desplegable en `columnas` (filas 2..fila_fin)."""
        self._lista_valores(pestana, columnas, valores, fila_fin)

    def validar_fecha(self, pestana, columnas, fila_fin=1000):
        self._validar_fecha(pestana, columnas, fila_fin)

    def agrupar_columnas(self, pestana, columnas):
        """Agrupa y contrae cada columna (se abre con el «+» de arriba). No duplica grupos. Devuelve cuántas agrupó."""
        return self._agrupar_columnas(pestana, sorted(set(columnas)))

    def anotar_historial(self, registros):
        """registros: dicts con las llaves de HISTORIAL_COLS (Fecha se llena sola)."""
        if not registros:
            return
        if TAB_HISTORIAL not in self.pestanas():
            self._reemplazar(TAB_HISTORIAL, [HISTORIAL_COLS], False)
        inicio = self.ultima_fila_con_datos(TAB_HISTORIAL, (1, 2)) + 1
        ahora = datetime.now()
        celdas = []
        for i, r in enumerate(registros):
            for j, k in enumerate(HISTORIAL_COLS, start=1):
                v = ahora if k == "Fecha" else r.get(k, "")
                if k == "Origen" and not v:
                    v = self.origen
                celdas.append((inicio + i, j, _texto_historial(v)))
        self.escribir(TAB_HISTORIAL, celdas)
        try:                                   # la fecha con formato de fecha y hora (no como número)
            self.formato_columnas(TAB_HISTORIAL, {1: "fechahora"})
        except Exception:
            pass


def _cambiar_texto(obj, cambios):
    """Aplica los cambios a todos los textos de una estructura (dict/list/str)."""
    if isinstance(obj, str):
        for a, b in cambios:
            obj = b.strip('"') if obj == a.strip('"') else obj.replace(a, b)
        return obj
    if isinstance(obj, list):
        return [_cambiar_texto(x, cambios) for x in obj]
    if isinstance(obj, dict):
        return {k: _cambiar_texto(v, cambios) for k, v in obj.items()}
    return obj


def _texto_historial(v):
    if isinstance(v, (datetime, date)) or v is None or isinstance(v, (int, float)):
        return "" if v is None else v
    return str(v)


# =============================================================================
# Google Sheets
# =============================================================================
_MES_ANIO = re.compile(r"(?i)^\s*(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|setiembre|"
                       r"octubre|noviembre|diciembre)\s+\d{4}\s*$")


def _como_texto(v):
    """Google interpreta "JUNIO 2026" como la fecha 01/06/2026; con apóstrofo se queda como texto."""
    return "'" + v if isinstance(v, str) and _MES_ANIO.match(v) else v


class BaseGoogle(_BaseComun):
    def __init__(self, info_cuenta_servicio, url):
        import gspread
        from gspread.http_client import BackOffHTTPClient
        self._gs = gspread
        # BackOffHTTPClient reintenta solo (con espera creciente) si Google responde 429 por cuota
        self.cliente = gspread.service_account_from_dict(dict(info_cuenta_servicio), http_client=BackOffHTTPClient)
        self.libro = self.cliente.open_by_url(url)
        self.url = url
        self.titulo = self.libro.title
        self._cache = {}
        self._hojas = None

    def _mapa(self):
        """{título: Worksheet}. Se pide a Google una sola vez por ejecución (cada consulta cuenta en la cuota)."""
        if self._hojas is None:
            self._hojas = {w.title: w for w in self.libro.worksheets()}
        return self._hojas

    def pestanas(self):
        return list(self._mapa())

    def _hoja(self, pestana):
        ws = self._mapa().get(pestana)
        if ws is None:                      # quizá se creó después de leer la lista
            self._hojas = None
            ws = self._mapa().get(pestana)
        if ws is None:
            raise ValueError(f"No existe la pestaña '{pestana}' en el Sheet.")
        return ws

    def _valores(self, pestana, solo_encabezado=False):
        clave = (pestana, solo_encabezado)
        if clave not in self._cache:
            from gspread.utils import ValueRenderOption, DateTimeOption
            ws = self._hoja(pestana)
            rango = "1:1" if solo_encabezado else None
            self._cache[clave] = ws.get(rango, value_render_option=ValueRenderOption.unformatted,
                                        date_time_render_option=DateTimeOption.serial_number,
                                        pad_values=True)
        return self._cache[clave]

    def _limpiar_cache(self, pestana):
        self._cache = {k: v for k, v in self._cache.items() if k[0] != pestana}

    def _escribir(self, pestana, celdas):
        from gspread.utils import ValueInputOption
        ws = self._hoja(pestana)
        max_fila = max(f for f, _, _ in celdas)
        max_col = max(c for _, c, _ in celdas)
        if max_fila > ws.row_count or max_col > ws.col_count:
            ws.resize(rows=max(ws.row_count, max_fila + 50), cols=max(ws.col_count, max_col))
        datos = [{"range": f"{letra(c)}{f}", "values": [["" if v is None else _como_texto(v)]]} for f, c, v in celdas]
        for i in range(0, len(datos), 500):            # lotes para no rebasar la cuota de la API
            ws.batch_update(datos[i:i + 500], value_input_option=ValueInputOption.user_entered)
        self._limpiar_cache(pestana)

    def _reemplazar(self, pestana, tabla, oculta):
        from gspread.utils import ValueInputOption
        ncols = max(len(f) for f in tabla)
        if pestana in self.pestanas():
            ws = self._hoja(pestana)
            ws.clear()
            if ws.row_count < len(tabla) or ws.col_count < ncols:
                ws.resize(rows=max(ws.row_count, len(tabla) + 50), cols=max(ws.col_count, ncols))
        else:
            ws = self.libro.add_worksheet(pestana, rows=len(tabla) + 50, cols=ncols)
            self._hojas = None
        tabla = [["" if v is None else _como_texto(v) for v in f] + [""] * (ncols - len(f)) for f in tabla]
        ws.update(tabla, "A1", value_input_option=ValueInputOption.user_entered)
        if oculta:
            ws.hide()
        else:
            try:                                   # el filtro debe abarcar las filas nuevas
                self._filtro_completo(pestana, solo_si_hay=True)
            except Exception:
                pass
        self._limpiar_cache(pestana)

    def _formulas(self, pestana):
        from gspread.utils import ValueRenderOption
        return self._hoja(pestana).get(value_render_option=ValueRenderOption.formula, pad_values=True)

    def _reemplazar_en_formatos(self, pestana, cambios, filas_min):
        ws = self._hoja(pestana)
        meta = self.libro.fetch_sheet_metadata(
            params={"fields": "sheets(properties(sheetId),conditionalFormats)"})
        hoja = next((h for h in meta.get("sheets", []) if h.get("properties", {}).get("sheetId") == ws.id), {})
        pedidos = []
        for i, regla in enumerate(hoja.get("conditionalFormats", [])):
            nueva = _cambiar_texto(regla, cambios)
            for r in nueva.get("ranges", []):          # que el color alcance a las filas nuevas
                if filas_min and r.get("endRowIndex") is not None and r["endRowIndex"] < filas_min:
                    r["endRowIndex"] = filas_min
            if nueva == regla:
                continue
            pedidos.append({"updateConditionalFormatRule": {"sheetId": ws.id, "index": i, "rule": nueva}})
        if pedidos:
            self.libro.batch_update({"requests": pedidos})
        return len(pedidos)

    def _dar_formato(self, pestana, ncols, numeros):
        ws = self._hoja(pestana)
        filas, ncols = ws.row_count, max(ncols, 1)
        rango = lambda f0, f1, c0, c1: {"sheetId": ws.id, "startRowIndex": f0, "endRowIndex": f1,
                                        "startColumnIndex": c0, "endColumnIndex": c1}
        texto = lambda color, negrita: {"fontFamily": ESTILO["fuente"], "fontSize": ESTILO["tamano"], "bold": negrita,
                                        "italic": False, "strikethrough": False, "underline": False,
                                        "foregroundColor": _rgb(color)}
        campos = "userEnteredFormat(textFormat,backgroundColor,horizontalAlignment,verticalAlignment,wrapStrategy,numberFormat)"
        pedidos = [
            {"repeatCell": {"range": rango(1, filas, 0, ws.col_count), "fields": campos, "cell": {"userEnteredFormat": {
                "textFormat": texto(ESTILO["texto"], False), "backgroundColor": _rgb(ESTILO["fondo"]),
                "horizontalAlignment": "LEFT", "verticalAlignment": "MIDDLE", "wrapStrategy": "CLIP"}}}},
            {"repeatCell": {"range": rango(0, 1, 0, ws.col_count), "fields": campos, "cell": {"userEnteredFormat": {
                "textFormat": texto(ESTILO["enc_texto"], True), "backgroundColor": _rgb(ESTILO["enc_fondo"]),
                "horizontalAlignment": "CENTER", "verticalAlignment": "MIDDLE", "wrapStrategy": "WRAP"}}}},
            {"updateSheetProperties": {"properties": {"sheetId": ws.id, "gridProperties": {"frozenRowCount": 1}},
                                       "fields": "gridProperties.frozenRowCount"}},
        ]
        for col, tipo in numeros.items():
            t, patron, _ = FORMATO_NUMERO[tipo]
            pedidos.append({"repeatCell": {
                "range": rango(1, filas, col - 1, col), "fields": "userEnteredFormat(numberFormat,horizontalAlignment)",
                "cell": {"userEnteredFormat": {"numberFormat": {"type": t, "pattern": patron},
                                               "horizontalAlignment": "LEFT" if tipo == "fecha" else "RIGHT"}}}})
        self.libro.batch_update({"requests": pedidos})

    def _rango(self, ws, f0, f1, c0, c1):
        """GridRange con índices 0-based y fin exclusivo."""
        return {"sheetId": ws.id, "startRowIndex": f0, "endRowIndex": f1, "startColumnIndex": c0, "endColumnIndex": c1}

    def _meta(self, ws, campos):
        meta = self.libro.fetch_sheet_metadata(params={"fields": f"sheets(properties(sheetId),{campos})"})
        return next((h for h in meta.get("sheets", []) if h.get("properties", {}).get("sheetId") == ws.id), {})

    def _copiar_formato_columna(self, pestana, de, a):
        ws = self._hoja(pestana)
        filas = ws.row_count
        if a > ws.col_count:
            ws.resize(cols=a)
        self.libro.batch_update({"requests": [{"copyPaste": {
            "source": self._rango(ws, 0, filas, de - 1, de), "destination": self._rango(ws, 0, filas, a - 1, a),
            "pasteType": "PASTE_FORMAT"}}]})

    def _formato_columnas(self, pestana, numeros, fila_fin):
        ws = self._hoja(pestana)
        fin = min(fila_fin or ws.row_count, ws.row_count)
        pedidos = []
        for col, tipo in numeros.items():
            t, patron, _ = FORMATO_NUMERO[tipo]
            pedidos.append({"repeatCell": {"range": self._rango(ws, 1, fin, col - 1, col),
                                           "fields": "userEnteredFormat.numberFormat",
                                           "cell": {"userEnteredFormat": {"numberFormat": {"type": t, "pattern": patron}}}}})
        self.libro.batch_update({"requests": pedidos})

    def _ocultar(self, pestana):
        self._hoja(pestana).hide()

    def _borrar(self, pestana):
        self.libro.del_worksheet(self._hoja(pestana))
        self._hojas = None
        self._limpiar_cache(pestana)

    def _ordenar(self, orden):
        mapa = self._mapa()
        self.libro.batch_update({"requests": [{"updateSheetProperties": {
            "properties": {"sheetId": mapa[p].id, "index": i}, "fields": "index"}} for i, p in enumerate(orden)]})
        self._hojas = None

    def _borrar_columna(self, pestana, col):
        ws = self._hoja(pestana)
        self.libro.batch_update({"requests": [{"deleteDimension": {"range": {
            "sheetId": ws.id, "dimension": "COLUMNS", "startIndex": col - 1, "endIndex": col}}}]})
        self._hojas = None                  # el número de columnas de la hoja cambió

    def _mover_columna(self, pestana, de, a):
        ws = self._hoja(pestana)
        self.libro.batch_update({"requests": [{"moveDimension": {
            "source": {"sheetId": ws.id, "dimension": "COLUMNS", "startIndex": de - 1, "endIndex": de},
            "destinationIndex": a - 1 if a < de else a}}]})

    def _congelar(self, pestana, filas, columnas):
        ws = self._hoja(pestana)
        self.libro.batch_update({"requests": [{"updateSheetProperties": {"properties": {
            "sheetId": ws.id, "gridProperties": {"frozenRowCount": filas, "frozenColumnCount": columnas}},
            "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}}]})

    def _proteger(self, pestana, columnas, descripcion):
        ws = self._hoja(pestana)
        hechas = {r.get("description") for r in self._meta(ws, "protectedRanges").get("protectedRanges", [])}
        if descripcion in hechas:
            return False
        if columnas:
            rangos = [self._rango(ws, 1, ws.row_count, c - 1, c) for c in columnas]
        else:
            rangos = [{"sheetId": ws.id}]
        self.libro.batch_update({"requests": [{"addProtectedRange": {"protectedRange": {
            "range": r, "description": descripcion, "warningOnly": True}}} for r in rangos]})
        return True

    def _filtro_completo(self, pestana, solo_si_hay):
        ws = self._hoja(pestana)
        actual = self._meta(ws, "basicFilter").get("basicFilter")
        if solo_si_hay and not actual:
            return False
        rango = (actual or {}).get("range", {})
        if actual and set(rango) <= {"sheetId"}:
            return False                          # ya cubre toda la hoja
        self.libro.batch_update({"requests": [{"setBasicFilter": {"filter": {"range": {"sheetId": ws.id}}}}]})
        return True

    def _vista_filtrada(self, pestana, titulo, columna, valor):
        ws = self._hoja(pestana)
        if titulo in {v.get("title") for v in self._meta(ws, "filterViews").get("filterViews", [])}:
            return False
        self.libro.batch_update({"requests": [{"addFilterView": {"filter": {
            "title": titulo, "range": {"sheetId": ws.id},
            "criteria": {str(columna - 1): {"condition": {"type": "TEXT_EQ", "values": [{"userEnteredValue": valor}]}}}}}}]})
        return True

    def _regla_color(self, pestana, columnas, formula, color, fila_fin, letra_color=None):
        ws = self._hoja(pestana)
        fila_fin = min(fila_fin, ws.row_count)          # un rango más allá de la hoja lo rechaza la API
        for regla in self._meta(ws, "conditionalFormats").get("conditionalFormats", []):
            valores = regla.get("booleanRule", {}).get("condition", {}).get("values", [])
            if any(v.get("userEnteredValue") == formula for v in valores):
                return False
        self.libro.batch_update({"requests": [{"addConditionalFormatRule": {"index": 0, "rule": {
            "ranges": [self._rango(ws, 1, fila_fin, c - 1, c) for c in columnas],
            "booleanRule": {"condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": formula}]},
                            "format": dict({"backgroundColor": _rgb(color)}, **(
                                {"textFormat": {"foregroundColor": _rgb(letra_color)}} if letra_color else {}))}}}}]})
        return True

    def _lista_valores(self, pestana, columnas, valores, fila_fin):
        ws = self._hoja(pestana)
        fila_fin = min(fila_fin, ws.row_count)
        regla = {"condition": {"type": "ONE_OF_LIST", "values": [{"userEnteredValue": v} for v in valores]},
                 "showCustomUi": True, "strict": False}
        columnas = [c for c in columnas if self._lista_actual(ws, c) != list(valores)]
        if columnas:            # si la lista ya es la misma no se toca (así no se pierden los colores de los chips)
            self.libro.batch_update({"requests": [{"setDataValidation": {
                "range": self._rango(ws, 1, fila_fin, c - 1, c), "rule": regla}} for c in columnas]})

    def _lista_actual(self, ws, col):
        """Opciones de la lista desplegable de la fila 2 de esa columna (o None)."""
        try:
            meta = self.libro.fetch_sheet_metadata(params={
                "ranges": f"'{ws.title}'!{letra(col)}2",
                "fields": "sheets(data(rowData(values(dataValidation))))"})
            celda = meta["sheets"][0]["data"][0]["rowData"][0]["values"][0]
            cond = celda.get("dataValidation", {}).get("condition", {})
            if cond.get("type") != "ONE_OF_LIST":
                return None
            return [v.get("userEnteredValue") for v in cond.get("values", [])]
        except (KeyError, IndexError, TypeError):
            return None

    def _validar_fecha(self, pestana, columnas, fila_fin):
        ws = self._hoja(pestana)
        fila_fin = min(fila_fin, ws.row_count)
        self.libro.batch_update({"requests": [{"setDataValidation": {
            "range": self._rango(ws, 1, fila_fin, c - 1, c),
            "rule": {"condition": {"type": "DATE_IS_VALID"}, "strict": False}}} for c in columnas]})

    def _agrupar_columnas(self, pestana, columnas):
        ws = self._hoja(pestana)
        agrupadas = set()
        for g in self._meta(ws, "columnGroups").get("columnGroups", []):
            agrupadas.update(range(g["range"].get("startIndex", 0) + 1, g["range"].get("endIndex", 0) + 1))
        nuevos = [c for c in columnas if c not in agrupadas]
        if not nuevos:
            return 0
        tramos = []                                # columnas seguidas en un solo grupo (Google une las vecinas)
        for c in nuevos:
            if tramos and c == tramos[-1][1] + 1:
                tramos[-1][1] = c
            else:
                tramos.append([c, c])
        rango = lambda a, b: {"sheetId": ws.id, "dimension": "COLUMNS", "startIndex": a - 1, "endIndex": b}
        self.libro.batch_update({"requests": [{"addDimensionGroup": {"range": rango(a, b)}} for a, b in tramos]})
        try:                                       # contraerlos (si Google los unió distinto, quedan abiertos)
            grupos = self._meta(ws, "columnGroups").get("columnGroups", [])
            self.libro.batch_update({"requests": [
                {"updateDimensionGroup": {"dimensionGroup": dict(g, collapsed=True), "fields": "collapsed"}}
                for g in grupos if any(g["range"].get("startIndex", 0) < c <= g["range"].get("endIndex", 0)
                                       for c in nuevos)]})
        except Exception:
            pass
        return len(nuevos)

    def _marcar(self, pestana, celdas, color):
        r, g, b = (int(color.lstrip("#")[i:i + 2], 16) / 255 for i in (0, 2, 4))
        formatos = [{"range": f"{letra(c)}{f}", "format": {"backgroundColor": {"red": r, "green": g, "blue": b}}}
                    for f, c in celdas]
        self._hoja(pestana).batch_format(formatos)

    def _quitar_notas(self, pestana, celdas):
        self._hoja(pestana).clear_notes([f"{letra(c)}{f}" for f, c in celdas])

    def respaldar(self, pestana):
        """Copia oculta de la pestaña antes de escribir; conserva las últimas RESPALDOS_A_GUARDAR."""
        prefijo = f"_resp {pestana[:20]} "
        nombre = prefijo + datetime.now().strftime("%Y-%m-%d %H%M%S")
        existentes = {w.title for w in self.libro.worksheets()}
        n = 2
        while nombre in existentes:          # dos respaldos en el mismo segundo: Google no admite nombres repetidos
            nombre = prefijo + datetime.now().strftime("%Y-%m-%d %H%M%S") + f" ({n})"
            n += 1
        copia = self._hoja(pestana).duplicate(new_sheet_name=nombre)
        copia.hide()
        viejas = sorted([w for w in self.libro.worksheets() if w.title.startswith(prefijo)], key=lambda w: w.title)
        for w in viejas[:-RESPALDOS_A_GUARDAR]:
            self.libro.del_worksheet(w)
        self._hojas = None
        return nombre

    def refrescar(self):
        self._cache = {}
        self._hojas = None


# =============================================================================
# Excel (modo sin conexión / pruebas)
# =============================================================================
class BaseExcel(_BaseComun):
    def __init__(self, contenido, nombre="Seguimiento.xlsx"):
        import openpyxl
        self._op = openpyxl
        self.libro = openpyxl.load_workbook(io.BytesIO(contenido))
        self.titulo = nombre
        self.url = None
        self.origen = "App (Excel)"

    def pestanas(self):
        return self.libro.sheetnames

    def _valores(self, pestana, solo_encabezado=False):
        ws = self.libro[pestana]
        filas = ws.iter_rows(min_row=1, max_row=1 if solo_encabezado else ws.max_row, values_only=True)
        out = []
        for f in filas:
            out.append([None if isinstance(v, str) and v.startswith("=") else v for v in f])
        return out

    def _escribir(self, pestana, celdas):
        ws = self.libro[pestana]
        for f, c, v in celdas:
            if isinstance(v, (int, float)) and _es_fecha(ws.cell(1, c).value) and 20000 < v < 80000:
                v = de_serial(v)
            ws.cell(f, c).value = v

    def _reemplazar(self, pestana, tabla, oculta):
        formatos = None
        if pestana in self.libro.sheetnames:
            idx = self.libro.sheetnames.index(pestana)
            formatos = self.libro[pestana].conditional_formatting      # se conservan los colores condicionales
            validaciones = self.libro[pestana].data_validations          # y las listas desplegables
            del self.libro[pestana]
            ws = self.libro.create_sheet(pestana, idx)
            ws.conditional_formatting = formatos
            ws.data_validations = validaciones
        else:
            ws = self.libro.create_sheet(pestana)
        for fila in tabla:
            ws.append([de_serial(v) if isinstance(v, (int, float)) and _es_fecha(tabla[0][i] if i < len(tabla[0]) else "")
                       else v for i, v in enumerate(fila)])
        if oculta:
            ws.sheet_state = "hidden"

    def _marcar(self, pestana, celdas, color):
        from openpyxl.styles import PatternFill
        relleno = (PatternFill(fill_type=None) if color.upper() in ("#FFFFFF", "FFFFFF")
                   else PatternFill("solid", fgColor="FF" + color.lstrip("#").upper()))
        ws = self.libro[pestana]
        for f, c in celdas:
            ws.cell(f, c).fill = relleno

    def _formulas(self, pestana):
        return [[c.value for c in fila] for fila in self.libro[pestana].iter_rows()]

    def _dar_formato(self, pestana, ncols, numeros):
        from openpyxl.styles import Alignment, Font, PatternFill
        ws = self.libro[pestana]
        sin = lambda c: c.lstrip("#").upper()
        letra_normal = Font(name=ESTILO["fuente"], size=ESTILO["tamano"], color="FF" + sin(ESTILO["texto"]))
        letra_enc = Font(name=ESTILO["fuente"], size=ESTILO["tamano"], bold=True, color="FF" + sin(ESTILO["enc_texto"]))
        fondo_enc = PatternFill("solid", fgColor="FF" + sin(ESTILO["enc_fondo"]))
        for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max(ws.max_column, ncols)):
            for c in fila:
                if c.row == 1:
                    c.font, c.fill = letra_enc, fondo_enc
                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    c.font, c.fill = letra_normal, PatternFill(fill_type=None)
                    tipo = numeros.get(c.column)
                    c.number_format = FORMATO_NUMERO[tipo][2] if tipo else "General"
                    c.alignment = Alignment(horizontal="right" if tipo in ("entero", "moneda") else "left",
                                            vertical="center")
        ws.freeze_panes = ws.freeze_panes or "A2"

    def _copiar_formato_columna(self, pestana, de, a):
        from copy import copy
        ws = self.libro[pestana]
        for f in range(1, ws.max_row + 1):
            origen, destino = ws.cell(f, de), ws.cell(f, a)
            if origen.has_style:
                destino.font, destino.fill = copy(origen.font), copy(origen.fill)
                destino.border, destino.alignment = copy(origen.border), copy(origen.alignment)
                destino.number_format = origen.number_format
        ancho = ws.column_dimensions[letra(de)].width
        if ancho:
            ws.column_dimensions[letra(a)].width = ancho

    def _formato_columnas(self, pestana, numeros, fila_fin):
        ws = self.libro[pestana]
        for col, tipo in numeros.items():
            for f in range(2, (fila_fin or ws.max_row) + 1):
                ws.cell(f, col).number_format = FORMATO_NUMERO[tipo][2]

    def _ocultar(self, pestana):
        self.libro[pestana].sheet_state = "hidden"

    def _borrar(self, pestana):
        self.libro.remove(self.libro[pestana])

    def _ordenar(self, orden):
        self.libro._sheets = [self.libro[p] for p in orden]
        visibles = [i for i, ws in enumerate(self.libro._sheets) if ws.sheet_state == "visible"]
        if visibles:
            self.libro.active = visibles[0]

    def _borrar_columna(self, pestana, col):
        n = self.libro[pestana].max_column
        _permutar_excel(self.libro[pestana], {c: (None if c == col else c - (c > col)) for c in range(1, n + 1)})

    def _mover_columna(self, pestana, de, a):
        orden = list(range(1, self.libro[pestana].max_column + 1))
        orden.insert(a - 1, orden.pop(de - 1))
        _permutar_excel(self.libro[pestana], {viejo: nuevo for nuevo, viejo in enumerate(orden, start=1)})

    def _congelar(self, pestana, filas, columnas):
        from openpyxl.utils import get_column_letter
        self.libro[pestana].freeze_panes = f"{get_column_letter(columnas + 1)}{filas + 1}"

    def _proteger(self, pestana, columnas, descripcion):
        return False

    def _vista_filtrada(self, pestana, titulo, columna, valor):
        return False

    def _filtro_completo(self, pestana, solo_si_hay):
        ws = self.libro[pestana]
        if solo_si_hay and not ws.auto_filter.ref:
            return False
        ref = f"A1:{letra(ws.max_column)}{ws.max_row}"
        if ws.auto_filter.ref == ref:
            return False
        ws.auto_filter.ref = ref
        return True

    def _regla_color(self, pestana, columnas, formula, color, fila_fin, letra_color=None):
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Font, PatternFill
        ws = self.libro[pestana]
        f_sin = formula.lstrip("=")
        for rango in ws.conditional_formatting:
            for r in rango.rules:
                if any(str(x).lstrip("=") == f_sin for x in (r.formula or [])):
                    return False
        sqref = " ".join(f"{letra(c)}2:{letra(c)}{fila_fin}" for c in columnas)
        relleno = PatternFill("solid", fgColor="FF" + color.lstrip("#").upper(), bgColor="FF" + color.lstrip("#").upper())
        letra_ = Font(color="FF" + letra_color.lstrip("#").upper()) if letra_color else None
        ws.conditional_formatting.add(sqref, FormulaRule(formula=[f_sin], fill=relleno, font=letra_))
        return True

    def _quitar_validaciones(self, ws, columnas, fila_fin):
        from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
        quitar = [CellRange(f"{letra(c)}2:{letra(c)}{fila_fin}") for c in columnas]
        for dv in list(ws.data_validations.dataValidation):
            quedan = [r for r in dv.sqref.ranges if not any(r.min_col == q.min_col and r.max_col == q.max_col
                                                            for q in quitar)]
            if len(quedan) != len(dv.sqref.ranges):
                if quedan:
                    dv.sqref = MultiCellRange(" ".join(r.coord for r in quedan))
                else:
                    ws.data_validations.dataValidation.remove(dv)

    def _lista_valores(self, pestana, columnas, valores, fila_fin):
        from openpyxl.worksheet.datavalidation import DataValidation
        ws = self.libro[pestana]
        self._quitar_validaciones(ws, columnas, fila_fin)
        dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True,
                            showErrorMessage=False)
        dv.sqref = " ".join(f"{letra(c)}2:{letra(c)}{fila_fin}" for c in columnas)
        ws.add_data_validation(dv)

    def _validar_fecha(self, pestana, columnas, fila_fin):
        from openpyxl.worksheet.datavalidation import DataValidation
        ws = self.libro[pestana]
        self._quitar_validaciones(ws, columnas, fila_fin)
        dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                            showErrorMessage=False)
        dv.sqref = " ".join(f"{letra(c)}2:{letra(c)}{fila_fin}" for c in columnas)
        ws.add_data_validation(dv)

    def _agrupar_columnas(self, pestana, columnas):
        from copy import copy
        ws = self.libro[pestana]
        n = 0
        for c in columnas:
            # openpyxl guarda a veces varias columnas en una sola dimensión (P:Q); se separa para no ocultar de más
            for k, d in list(ws.column_dimensions.items()):
                if d.min and d.max and d.min <= c <= d.max and d.min != d.max:
                    for j in range(d.min, d.max + 1):
                        nueva = copy(d)
                        nueva.index, nueva.min, nueva.max = letra(j), j, j
                        ws.column_dimensions[letra(j)] = nueva
                    break
            d = ws.column_dimensions[letra(c)]
            if not d.outline_level:
                d.outline_level, d.hidden = 1, True
                n += 1
        return n

    def _reemplazar_en_formatos(self, pestana, cambios, filas_min):
        n = 0
        for rango in self.libro[pestana].conditional_formatting:
            for regla in rango.rules:
                nuevas = [_cambiar_texto(f, cambios) for f in (regla.formula or [])]
                if nuevas != list(regla.formula or []):
                    regla.formula = nuevas
                    n += 1
        return n

    def _quitar_notas(self, pestana, celdas):
        ws = self.libro[pestana]
        for f, c in celdas:
            ws.cell(f, c).comment = None

    def respaldar(self, pestana):
        prefijo = f"_resp {pestana[:12]} "
        nombre = prefijo + datetime.now().strftime("%m%d %H%M%S")
        copia = self.libro.copy_worksheet(self.libro[pestana])
        copia.title = nombre[:31]
        copia.sheet_state = "hidden"
        viejas = [ws for ws in self.libro.worksheets if ws.title.startswith(prefijo) and ws is not copia]
        for ws in viejas[:max(0, len(viejas) - (RESPALDOS_A_GUARDAR - 1))]:
            self.libro.remove(ws)
        return copia.title

    def refrescar(self):
        pass

    def guardar(self):
        buf = io.BytesIO()
        self.libro.save(buf)
        return buf.getvalue()


# =============================================================================
# Columnas reacomodadas: referencias de celdas
# =============================================================================
_REF_CELDA = re.compile(r"(?<![A-Za-z_!.'\d$])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(])")
_REF_OTRA_HOJA = re.compile(r"((?:'[^']+'|[A-Za-z_][\w.]*)!\$?[A-Z]{1,3}\$?\d*(?::\$?[A-Z]{1,3}\$?\d*)?)")
_REF_COLUMNAS = re.compile(r"(?<![A-Za-z_!.'\d$])(\$?)([A-Z]{1,3}):(\$?)([A-Z]{1,3})(?![A-Za-z\d(])")


def _num_col(letras):
    n = 0
    for ch in letras:
        n = n * 26 + ord(ch) - 64
    return n


def cambiar_referencias(texto, mapa):
    """Cambia las letras de columna de referencias sin pestaña («C5», «$D$2:$D9», «C:C») según mapa
    {columna vieja: columna nueva o None si se borró}. No toca lo que está entre comillas."""
    if not isinstance(texto, str):
        return texto
    m = lambda l: (lambda n: "#REF!" if n is None else letra(n))(mapa.get(_num_col(l), _num_col(l)))
    def cambiar(t):
        t = _REF_COLUMNAS.sub(lambda x: f"{x.group(1)}{m(x.group(2))}:{x.group(3)}{m(x.group(4))}", t)
        return _REF_CELDA.sub(lambda x: f"{x.group(1)}{m(x.group(2))}{x.group(3)}{x.group(4)}", t)

    partes = texto.split('"')
    for i in range(0, len(partes), 2):        # fuera de comillas; lo que apunta a otra pestaña (Hoja!A1:B2) no
        trozos = _REF_OTRA_HOJA.split(partes[i])
        partes[i] = "".join(t if j % 2 else cambiar(t) for j, t in enumerate(trozos))
    return '"'.join(partes)


def _permutar_excel(ws, mapa):
    """Reacomoda las columnas de una hoja de openpyxl: mapa {columna vieja: nueva o None (se borra)}. Mueve valores,
    estilos, anchos, colores condicionales y validaciones, y ajusta las fórmulas de la hoja."""
    from copy import copy
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.styles.cell_style import StyleArray
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    n, filas = ws.max_column, ws.max_row
    celdas = {(c.row, c.column): (c.value, copy(c._style)) for fila in ws.iter_rows(1, filas, 1, n) for c in fila}
    anchos = {c: ws.column_dimensions[letra(c)].width for c in range(1, n + 1)}
    for (f, c), (v, _) in celdas.items():
        ws.cell(f, c).value, ws.cell(f, c)._style = None, StyleArray()
    for (f, c), (v, estilo) in celdas.items():
        nueva = mapa.get(c, c)
        if nueva is None:
            continue
        celda = ws.cell(f, nueva)
        celda.value = cambiar_referencias(v, mapa) if isinstance(v, str) and v.startswith("=") else v
        celda._style = estilo
    for c, w in anchos.items():
        nueva = mapa.get(c, c)
        if nueva is not None:
            ws.column_dimensions[letra(nueva)].width = w

    def rangos(sqref):
        out = []
        for r in sqref.ranges:
            for c in range(r.min_col, r.max_col + 1):
                nueva = mapa.get(c, c)
                if nueva is not None:
                    out.append(CellRange(min_col=nueva, max_col=nueva, min_row=r.min_row, max_row=r.max_row).coord)
        return " ".join(out)

    nuevas = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        rng = rangos(cf.sqref)
        if rng:
            for regla in cf.rules:
                regla.formula = [cambiar_referencias(f, mapa) for f in (regla.formula or [])]
                nuevas.add(rng, regla)
    ws.conditional_formatting = nuevas
    for dv in list(ws.data_validations.dataValidation):
        rng = rangos(dv.sqref)
        if rng:
            dv.sqref = MultiCellRange(rng)
            dv.formula1 = cambiar_referencias(dv.formula1, mapa)
        else:
            ws.data_validations.dataValidation.remove(dv)
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{letra(max(v for v in mapa.values() if v))}{filas}"


def a_excel(base, pestanas):
    """Arma un .xlsx con las pestañas indicadas (lo usan los módulos de cuadre que leen archivos)."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for p in pestanas:
        ws = wb.create_sheet(p[:31])
        valores = base._valores(p)
        enc = valores[0] if valores else []
        fechas = {i for i, h in enumerate(enc) if _es_fecha(h)}
        for fila in valores:
            ws.append([de_serial(v) if i in fechas and isinstance(v, (int, float)) else v
                       for i, v in enumerate(fila)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
