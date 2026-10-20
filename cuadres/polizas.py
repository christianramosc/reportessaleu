# -*- coding: utf-8 -*-
"""
polizas.py  (también funciona suelto en Colab como extraer_polizas.py)
----------
Extrae de pólizas de seguro de auto (PDF) los datos:
NOMBRE, DATOS DE VEHÍCULO, SERIE, MOTOR, RFC, DIRECCIÓN, CP y VIGENCIA
y además: folio de crédito, tipo de persona, fecha de nacimiento y edad (del RFC), sexo, municipio,
estado, modelo MG, paquete, suma asegurada, deducibles, RC, prima neta/total/anual, forma de pago,
tipo de póliza (anual/multianual), beneficiario preferente y agencia.
Cada PDF trae las 2 pólizas del crédito: la 1 (primer año, de contado) y la 2 (resto del plazo).
Con ellas calcula el plazo del crédito y TODAS las oportunidades (aniversarios) para que el cliente
contrate póliza externa o pague de contado su seguro.
Genera un Excel con 2 hojas:
  - Contacto: una fila por cliente con su PRÓXIMA oportunidad, fechas y estatus del día
    (las columnas de seguimiento se conservan si vuelves a subir el Excel anterior).
  - Base de datos: una fila por crédito con todos los datos importantes de sus 2 pólizas.
  (Solo si algún PDF no se leyó completo aparece una hoja "Por revisar".)

Aseguradoras reconocidas: GNP, HDI, Quálitas, Chubb y Seguros Inbursa.
Cualquier otra se intenta con un lector genérico y se marca "REVISAR".

Entrada (cualquiera de estas):
  - un .zip (aunque pese varios GB: se lee archivo por archivo, sin descomprimir todo)
  - varios .zip (por ejemplo, si partiste el archivo grande en partes)
  - una carpeta con PDFs y/o zips (incluye subcarpetas y zips dentro de zips)

Requisitos:
  Colab:   !apt-get -qq install -y poppler-utils
           !pip -q install openpyxl
  Windows: instalar Poppler (https://github.com/oschwartz10612/poppler-windows/releases)
           y agregar su carpeta "bin" al PATH;  pip install openpyxl

Uso:
  Colab  -> ejecuta la celda, da clic en "Elegir archivos" y sube los PDFs / zip
            (y si quieres, el Excel anterior). El resultado se descarga solo.
  Consola-> python extraer_polizas.py seguros.zip [otro.zip | carpeta ...] -o Polizas.xlsx
"""

import io
import os
import re
import sys
import shutil
import zipfile
import tempfile
import subprocess
from datetime import date
from concurrent.futures import ThreadPoolExecutor, as_completed

# ============================ CONFIGURACIÓN (Colab) ============================
# En Colab, al ejecutar la celda aparece el botón "Elegir archivos" para subir:
#   - los PDFs de las pólizas y/o uno o varios .zip con ellos
#   - (opcional) el Polizas_extraidas.xlsx anterior, para conservar tu seguimiento
# Al terminar, el Excel se descarga solo a tu computadora.
CARPETA_SUBIDA = "/content/polizas_subidas"
ARCHIVO_SALIDA = "/content/Polizas_extraidas.xlsx"
HILOS = 4            # PDFs que se procesan en paralelo

# Reglas del contrato (días naturales ANTES de cada aniversario de la póliza).
# Cómo funciona: cada crédito trae 2 pólizas -> la 1 (de contado al comprar) y la 2 (resto del plazo).
# En CADA aniversario antes del fin del crédito el cliente puede:
#   - contratar póliza externa: a más tardar DIAS_POLIZA_EXTERNA días antes del aniversario (luego se cierra)
#   - pagar de contado su seguro: desde DIAS_PAGO_CONTADO días antes hasta el día anterior
# Si no hace ninguna, ese año el seguro se financia y le sube la mensualidad.
# También se pueden cambiar después en la hoja "Parámetros" del Excel.
DIAS_POLIZA_EXTERNA = 45   # último día para póliza externa (días antes del aniversario)
DIAS_PAGO_CONTADO = 30     # abre la ventana para pagar de contado
DIAS_AVISO_PREVIO = 7      # llamar al menos N días antes del límite de póliza externa
# ==============================================================================

RE_RFC = r"[A-ZÑ&]{3,4}\d{6}[A-Z0-9]{3}"
RE_VIN = r"\b[A-HJ-NPR-Z0-9]{17}\b"
MESES = {"ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6, "JUL": 7,
         "AGO": 8, "SEP": 9, "SET": 9, "OCT": 10, "NOV": 11, "DIC": 12,
         "JAN": 1, "APR": 4, "AUG": 8, "DEC": 12}

CAMPOS = ["ARCHIVO", "FOLIO CRÉDITO", "ASEGURADORA", "NO. PÓLIZA", "PÓLIZA DEL CRÉDITO", "NOMBRE", "TIPO PERSONA", "RFC",
          "FECHA NAC. (RFC)", "EDAD", "SEXO", "TELÉFONO", "CORREO", "DIRECCIÓN", "MUNICIPIO", "ESTADO", "CP",
          "DATOS DE VEHÍCULO", "MODELO MG", "AÑO MODELO", "SERIE", "MOTOR", "PAQUETE", "SUMA ASEGURADA DM",
          "DEDUCIBLE DM", "DEDUCIBLE RT", "RC TERCEROS", "PRIMA NETA", "PRIMA TOTAL", "FORMA DE PAGO",
          "VIGENCIA DESDE", "VIGENCIA HASTA", "AÑOS DE VIGENCIA", "TIPO DE PÓLIZA", "PRIMA ANUAL EST.",
          "BENEFICIARIO PREFERENTE", "AGENCIA", "OBSERVACIONES"]
OBLIGATORIOS = ["NOMBRE", "DATOS DE VEHÍCULO", "SERIE", "MOTOR", "RFC",
                "DIRECCIÓN", "CP", "VIGENCIA DESDE", "VIGENCIA HASTA"]


# ------------------------------------------------------------------ utilidades
def limpiar(s):
    return re.sub(r"\s+", " ", s or "").strip(" ,;-")


def a_fecha(txt):
    """'24/NOV/2023' o '25/11/2023' -> date"""
    if not txt:
        return None
    txt = txt.strip().upper()
    m = re.match(r"(\d{1,2})[/-]([A-Z]{3})[/-](\d{4})", txt)
    if m and m.group(2) in MESES:
        return date(int(m.group(3)), MESES[m.group(2)], int(m.group(1)))
    m = re.match(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", txt)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def sumar_anios(f, n):
    try:
        return f.replace(year=f.year + n)
    except ValueError:  # 29 de febrero
        return f.replace(year=f.year + n, day=28)


def buscar(patron, texto, grupo=1, flags=0):
    m = re.search(patron, texto, flags)
    return limpiar(m.group(grupo)) if m else ""


def pdf_a_paginas(ruta_pdf):
    """Texto con layout de cada página (usa pdftotext de Poppler)."""
    r = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", ruta_pdf, "-"],
                       capture_output=True)
    texto = r.stdout.decode("utf-8", errors="ignore")
    if r.returncode != 0 and not texto.strip():
        raise RuntimeError("PDF dañado o protegido")
    return texto.split("\f")


# -------------------------------------------------------------- lectores por aseguradora
def es_gnp(p):
    return "No. Póliza" in p and "Carátula" in p and "Contrato" in p


def leer_gnp(p):
    lin = p.splitlines()
    d = {"ASEGURADORA": "GNP"}
    d["NO. PÓLIZA"] = buscar(r"No\. Póliza\s+(\S+)", p)
    d["VIGENCIA DESDE"] = a_fecha(buscar(r"Desde las 12 hrs\s+(\S+)", p))
    d["VIGENCIA HASTA"] = a_fecha(buscar(r"Hasta las 12 hrs\s+(\S+)", p))

    i_nom = next((i for i, l in enumerate(lin) if re.search(r"\sContrato\s", l)), None)
    i_ref = next((i for i, l in enumerate(lin) if "Referencia Cliente" in l), None)
    if i_nom is not None:
        d["NOMBRE"] = limpiar(lin[i_nom].split("Contrato")[0])
    if i_nom is not None and i_ref is not None:
        bloque = []
        for l in lin[i_nom + 1:i_ref]:
            l = re.split(r"\s{2,}(?:Desde las|Hasta las|Duración)", l)[0]
            if l.strip():
                bloque.append(l.strip())
        bloque = " ".join(bloque)
        rfc = re.match(r"(" + RE_RFC + r")\s+(.*)", bloque)
        if rfc:
            d["RFC"] = rfc.group(1)
            bloque = rfc.group(2)
        d["CP"] = buscar(r"C\.?\s?P\.?\s*(\d{5})", bloque)
        d["DIRECCIÓN"] = limpiar(re.sub(r",?\s*C\.?\s?P\.?\s*\d{5}", "", bloque))

    # vehículo: renglón con la serie (VIN), renglón siguiente (versión), renglón con año y motor
    if i_ref is not None:
        for j in range(i_ref, min(i_ref + 15, len(lin))):
            m = re.search(RE_VIN, lin[j])
            if m:
                d["SERIE"] = m.group(0)
                desc1 = lin[j][:m.start()]
                desc2 = re.split(r"\s{2,}\$|\s{3,}", lin[j + 1].strip())[0] if j + 1 < len(lin) else ""
                desc = limpiar(desc1 + " " + desc2)
                desc = re.sub(r"^[A-Z0-9]*\d[A-Z0-9]*\s+\d+\s+(?:AUTO|PICK ?UP|CAMION\w*)?\s*", "", desc)
                d["DATOS DE VEHÍCULO"] = desc
                for k in range(j + 1, min(j + 6, len(lin))):
                    ma = re.match(r"\s*((?:19|20)\d{2})\s{2,}.*?(\S+)\s*$", lin[k])
                    if ma:
                        d["AÑO MODELO"] = ma.group(1)
                        d["MOTOR"] = ma.group(2)
                        break
                break
    return d


def es_hdi(p):
    return "HDI" in p and "Núm. de Motor:" in p and "Vigencia: Desde" in p


def leer_hdi(p):
    lin = p.splitlines()
    d = {"ASEGURADORA": "HDI"}
    d["NO. PÓLIZA"] = buscar(r"Póliza:\s*(\S+)", p)
    d["VIGENCIA DESDE"] = a_fecha(buscar(r"Desde las [\d:]+ hrs\. del\s+(\d{2}/\d{2}/\d{4})", p))
    d["VIGENCIA HASTA"] = a_fecha(buscar(r"Hasta las [\d:]+ hrs\. del\s+(\d{2}/\d{2}/\d{4})", p))
    d["RFC"] = buscar(r"RFC:\s*(" + RE_RFC + ")", p)
    d["SERIE"] = buscar(r"Serie:\s*(\S+)", p)
    d["MOTOR"] = buscar(r"Núm\. de Motor:\s*(\S+)", p)

    i_rfc = next((i for i, l in enumerate(lin) if "RFC:" in l), None)
    i_pol = next((i for i, l in enumerate(lin) if re.search(r"^\s*Póliza:", l)), None)
    if i_rfc is not None:
        # nombre: último renglón con texto arriba del RFC (o en el mismo, a la izquierda)
        izq = lin[i_rfc].split("RFC:")[0].strip()
        if izq:
            d["NOMBRE"] = limpiar(izq)
        else:
            for l in reversed(lin[:i_rfc]):
                if l.strip():
                    d["NOMBRE"] = limpiar(l)
                    break
    if i_rfc is not None and i_pol is not None:
        dom = " ".join(l.replace("Datos Generales", "").strip()
                       for l in lin[i_rfc + 1:i_pol] if l.strip())
        dom = re.sub(r"^(?:Datos|Generales)\s+", "", limpiar(dom))
        d["CP"] = buscar(r"C\.?\s?P\.?\s*(\d{5})", dom)
        dom = re.sub(r"C\.?\s?P\.?\s*\d{5}", "", dom)
        d["TELÉFONO"] = re.sub(r"\D", "", buscar(r"Tel\s*:?\s*([\d\- ]{7,})", dom))
        dom = re.sub(r"Tel\s*:?\s*[\d\- ]+", "", dom)
        d["DIRECCIÓN"] = limpiar(dom)

    i_cla = next((i for i, l in enumerate(lin) if "Clave:" in l), None)
    if i_cla is not None:
        marca = limpiar(lin[i_cla].split("Clave:")[0])
        d["AÑO MODELO"] = buscar(r"((?:19|20)\d{2})\s*$", marca)
        version = ""
        for l in lin[i_cla + 1:i_cla + 3]:
            m = re.search(r"Versión:\s*(.+?)(?:\s{2,}|$)", l)
            if m:
                version = m.group(1)
                break
        marca = re.sub(r"\s*(?:19|20)\d{2}\s*$", "", marca).replace(",", "")
        d["DATOS DE VEHÍCULO"] = limpiar(marca + " " + version)
    return d


def es_qualitas(p):
    return "PÓLIZA DE SEGURO DE AUTOMÓVILES" in p and "DESCRIPCIÓN DEL VEHÍCULO" in p and "Serie:" in p


def leer_qualitas(p):
    lin = p.splitlines()
    d = {"ASEGURADORA": "QUÁLITAS"}
    i_hdr = next((i for i, l in enumerate(lin) if re.search(r"PÓLIZA\s+ENDOSO\s+INCISO", l)), None)
    if i_hdr is not None:
        for l in lin[i_hdr:i_hdr + 3]:
            m = re.search(r"\b(\d{8,})\b", l)
            if m:
                d["NO. PÓLIZA"] = m.group(1)
                break
    i_aseg = next((i for i, l in enumerate(lin) if "INFORMACIÓN DEL ASEGURADO" in l), None)
    if i_aseg is not None:
        for l in lin[i_aseg + 1:i_aseg + 3]:
            if l.strip():
                d["NOMBRE"] = limpiar(l)
                break
    d["RFC"] = buscar(r"R\.F\.C\.?:?\s*(" + RE_RFC + ")", p)  # formato nuevo "R.F.C.:" y viejo "R.F.C."
    calle = buscar(r"Domicilio:\s*(.*?)(?:\s+Número:|\s{2,}|$)", p, flags=re.M)
    num = buscar(r"Número:\s*(.*?)(?:\s+Interior:|\s{2,}|$)", p, flags=re.M)
    inte = buscar(r"Interior:(?: ?(\S.*?))?(?:\s{2,}|$)", p, flags=re.M)  # vacío si no hay interior
    col = buscar(r"Colonia:\s*(.*?)(?:\s{2,}|$)", p, flags=re.M)
    mun = buscar(r"Municipio:\s*(.*?)(?:\s+Estado:|\s{2,}|$)", p, flags=re.M)
    edo = buscar(r"Estado:\s*(.*?)(?:\s{2,}|$)", p, flags=re.M)
    partes = [limpiar(f"{calle} {num}" + (f" INT {inte}" if inte else ""))]
    partes += [f"COL {col}" if col else "", mun, edo]
    d["DIRECCIÓN"] = ", ".join(x for x in partes if x)
    d["CP"] = buscar(r"C\.P\.:\s*(\d{5})", p)

    i_veh = next((i for i, l in enumerate(lin) if "DESCRIPCIÓN DEL VEHÍCULO" in l), None)
    if i_veh is not None:
        for l in lin[i_veh + 1:i_veh + 3]:
            if l.strip():
                desc = re.sub(r"^\s*\d+\s+", "", l)
                d["DATOS DE VEHÍCULO"] = limpiar(re.sub(r"\s+(?:19|20)\d{2}\s*$", "", desc))
                break
    d["AÑO MODELO"] = buscar(r"Modelo:\s*(\d{4})", p)
    d["SERIE"] = buscar(r"Serie:\s*(\S+)", p)
    d["MOTOR"] = buscar(r"Motor:\s*(\S+)", p)
    d["VIGENCIA DESDE"] = a_fecha(buscar(r"DESDE LAS .*?DEL\s+(\d{2}/\d{2}/\d{4})", p))
    d["VIGENCIA HASTA"] = a_fecha(buscar(r"HASTA LAS .*?DEL\s+(\d{2}/\d{2}/\d{4})", p))
    return d


def es_chubb(p):
    return "Chubb Seguros" in p and "Datos del asegurado" in p and "Descripción del vehículo" in p


def leer_chubb(p):
    lin = p.splitlines()
    d = {"ASEGURADORA": "CHUBB"}
    d["NO. PÓLIZA"] = buscar(r"Póliza:\s*([A-Z]{0,3}\s?\d+)", p)
    d["VIGENCIA DESDE"] = a_fecha(buscar(r"Vigencia:\s*Del\s+(\S+)", p))
    d["VIGENCIA HASTA"] = a_fecha(buscar(r"Vigencia:\s*Del\s+\S+\s+[\d:]+\s+horas\s+al\s+(\S+)", p))
    d["NOMBRE"] = buscar(r"^\s*Asegurado:\s*(.+?)\s*$", p, flags=re.M)
    d["RFC"] = buscar(r"RFC:\s*(" + RE_RFC + ")", p)
    d["CP"] = buscar(r"C\.P\.:\s*(\d{5})", p)
    d["TELÉFONO"] = re.sub(r"\D", "", buscar(r"Teléfono:\s*([\d\- ]{7,})", p))

    # domicilio: renglón "Domicilio:" + siguientes hasta "Datos generales", solo la columna izquierda
    i_dom = next((i for i, l in enumerate(lin) if l.strip().startswith("Domicilio:")), None)
    if i_dom is not None:
        partes = []
        for l in lin[i_dom:i_dom + 4]:
            if "Datos generales" in l:
                break
            izq = re.split(r"\s{3,}(?:C\.P\.|Teléfono|RFC)", l)[0]
            partes.append(izq.replace("Domicilio:", "").strip())
        d["DIRECCIÓN"] = limpiar(", ".join(x for x in partes if x))

    veh = buscar(r"Descripción del vehículo\*?:\s*(.+?)(?:\s{2,}|$)", p, flags=re.M)
    marca = buscar(r"Marca:\s*(.+?)(?:\s{2,}|$)", p, flags=re.M)
    d["DATOS DE VEHÍCULO"] = limpiar(f"{marca} {veh}") if veh else ""
    d["AÑO MODELO"] = buscar(r"Modelo:\s*(\d{4})", p)
    d["SERIE"] = buscar(r"Serie:\s*(\S+)", p)
    d["MOTOR"] = buscar(r"Motor:\s*(\S+)", p)
    return d


def es_inbursa(p):
    return "Seguros Inbursa" in p and "CARÁTULA DE LA PÓLIZA" in p and "DATOS DEL CONTRATANTE" in p


def leer_inbursa(p):
    lin = p.splitlines()
    d = {"ASEGURADORA": "INBURSA"}
    d["NO. PÓLIZA"] = buscar(r"PÓLIZA\s+(\d+\s+\d+)", p)
    d["VIGENCIA DESDE"] = a_fecha(buscar(r"Desde las [\d:]+ horas\s+(\S+)", p))
    d["VIGENCIA HASTA"] = a_fecha(buscar(r"hasta las [\d:]+ horas\s+(\S+)", p))

    i_nom = next((i for i, l in enumerate(lin) if re.match(r"\s*NOMBRE:", l)), None)
    if i_nom is not None:
        for l in lin[i_nom + 1:i_nom + 3]:
            if l.strip():
                d["NOMBRE"] = limpiar(re.split(r"\s{3,}", l.strip())[0])
                break
    i_dir = next((i for i, l in enumerate(lin) if re.match(r"\s*DIRECCIÓN:", l)), None)
    i_prod = next((i for i, l in enumerate(lin) if re.match(r"\s*PRODUCTO:", l)), None)
    if i_dir is not None and i_prod is not None:
        partes = []
        for l in lin[i_dir + 1:i_prod]:
            if not l.strip() or re.match(r"\s{30,}", l):  # renglones solo con la columna derecha
                continue
            partes.append(re.split(r"\s{3,}", l.strip())[0])
        dom = limpiar(" ".join(partes))
        d["CP"] = buscar(r"C\.?\s?P\.?\s*(\d{5})", dom)
        dom = re.sub(r"\s*C\.?\s?P\.?\s*\d{5}", "", dom)
        dom = re.sub(r"\s+NUM INT 0\b", "", dom)
        d["DIRECCIÓN"] = limpiar(dom)
    d["RFC"] = buscar(r"R\.F\.C\.:.*?\n.*?\b(" + RE_RFC + r")\b", p, flags=re.S)
    comp = buscar(r"COMPLEMENTO:\s*(.+?)(?:\s{2,}|$)", p, flags=re.M)
    d["DATOS DE VEHÍCULO"] = limpiar(re.sub(r"^[A-Z0-9]*\d[A-Z0-9]*\s+", "", comp))  # quita la clave
    d["AÑO MODELO"] = buscar(r"MODELO:\s*(\d{4})", p)
    d["SERIE"] = buscar(r"NÚMERO DE SERIE:\s*(\S+)", p)
    d["MOTOR"] = buscar(r"NÚMERO DE MOTOR:\s*(\S+)", p)
    return d


def leer_generico(paginas):
    """Para aseguradoras no reconocidas: rescata lo que se pueda con patrones generales."""
    p = "\n".join(paginas[:3])
    d = {"ASEGURADORA": "NO RECONOCIDA"}
    m = re.search(RE_VIN, p)
    d["SERIE"] = m.group(0) if m else ""
    d["RFC"] = buscar(r"(?:R\.?F\.?C\.?\s*:?\s*)(" + RE_RFC + ")", p) or buscar(r"\b(" + RE_RFC + r")\b", p)
    d["MOTOR"] = buscar(r"Motor\s*:?\s*([A-Z0-9]{6,})", p, flags=re.I)
    d["CP"] = buscar(r"C\.?\s?P\.?\s*:?\s*(\d{5})", p)
    fechas = re.findall(r"\b\d{1,2}/(?:\d{1,2}|[A-Za-z]{3})/\d{4}\b", p)
    if len(fechas) >= 2:
        d["VIGENCIA DESDE"], d["VIGENCIA HASTA"] = a_fecha(fechas[0]), a_fecha(fechas[1])
    return d


LECTORES = [(es_gnp, leer_gnp), (es_hdi, leer_hdi), (es_qualitas, leer_qualitas),
            (es_chubb, leer_chubb), (es_inbursa, leer_inbursa)]

DOMINIOS_NO_CLIENTE = r"gnp|qualitas|hdi|chubb|condusef|marsh|axa|mapfre|zurich|segur|inbursa|cetelem|banorte|afirme|ana\.com|gob\.mx"


def correo_cliente(texto):
    """Primer correo del PDF que no sea de la aseguradora / agente / autoridad."""
    for c in re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", texto):
        if not re.search(DOMINIOS_NO_CLIENTE, c, re.I):
            return c.upper()
    return ""


# ---------------------------------------------------------------- datos adicionales
MUNICIPIOS = ["VILLA DE ALVAREZ", "VILLA DE ÁLVAREZ", "TECOMAN", "TECOMÁN", "MANZANILLO", "ARMERIA", "ARMERÍA",
              "COMALA", "CUAUHTEMOC", "CUAUHTÉMOC", "IXTLAHUACAN", "IXTLAHUACÁN", "COQUIMATLAN", "COQUIMATLÁN",
              "MINATITLAN", "MINATITLÁN", "COLIMA", "GUADALAJARA", "ZAPOPAN", "TLAJOMULCO", "TLAQUEPAQUE",
              "TONALA", "TONALÁ", "EL SALTO", "ZAPOTLAN EL GRANDE", "CIUDAD GUZMAN", "CD GUZMAN", "TOLIMAN",
              "TOLIMÁN", "TUXPAN", "AUTLAN", "AUTLÁN", "CIHUATLAN", "CIHUATLÁN", "TAMAZULA", "PIHUAMO",
              "TECALITLAN", "TECALITLÁN", "ZAPOTILTIC", "TONILA", "SAYULA", "MORELIA", "URUAPAN", "COALCOMAN",
              "AQUILA", "TEPALCATEPEC"]
ESTADOS = ["COLIMA", "JALISCO", "MICHOACAN", "MICHOACÁN", "NAYARIT", "GUANAJUATO", "CIUDAD DE MEXICO", "ESTADO DE MEXICO"]
SIN_ACENTO = str.maketrans("ÁÉÍÓÚ", "AEIOU")


def numero(s):
    try:
        return float(str(s).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def nacimiento_rfc(rfc):
    """Persona física (13 caracteres): AAAA + AAMMDD."""
    m = re.match(r"^[A-ZÑ&]{4}(\d{2})(\d{2})(\d{2})", rfc or "")
    if not m:
        return None
    yy, mm, dd = int(m.group(1)), int(m.group(2)), int(m.group(3))
    anio = 2000 + yy if 2000 + yy <= date.today().year - 16 else 1900 + yy
    try:
        return date(anio, mm, dd)
    except ValueError:
        return None


def modelo_mg(desc):
    """Nombre corto del modelo: MG5, MG3, ONE, HS, ZS, GT, RX5..."""
    d = (desc or "").upper()
    m = re.search(r"\b(?:MG\s*)?(MG\s?\d|EHS|ONE|GT|HS|ZS|RX5|RX8|RX9|CYBERSTER|MARVEL)\b", d.replace("MG MOTOR", "MG"))
    if not m:
        return ""
    v = m.group(1).replace(" ", "")
    return v if v.startswith("MG") or not v.isdigit() else "MG" + v


def montos(linea):
    return [numero(x) for x in re.findall(r"\d[\d,]*\.\d{2}", linea)]


def datos_extra(r, texto):
    """Agrega al registro r los datos adicionales que se pueden leer del texto completo de la póliza."""
    a = r.get("ASEGURADORA", "")
    lin = texto.splitlines()
    up = texto.upper()

    m = re.search(r"Poliza_(\d{8})_", r.get("ARCHIVO", ""))
    r["FOLIO CRÉDITO"] = m.group(1) if m else ""

    rfc = r.get("RFC", "")
    r["TIPO PERSONA"] = "MORAL" if len(rfc) == 12 else ("FÍSICA" if len(rfc) == 13 else "")
    r["FECHA NAC. (RFC)"] = nacimiento_rfc(rfc) if r["TIPO PERSONA"] == "FÍSICA" else None
    sexo = buscar(r"\d{2}\s+Años\s+(MASCULINO|FEMENINO)", texto)
    r["SEXO"] = sexo

    dom = (r.get("DIRECCIÓN", "") + " " + r.get("CP", "")).upper().translate(SIN_ACENTO)
    r["MUNICIPIO"] = next((mu.translate(SIN_ACENTO) for mu in sorted(MUNICIPIOS, key=len, reverse=True)
                           if re.search(r"\b" + mu.translate(SIN_ACENTO) + r"\b", dom)), "")
    est = next((e.translate(SIN_ACENTO) for e in ESTADOS if re.search(r"\b" + e.translate(SIN_ACENTO) + r"\b", dom)), "")
    if not est and re.search(r"\bJAL\b", dom):
        est = "JALISCO"
    if not est and re.search(r"\bCOL\b\s*$", dom.strip()) or (not est and r.get("CP", "").startswith("28")):
        est = "COLIMA"
    if not est and r.get("CP", "")[:2] in ("44", "45", "46", "47", "48", "49"):
        est = "JALISCO"
    r["ESTADO"] = est
    r["MODELO MG"] = modelo_mg(r.get("DATOS DE VEHÍCULO", ""))

    # paquete / plan
    paq = (buscar(r"Paquete:\s*([A-ZÁÉÍÓÚ ]+?)(?:\s{2,}|$)", texto, flags=re.M)
           or buscar(r"PLAN:\s*([A-ZÁÉÍÓÚ ]+?)(?:\s{2,}|$)", texto, flags=re.M)
           or buscar(r"AGENCIAS\s*-\s*([A-ZÁÉÍÓÚ ]+?)(?:\s{2,}|$)", texto, flags=re.M))
    if not paq and a == "INBURSA" and re.search(r"Daños materiales", texto, re.I):
        paq = "AMPLIA (AUTOTAL)"
    r["PAQUETE"] = paq

    # coberturas principales: suma asegurada y deducibles
    l_dm = next((l for l in lin if re.match(r"\s*Daños Materiales(?: Pérdida Parcial)?\s{2,}", l, re.I)), "")
    l_rt = next((l for l in lin if re.match(r"\s*Robo Total\s{2,}", l, re.I)), "")
    sa = montos(l_dm)
    if "VALOR FACTURA" in l_dm.upper() and not re.search(r"\d{3},\d{3}\.\d\d", l_dm.split("Valor")[0] if "Valor" in l_dm else l_dm):
        r["SUMA ASEGURADA DM"] = "VALOR FACTURA"  # Inbursa: no trae monto, cubre a valor factura
    else:
        r["SUMA ASEGURADA DM"] = sa[0] if sa and sa[0] and sa[0] > 20000 else ""

    def deducible(l):
        m = re.search(r"(\d{1,2}(?:\.\d+)?)\s?%", l)
        if m:
            return float(m.group(1)) / 100
        m = re.search(r"Valor Factura\s+(?:\d\s+años\s+)?(\d{1,2})\s+\d+", l, re.I)  # Inbursa: % y UMA en columnas
        return int(m.group(1)) / 100 if m else None
    r["DEDUCIBLE DM"] = deducible(l_dm)
    r["DEDUCIBLE RT"] = deducible(l_rt)
    l_rc = next((l for l in lin if re.match(r"\s*Responsabilidad Civil(?: por Daños a Terceros| Por Daños a Terceros| \(Límite Único| a Terceros catastrófica|$)", l, re.I)
                 and montos(l)), "")
    rc = [x for x in montos(l_rc) if x and x >= 100000]
    r["RC TERCEROS"] = rc[0] if rc else None
    if a == "INBURSA":  # RC dividida en bienes + personas + catastrófica: se suma el total
        i_rc = [i for i, l in enumerate(lin) if re.match(r"\s*Responsabilidad civil (?:por daños a Terceros|a Terceros catastrófica)", l, re.I)]
        tot = 0
        for i in i_rc[:3]:
            v = [x for x in montos(lin[i]) + (montos(lin[i + 1]) if i + 1 < len(lin) and not montos(lin[i]) else []) if x and x >= 100000]
            tot += v[0] if v else 0
        r["RC TERCEROS"] = tot or r["RC TERCEROS"]

    # primas y forma de pago
    neta = total = None
    fp = ""
    if a == "GNP":
        total = numero(buscar(r"Importe Total Actual\s+\$\s*([\d,]+\.\d\d)", texto))
        i_ref = next((i for i, l in enumerate(lin) if "Referencia Cliente" in l), None)
        if i_ref is not None:
            for l in lin[i_ref:i_ref + 12]:
                v = re.search(r"\$\s*([\d,]+\.\d\d)\s*$", l)
                if v:
                    neta = numero(v.group(1))
                    break
        fp = "ANUAL" if re.search(r"Intermediario\s+Anual", texto) else buscar(r"Intermediario\s+(\w+)", texto).upper()
    elif a == "QUÁLITAS":
        neta = numero(buscar(r"Prima Neta\s+\$\s*([\d,]+\.\d\d)", texto))
        total = numero(buscar(r"IMPORTE TOTAL\.?\s+\$\s*([\d,]+\.\d\d)", texto))
        fp = buscar(r"Forma de Pago:\s*(\S+)", texto)
    elif a == "CHUBB":
        neta = numero(buscar(r"Prima neta\s+([\d,]+\.\d\d)", texto))
        total = numero(buscar(r"Prima Total:\s*([\d,]+\.\d\d)", texto))
        fp = buscar(r"Forma de pago:\s*([A-Z ]+?)(?:\s{2,}|$)", texto, flags=re.M)
    elif a == "HDI":
        i_h = next((i for i, l in enumerate(lin) if "Prima Neta" in l and "Descuento" in l), None)
        if i_h is not None:
            for l in lin[i_h + 1:i_h + 4]:
                v = montos(l)
                if len(v) >= 3:
                    neta, total = v[0], v[-1]
                    break
        fp = buscar(r"^\s*(CONTADO[ A-Z]*?|ANUAL|SEMESTRAL|TRIMESTRAL|MENSUAL)\s*$", texto, flags=re.M)
    elif a == "INBURSA":
        i_n = next((i for i, l in enumerate(lin) if "PRIMA NETA:" in l), None)
        if i_n is not None and i_n + 1 < len(lin):
            v = montos(lin[i_n + 1])
            neta = v[-1] if v else None
        total = numero(buscar(r"(?:INICIAL|RENOVACI[OÓ]N|ENDOSO)\s+([\d,]+\.\d\d)", texto))
        i_f = next((i for i, l in enumerate(lin) if "FORMA DE PAGO:" in l), None)
        if i_f is not None and i_f + 1 < len(lin):
            fp = re.split(r"\s{3,}", lin[i_f + 1].strip())[0]
    r["PRIMA NETA"] = neta
    r["PRIMA TOTAL"] = total
    r["FORMA DE PAGO"] = limpiar(fp).upper()

    d0, d1 = r.get("VIGENCIA DESDE"), r.get("VIGENCIA HASTA")
    anios = max(1, round((d1 - d0).days / 365.25)) if d0 and d1 else None
    r["AÑOS DE VIGENCIA"] = anios
    r["TIPO DE PÓLIZA"] = "" if not anios else ("ANUAL" if anios == 1 else f"MULTIANUAL ({anios} AÑOS)")
    r["PRIMA ANUAL EST."] = round(total / anios, 2) if total and anios else None

    r["BENEFICIARIO PREFERENTE"] = ("AUTOEXPRESS INBURSA" if "AUTOEXPRESS INBURSA" in up
                                    else "CETELEM" if "CETELEM" in up else "")
    ag = (buscar(r"agencia de origen\s+(.+?)\s+\d{6,}", texto)
          or buscar(r"Agencia:\s*\d+\s+(.+?)(?:\s{2,}|$)", texto, flags=re.M)
          or buscar(r"Conducto:\s*\d+\s*-\s*(.+?)(?:\s{2,}|$)", texto, flags=re.M)
          or buscar(r"DISTRIBUIDORA:\s*\d+\s+(.+?)(?:\s{2,}|$)", texto, flags=re.M)
          or buscar(r'preferencia con "([^"]+)"', texto))
    r["AGENCIA"] = ag
    return r


# -------------------------------------------------------------- procesamiento
def procesar_pdf(ruta_pdf, nombre_archivo):
    try:
        paginas = pdf_a_paginas(ruta_pdf)
    except Exception as e:
        return [], f"No se pudo leer ({e})"
    registros = []
    for pag in paginas:
        for detecta, lee in LECTORES:
            if detecta(pag):
                try:
                    r = lee(pag)
                    if r.get("SERIE") or r.get("NO. PÓLIZA"):
                        registros.append(r)
                except Exception as e:  # una página rara no detiene todo
                    registros.append({"OBSERVACIONES": f"Error leyendo página: {e}"})
                break
    if not registros:
        if not "".join(paginas).strip():
            return [], "PDF sin texto (escaneado, requiere OCR)"
        registros = [leer_generico(paginas)]

    # Una misma póliza aparece varias veces (copias); se queda la versión más completa.
    mejores = {}
    for r in registros:
        clave = (r.get("ASEGURADORA"), r.get("NO. PÓLIZA") or r.get("SERIE"))
        llenos = sum(1 for c in OBLIGATORIOS if r.get(c))
        if clave not in mejores or llenos > mejores[clave][0]:
            mejores[clave] = (llenos, r)
    correo = correo_cliente("\n".join(paginas))
    salida = []
    texto_completo = "\n".join(paginas)
    for _, r in mejores.values():
        r["ARCHIVO"] = nombre_archivo
        # datos extra solo con las páginas de ESTA póliza (cada PDF trae 2 pólizas)
        num = (r.get("NO. PÓLIZA") or "").split()[-1] if r.get("NO. PÓLIZA") else ""
        propias = [pg for pg in paginas if num and num in pg] if len(mejores) > 1 else []
        try:
            datos_extra(r, "\n".join(propias) if propias else texto_completo)
        except Exception as e:  # los datos extra nunca detienen la extracción
            r["OBSERVACIONES"] = (r.get("OBSERVACIONES", "") + f" Datos extra incompletos: {e}").strip()
        if correo and not r.get("CORREO"):
            r["CORREO"] = correo
        faltan = [c for c in OBLIGATORIOS if not r.get(c)]
        obs = []
        if r.get("ASEGURADORA") == "NO RECONOCIDA":
            obs.append("Formato no reconocido: REVISAR")
        if faltan:
            obs.append("Falta: " + ", ".join(faltan))
        if r.get("SERIE") and not re.fullmatch(RE_VIN, r["SERIE"]):
            obs.append("Serie con formato inusual")
        if r.get("OBSERVACIONES"):
            obs.insert(0, r["OBSERVACIONES"])
        r["OBSERVACIONES"] = "; ".join(obs)
        salida.append(r)
    return salida, None


def iterar_pdfs(rutas, tmpdir):
    """Genera (ruta_temporal_pdf, nombre_mostrado). Lee zips miembro por miembro."""
    contador = [0]

    def desde_zip(zf, prefijo):
        for info in zf.infolist():
            if info.is_dir():
                continue
            nombre = info.filename
            low = nombre.lower()
            if low.endswith(".pdf"):
                contador[0] += 1
                destino = os.path.join(tmpdir, f"{contador[0]:07d}.pdf")
                with zf.open(info) as src, open(destino, "wb") as dst:
                    shutil.copyfileobj(src, dst, 1024 * 1024)
                yield destino, prefijo + nombre
            elif low.endswith(".zip"):
                with zf.open(info) as src:
                    datos = io.BytesIO(src.read())
                with zipfile.ZipFile(datos) as interno:
                    yield from desde_zip(interno, prefijo + nombre + "/")

    for ruta in rutas:
        if os.path.isdir(ruta):
            for raiz, _, archivos in os.walk(ruta):
                for a in sorted(archivos):
                    completo = os.path.join(raiz, a)
                    rel = os.path.relpath(completo, ruta)
                    if a.lower().endswith(".pdf"):
                        yield completo, rel
                    elif a.lower().endswith(".zip"):
                        with zipfile.ZipFile(completo) as zf:
                            yield from desde_zip(zf, rel + "/")
        elif ruta.lower().endswith(".zip"):
            with zipfile.ZipFile(ruta) as zf:
                yield from desde_zip(zf, "")
        elif ruta.lower().endswith(".pdf"):
            yield ruta, os.path.basename(ruta)
        else:
            print(f"  Omitido (no es pdf/zip/carpeta): {ruta}")


def extraer(rutas, hilos=4, zip_revisar=None):
    """Devuelve (filas, por_revisar). por_revisar = [(archivo, motivo)].
    Si zip_revisar tiene ruta, ahí se copian los PDFs que no se pudieron leer completos,
    listos para mandarlos y agregar su formato."""
    if shutil.which("pdftotext") is None:
        sys.exit("Falta pdftotext (Poppler). En Colab: !apt-get -qq install -y poppler-utils")
    filas, errores = [], []
    zrev = zipfile.ZipFile(zip_revisar, "w", zipfile.ZIP_DEFLATED) if zip_revisar else None
    tmpdir = tempfile.mkdtemp(prefix="polizas_")
    try:
        with ThreadPoolExecutor(max_workers=hilos) as ex:
            pendientes = {}
            n = 0

            def recoger(listos):
                nonlocal n
                for fut in listos:
                    ruta_tmp, nombre = pendientes.pop(fut)
                    regs, err = fut.result()
                    filas.extend(regs)
                    motivos = [err] if err else []
                    motivos += [r["OBSERVACIONES"] for r in regs if r.get("OBSERVACIONES")]
                    if motivos:
                        errores.append((nombre, " | ".join(dict.fromkeys(motivos))))
                        if zrev is not None:
                            zrev.write(ruta_tmp, nombre)
                    if ruta_tmp.startswith(tmpdir):
                        os.remove(ruta_tmp)  # libera espacio en cuanto se procesa
                    n += 1
                    if n % 25 == 0:
                        print(f"  {n} PDFs procesados...")

            for ruta_pdf, nombre in iterar_pdfs(rutas, tmpdir):
                fut = ex.submit(procesar_pdf, ruta_pdf, nombre)
                pendientes[fut] = (ruta_pdf, nombre)
                if len(pendientes) >= hilos * 4:  # no llenar el disco con temporales
                    hechos = [f for f in pendientes if f.done()] or [next(as_completed(pendientes))]
                    recoger(hechos)
            recoger(list(as_completed(pendientes)))
            print(f"  {n} PDFs procesados en total.")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
        if zrev is not None:
            zrev.close()
            if not errores:
                os.remove(zip_revisar)
    filas.sort(key=lambda r: (r.get("NOMBRE") or "", r.get("VIGENCIA DESDE") or date.min))
    return filas, errores


# -------------------------------------------------------------- calendario de contacto
def agrupar_creditos(filas):
    """Junta las pólizas de un mismo vehículo (serie) = un crédito. Marca cuál es la 1 y cuál la 2."""
    grupos = {}
    for r in filas:
        if not (r.get("VIGENCIA DESDE") and r.get("VIGENCIA HASTA")) or r.get("ASEGURADORA") == "NO RECONOCIDA":
            continue
        grupos.setdefault(r.get("SERIE") or r.get("NOMBRE"), []).append(r)
    creditos = []
    for clave, pols in grupos.items():
        pols.sort(key=lambda r: r["VIGENCIA DESDE"])
        if len(pols) == 1:
            pols[0]["PÓLIZA DEL CRÉDITO"] = "ÚNICA EN EL PDF"
        else:
            pols[0]["PÓLIZA DEL CRÉDITO"] = "1 · CONTADO (INICIO)"
            for r in pols[1:]:
                r["PÓLIZA DEL CRÉDITO"] = "2 · RESTO DEL PLAZO"
        inicio = pols[0]["VIGENCIA DESDE"]
        fin = max(r["VIGENCIA HASTA"] for r in pols)
        plazo = max(1, round((fin - inicio).days / 365.25))
        # oportunidades = cada aniversario antes del fin del crédito
        # la 1a oportunidad es cuando vence la póliza de contado (normalmente 1 año, en créditos viejos a veces 2)
        base = pols[0]["VIGENCIA HASTA"] if len(pols) > 1 else sumar_anios(inicio, 1)
        anivs = []
        n = 0
        while sumar_anios(base, n) < fin:
            anivs.append(sumar_anios(base, n))
            n += 1
        nota = ""
        if len(pols) == 1 and plazo == 1:
            anivs = [fin]  # solo viene la póliza del 1er año: al menos su vencimiento es una oportunidad
            nota = "Solo viene la póliza del 1er año: no se conoce el fin del crédito"
        elif len(pols) == 1:
            nota = "Solo viene una póliza multianual en el PDF"
        creditos.append({"clave": clave, "polizas": pols, "inicio": inicio, "fin": fin, "plazo": plazo,
                         "aniversarios": anivs, "nota": nota})
    return creditos


def calendario(filas, hoy=None):
    """Una fila por cada OPORTUNIDAD que todavía no pasa: cada aniversario antes del fin del crédito."""
    hoy = hoy or date.today()
    eventos = []
    for c in agrupar_creditos(filas):
        pols, anivs = c["polizas"], c["aniversarios"]
        ult = pols[-1]
        total = len(anivs)
        for k, f in enumerate(anivs, start=1):
            if f <= hoy:  # el último día para actuar es el día anterior al aniversario
                continue
            vence = next((r for r in pols if r["VIGENCIA DESDE"] < f <= r["VIGENCIA HASTA"]), ult)
            sigue = next((r for r in pols if r["VIGENCIA DESDE"] <= f < r["VIGENCIA HASTA"]), vence)
            anio = (f.year - c["inicio"].year) + 1  # año del crédito que empieza en este aniversario
            eventos.append({
                "CLIENTE": ult.get("NOMBRE", ""), "TELÉFONO": ult.get("TELÉFONO", ""),
                "CORREO": ult.get("CORREO", ""),
                "VEHÍCULO": limpiar(f"{ult.get('DATOS DE VEHÍCULO', '')} {ult.get('AÑO MODELO', '')}"),
                "SERIE": ult.get("SERIE", ""), "ASEGURADORA": sigue.get("ASEGURADORA", ""),
                "PÓLIZA": sigue.get("NO. PÓLIZA", ""),
                "OPORTUNIDAD": f"{k} de {total}",
                "QUÉ SE DECIDE": ((f"Seguro del año {anio} de {c['plazo']}" if anio <= c["plazo"]
                                   else f"Seguro del año {anio} (plazo no conocido)")
                                  + (" · 1er año después del contado" if k == 1 else "")),
                "FECHA ANIVERSARIO": f, "FIN DEL CRÉDITO": c["fin"], "PLAZO (AÑOS)": c["plazo"],
                "RFC": ult.get("RFC", ""), "DIRECCIÓN": ult.get("DIRECCIÓN", ""), "CP": ult.get("CP", ""),
                "FOLIO CRÉDITO": ult.get("FOLIO CRÉDITO", ""), "MUNICIPIO": ult.get("MUNICIPIO", ""),
                "MODELO MG": ult.get("MODELO MG", ""), "PAQUETE": sigue.get("PAQUETE", ""),
                "SUMA ASEGURADA DM": sigue.get("SUMA ASEGURADA DM", ""),
                "PRIMA ANUAL EST.": sigue.get("PRIMA ANUAL EST."), "FECHA NAC. (RFC)": ult.get("FECHA NAC. (RFC)"),
                "NOTA": c["nota"],
            })
    eventos.sort(key=lambda e: (e["FECHA ANIVERSARIO"], e["CLIENTE"]))
    return eventos


# -------------------------------------------------------------- Excel
RESULTADOS = ["Interesado", "Cotización enviada", "Vendida (póliza externa)", "Pagará de contado",
              "Se le financiará", "No interesado", "Sin respuesta", "Número equivocado"]

# Estatus de la próxima oportunidad, en el orden en que van pasando: (texto, qué hacer, fondo, letra)
ESTATUS_LISTA = [
    ("PRÓXIMO", "Todavía no; esperar a la fecha LLAMAR DESDE", "DDEBF7", "1F3864"),
    ("LLAMAR YA · PÓLIZA EXTERNA", "Llamar y cerrar la póliza externa antes del ÚLTIMO DÍA PÓLIZA EXTERNA", "FFC7CE", "9C0006"),
    ("EXTERNA CERRADA · ESPERAR CONTADO", "Ya no hay póliza externa; recordarle que pague de contado cuando abra la ventana", "FCE4D6", "C55A11"),
    ("VENTANA DE CONTADO ABIERTA", "Que pague de contado antes del ÚLTIMO DÍA CONTADO para no financiar el seguro", "FFEB9C", "9C5700"),
    ("PRÓXIMO · TERMINA SU CRÉDITO", "Su crédito termina pronto: el auto será suyo y su póliza vence ese día", "E4DFEC", "5B2C83"),
    ("LLAMAR YA · TERMINA SU CRÉDITO", "Ofrecerle su nuevo seguro antes de que venza su póliza (ya sin restricciones del contrato)", "CDB4E0", "3F1A5C"),
    ("CRÉDITO TERMINADO · CLIENTE LIBRE", "El auto ya es suyo: ofrecer seguro en cualquier momento, sin contrato de por medio", "C9E7E1", "0B5345"),
    ("CONTACTADO", "Registrar el RESULTADO", "C6EFCE", "006100"),
]


def seguimiento_previo(salida):
    """Si el Excel ya existe, recupera lo capturado a mano en la hoja Contacto
    (Contactado, Fecha, Resultado, Notas) por número de serie."""
    previo = {}
    if not os.path.exists(salida):
        return previo
    try:
        from openpyxl import load_workbook
        wb = load_workbook(salida)
        if "Contacto" not in wb.sheetnames:
            return previo
        ws = wb["Contacto"]
        fila_hdr = next((r for r in range(1, 15) if ws.cell(r, 1).value == "ESTATUS HOY"), None)
        if not fila_hdr:
            return previo
        hdr = [c.value for c in ws[fila_hdr]]
        need = ["SERIE", "CONTACTADO", "FECHA CONTACTO", "RESULTADO", "NOTAS"]
        if not all(h in hdr for h in need):
            return previo
        ix = {h: hdr.index(h) for h in need}
        for row in ws.iter_rows(min_row=fila_hdr + 1, values_only=True):
            datos = tuple(row[ix[h]] for h in need[1:])
            if row[ix["SERIE"]] and any(datos):
                previo[row[ix["SERIE"]]] = datos
    except Exception as e:
        print(f"  Aviso: no se pudo recuperar el seguimiento anterior ({e})")
    return previo


def _roles(c):
    """Póliza 1 (contado) y póliza 2 (resto del plazo) de un crédito, y la fecha de la 1a oportunidad."""
    p = c["polizas"]
    if len(p) > 1:
        return p[0], p[-1], p[0]["VIGENCIA HASTA"], c["fin"]
    if c["plazo"] == 1:          # solo vino la póliza de contado: no se conoce el fin del crédito
        return p[0], None, p[0]["VIGENCIA HASTA"], None
    return None, p[0], sumar_anios(c["inicio"], 1), c["fin"]   # una sola póliza multianual


def guardar_excel(filas, errores, salida):
    """Excel con 2 hojas: 'Contacto' (a quién llamar y cuándo) y 'Base de datos' (todo lo importante)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter as L

    AZUL = PatternFill("solid", fgColor="1F3864")
    AMARILLO = PatternFill("solid", fgColor="FFF2CC")
    fino = Side(style="thin", color="BFBFBF")
    BORDE = Border(top=fino, bottom=fino, left=fino, right=fino)
    MONEDA = '"$"#,##0.00'
    F10 = Font(name="Arial", size=10)

    def relleno(c):
        return PatternFill("solid", fgColor=c)

    def tabla(ws, fila_hdr, cols, anchos, datos, fechas=(), texto=(), centro=(), moneda=(), pct=(), captura=()):
        for j, h in enumerate(cols, start=1):
            c = ws.cell(fila_hdr, j, h)
            c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            c.fill = AZUL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDE
            ws.column_dimensions[L(j)].width = anchos.get(h, 14)
        ws.row_dimensions[fila_hdr].height = 42
        for i, fila in enumerate(datos, start=fila_hdr + 1):
            for j, (h, v) in enumerate(zip(cols, fila), start=1):
                c = ws.cell(i, j, v)
                c.font = F10
                c.border = BORDE
                c.alignment = Alignment(vertical="center")
                if h in fechas:
                    c.number_format = "DD/MM/YYYY"
                    c.alignment = Alignment(horizontal="center", vertical="center")
                if h in texto:
                    c.number_format = "@"
                if h in moneda:
                    c.number_format = MONEDA
                if h in pct:
                    c.number_format = "0%"
                    c.alignment = Alignment(horizontal="center", vertical="center")
                if h in centro:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                if h in captura:
                    c.fill = AMARILLO
                    c.font = Font(name="Arial", size=10, color="0000FF")
        ultima = max(fila_hdr + len(datos), fila_hdr + 1)
        ws.auto_filter.ref = f"A{fila_hdr}:{L(len(cols))}{ultima}"
        return ultima

    previo = seguimiento_previo(salida)
    creditos = agrupar_creditos(filas)   # también marca "PÓLIZA DEL CRÉDITO" en cada póliza
    hoy = date.today()

    wb = Workbook()
    wc = wb.active
    wc.title = "Contacto"
    wbd = wb.create_sheet("Base de datos")

    # ======================= Hoja CONTACTO =======================
    # Parámetros arriba (celdas amarillas editables)
    wc["A1"] = "CONTACTO · PÓLIZAS DE SEGURO (una fila por cliente con su PRÓXIMA oportunidad: aniversario, fin de crédito o auto liberado)"
    wc["A1"].font = Font(name="Arial", size=13, bold=True, color="1F3864")
    params = [("A2", "Último día para póliza externa (días antes del aniversario):", "D2", DIAS_POLIZA_EXTERNA),
              ("E2", "Abre ventana de contado (días antes):", "H2", DIAS_PAGO_CONTADO),
              ("I2", "Llamar N días antes del límite de externa:", "K2", DIAS_AVISO_PREVIO)]
    for ce, et, cv, v in params:
        wc[ce] = et
        wc[ce].font = Font(name="Arial", size=9, bold=True)
        wc[cv] = v
        wc[cv].font = Font(name="Arial", size=10, bold=True, color="0000FF")
        wc[cv].fill = AMARILLO
        wc[cv].alignment = Alignment(horizontal="center")
    PE, PC, PA = "$D$2", "$H$2", "$K$2"

    cc = ["ESTATUS HOY", "QUÉ HACER", "DÍAS QUE QUEDAN", "CLIENTE", "TELÉFONO", "CORREO",
          "PRÓXIMA FECHA", "OPORTUNIDAD", "LLAMAR DESDE", "ÚLTIMO DÍA PÓLIZA EXTERNA",
          "ABRE CONTADO", "ÚLTIMO DÍA CONTADO", "CONTACTADO", "FECHA CONTACTO", "RESULTADO", "NOTAS",
          "ASEGURADORA", "PÓLIZA ACTUAL", "PRIMA ANUAL ACTUAL", "VEHÍCULO", "SERIE", "FOLIO CRÉDITO",
          "MUNICIPIO", "DIRECCIÓN", "TIPO DE OPORTUNIDAD", "SIGUIENTE ANIVERSARIO", "PRIMERA OPORTUNIDAD",
          "FIN DEL CRÉDITO"]
    X = {c: L(i + 1) for i, c in enumerate(cc)}
    HDR = 5
    filas_c, recuperados = [], 0

    def prox_py(base, fin):  # solo para ordenar: próxima fecha clave vista desde hoy
        if fin is not None and hoy >= fin:
            return hoy  # crédito terminado: cliente libre, se puede llamar ya
        n = 0
        while sumar_anios(base, n) <= hoy:
            n += 1
        a = sumar_anios(base, n)
        return a if (fin is None or a < fin) else fin

    orden = []
    for c in creditos:
        p1, p2, base, fin = _roles(c)
        a = prox_py(base, fin)
        orden.append((date.fromordinal(a.toordinal() - DIAS_POLIZA_EXTERNA - DIAS_AVISO_PREVIO) if a else date.max,
                      c, p1, p2, base, fin))
    orden.sort(key=lambda t: (t[0], t[1]["polizas"][-1].get("NOMBRE") or ""))

    for i, (_, c, p1, p2, base, fin) in enumerate(orden, start=HDR + 1):
        u = c["polizas"][-1]
        actual = p2 or p1
        B, FN, A = f'{X["PRIMERA OPORTUNIDAD"]}{i}', f'{X["FIN DEL CRÉDITO"]}{i}', f'{X["PRÓXIMA FECHA"]}{i}'
        SG, T = f'{X["SIGUIENTE ANIVERSARIO"]}{i}', f'{X["TIPO DE OPORTUNIDAD"]}{i}'
        CT, FC = f'{X["CONTACTADO"]}{i}', f'{X["FECHA CONTACTO"]}{i}'
        sig = f'=IF({B}>TODAY(),{B},EDATE({B},12*(DATEDIF({B},TODAY(),"y")+1)))'
        tipo = (f'=IF({FN}="","ANIVERSARIO",IF(TODAY()>={FN},"CRÉDITO TERMINADO",'
                f'IF({SG}>={FN},"FIN DE CRÉDITO","ANIVERSARIO")))')
        prox = f'=IF({T}="FIN DE CRÉDITO",{FN},{SG})'
        estatus = (f'=IF(AND({CT}="Sí",OR({FC}="",{FC}>EDATE({A},-12))),"CONTACTADO",'
                   f'IF({T}="CRÉDITO TERMINADO","CRÉDITO TERMINADO · CLIENTE LIBRE",'
                   f'IF({T}="FIN DE CRÉDITO",IF(TODAY()>={A}-{PE}-{PA},"LLAMAR YA · TERMINA SU CRÉDITO","PRÓXIMO · TERMINA SU CRÉDITO"),'
                   f'IF(TODAY()>={A}-{PC},"VENTANA DE CONTADO ABIERTA",'
                   f'IF(TODAY()>{A}-{PE},"EXTERNA CERRADA · ESPERAR CONTADO",'
                   f'IF(TODAY()>={A}-{PE}-{PA},"LLAMAR YA · PÓLIZA EXTERNA","PRÓXIMO"))))))')
        que = "=" + "".join(f'IF({X["ESTATUS HOY"]}{i}="{e}","{q}",' for e, q, _, _ in ESTATUS_LISTA) + '""' + ")" * len(ESTATUS_LISTA)
        n_tot = f'IF({FN}="","?",DATEDIF({B},{FN},"y")+(DAY({B})<>DAY({FN})))'
        oport = (f'=IF({T}="FIN DE CRÉDITO","Fin del crédito",IF({T}="CRÉDITO TERMINADO","Auto ya liberado",'
                 f'(DATEDIF({B},{A},"y")+1)&" de "&{n_tot}))')
        dias = (f'=IF({T}="CRÉDITO TERMINADO","",IF({T}="FIN DE CRÉDITO",{A}-TODAY(),'
                f'IF({A}-{PE}<TODAY(),"Externa cerrada",{A}-{PE}-TODAY())))')
        solo_aniv = lambda fx: f'=IF({T}="ANIVERSARIO",{fx},"No aplica")'
        seg = previo.get(u.get("SERIE"), (None, None, None, None))
        recuperados += any(seg)
        filas_c.append([
            estatus, que, dias,
            u.get("NOMBRE", ""), u.get("TELÉFONO", ""), u.get("CORREO", ""),
            prox, oport, f"={A}-{PE}-{PA}", solo_aniv(f"{A}-{PE}"), solo_aniv(f"{A}-{PC}"), solo_aniv(f"{A}-1"),
            *seg,
            actual.get("ASEGURADORA", ""), actual.get("NO. PÓLIZA", ""), actual.get("PRIMA ANUAL EST."),
            limpiar(f"{u.get('DATOS DE VEHÍCULO', '')} {u.get('AÑO MODELO', '')}"), u.get("SERIE", ""),
            u.get("FOLIO CRÉDITO", ""), u.get("MUNICIPIO", ""), u.get("DIRECCIÓN", ""), tipo, sig, base, fin])
    anchos_c = {"ESTATUS HOY": 34, "QUÉ HACER": 48, "DÍAS QUE QUEDAN": 12, "CLIENTE": 32,
                "TELÉFONO": 13, "CORREO": 22, "PRÓXIMA FECHA": 13, "TIPO DE OPORTUNIDAD": 18,
                "SIGUIENTE ANIVERSARIO": 13, "OPORTUNIDAD": 11, "LLAMAR DESDE": 12,
                "ÚLTIMO DÍA PÓLIZA EXTERNA": 13, "ABRE CONTADO": 12, "ÚLTIMO DÍA CONTADO": 12, "CONTACTADO": 11,
                "FECHA CONTACTO": 12, "RESULTADO": 22, "NOTAS": 34, "ASEGURADORA": 12, "PÓLIZA ACTUAL": 17,
                "PRIMA ANUAL ACTUAL": 14, "VEHÍCULO": 34, "SERIE": 20, "FOLIO CRÉDITO": 11, "MUNICIPIO": 18,
                "DIRECCIÓN": 50, "PRIMERA OPORTUNIDAD": 13, "FIN DEL CRÉDITO": 12}
    ultc = tabla(wc, HDR, cc, anchos_c, filas_c,
                 fechas=("PRÓXIMA FECHA", "SIGUIENTE ANIVERSARIO", "LLAMAR DESDE", "ÚLTIMO DÍA PÓLIZA EXTERNA", "ABRE CONTADO",
                         "ÚLTIMO DÍA CONTADO", "FECHA CONTACTO", "PRIMERA OPORTUNIDAD", "FIN DEL CRÉDITO"),
                 texto=("TELÉFONO", "SERIE", "PÓLIZA ACTUAL", "FOLIO CRÉDITO"),
                 centro=("DÍAS QUE QUEDAN", "OPORTUNIDAD", "ASEGURADORA", "CONTACTADO", "TIPO DE OPORTUNIDAD"),
                 moneda=("PRIMA ANUAL ACTUAL",),
                 captura=("CONTACTADO", "FECHA CONTACTO", "RESULTADO", "NOTAS"))
    for c in wc["A"][HDR:]:
        c.font = Font(name="Arial", size=10, bold=True)
    for c in wc[X["DÍAS QUE QUEDAN"]][HDR:]:
        c.number_format = "0"
    for c in wc[X["QUÉ HACER"]][HDR:]:
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.font = Font(name="Arial", size=9)
    for e, _, fondo, letra in ESTATUS_LISTA:
        wc.conditional_formatting.add(f"A{HDR + 1}:C{ultc}", FormulaRule(
            formula=[f'$A{HDR + 1}="{e}"'], fill=relleno(fondo), font=Font(color=letra, bold=True)))
    for col, color in (("ÚLTIMO DÍA PÓLIZA EXTERNA", "FFC7CE"), ("ÚLTIMO DÍA CONTADO", "FFEB9C")):
        for c in wc[X[col]][HDR:]:
            c.fill = relleno(color)
    dv_si = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    dv_res = DataValidation(type="list", formula1='"' + ",".join(RESULTADOS) + '"', allow_blank=True)
    wc.add_data_validation(dv_si)
    wc.add_data_validation(dv_res)
    dv_si.add(f'{X["CONTACTADO"]}{HDR + 1}:{X["CONTACTADO"]}{ultc}')
    dv_res.add(f'{X["RESULTADO"]}{HDR + 1}:{X["RESULTADO"]}{ultc}')
    # contador por estatus (fila 3)
    rango_est = f"$A${HDR + 1}:$A${ultc}"
    col = 1
    for e, _, fondo, letra in ESTATUS_LISTA:
        c = wc.cell(3, col, f'="{e}: "&COUNTIF({rango_est},"{e}")')
        c.font = Font(name="Arial", size=9, bold=True, color=letra)
        c.fill = relleno(fondo)
        col += 3
    wc["A4"] = ("Regla: póliza externa hasta 45 días antes del aniversario (después se cierra) · "
                "contado desde 30 días antes hasta el día anterior · llamar al menos 7 días antes del límite de externa · "
                "al terminar el crédito el auto es del cliente y puede contratar el seguro que quiera.")
    wc["A4"].font = Font(name="Arial", size=9, italic=True, color="595959")
    wc.freeze_panes = wc.cell(HDR + 1, 5)

    # ======================= Hoja BASE DE DATOS =======================
    cb = ["FOLIO CRÉDITO", "CLIENTE", "TIPO PERSONA", "RFC", "EDAD", "SEXO", "TELÉFONO", "CORREO",
          "DIRECCIÓN", "MUNICIPIO", "ESTADO", "CP",
          "VEHÍCULO", "MODELO MG", "AÑO MODELO", "SERIE", "MOTOR",
          "ASEGURADORA", "PAQUETE", "SUMA ASEGURADA DM", "DEDUCIBLE DM", "DEDUCIBLE RT", "RC TERCEROS",
          "PÓLIZA 1 (CONTADO)", "P1 DESDE", "P1 HASTA", "PRIMA P1",
          "PÓLIZA 2 (RESTO DEL PLAZO)", "P2 DESDE", "P2 HASTA", "PRIMA P2 (TOTAL)", "PRIMA ANUAL P2",
          "INICIO DEL CRÉDITO", "FIN DEL CRÉDITO", "PLAZO (AÑOS)", "ESTADO DEL CRÉDITO", "OPORTUNIDADES EN TOTAL",
          "BENEFICIARIO PREFERENTE", "AGENCIA", "ARCHIVO PDF", "NOTA"]
    B_ = {c: L(i + 1) for i, c in enumerate(cb)}
    datos = []
    creditos.sort(key=lambda c: (c["polizas"][-1].get("NOMBRE") or ""))
    for i, c in enumerate(creditos, start=2):
        p1, p2, base, fin = _roles(c)
        u = c["polizas"][-1]
        ref = p2 or p1
        nac = u.get("FECHA NAC. (RFC)")
        datos.append([
            u.get("FOLIO CRÉDITO", ""), u.get("NOMBRE", ""), u.get("TIPO PERSONA", ""), u.get("RFC", ""),
            f'=DATEDIF(DATE({nac.year},{nac.month},{nac.day}),TODAY(),"y")' if nac else "",
            u.get("SEXO", ""), u.get("TELÉFONO", ""), u.get("CORREO", ""),
            u.get("DIRECCIÓN", ""), u.get("MUNICIPIO", ""), u.get("ESTADO", ""), u.get("CP", ""),
            u.get("DATOS DE VEHÍCULO", ""), u.get("MODELO MG", ""), u.get("AÑO MODELO", ""), u.get("SERIE", ""),
            u.get("MOTOR", ""),
            ref.get("ASEGURADORA", ""), ref.get("PAQUETE", ""), ref.get("SUMA ASEGURADA DM") or (p1 or {}).get("SUMA ASEGURADA DM", ""),
            ref.get("DEDUCIBLE DM"), ref.get("DEDUCIBLE RT"), ref.get("RC TERCEROS"),
            p1.get("NO. PÓLIZA", "") if p1 else "", p1.get("VIGENCIA DESDE") if p1 else None,
            p1.get("VIGENCIA HASTA") if p1 else None, p1.get("PRIMA TOTAL") if p1 else None,
            p2.get("NO. PÓLIZA", "") if p2 else "", p2.get("VIGENCIA DESDE") if p2 else None,
            p2.get("VIGENCIA HASTA") if p2 else None, p2.get("PRIMA TOTAL") if p2 else None,
            p2.get("PRIMA ANUAL EST.") if p2 else None,
            c["inicio"], fin, c["plazo"] if fin else "",
            f'=IF({B_["FIN DEL CRÉDITO"]}{i}="","FIN NO CONOCIDO",IF(TODAY()>={B_["FIN DEL CRÉDITO"]}{i},"TERMINADO · AUTO LIBERADO",'
            f'IF({B_["FIN DEL CRÉDITO"]}{i}-TODAY()<=365,"TERMINA EN MENOS DE 1 AÑO","VIGENTE")))',
            (len(c["aniversarios"]) + 1) if fin else "",
            u.get("BENEFICIARIO PREFERENTE", "") or (p1 or {}).get("BENEFICIARIO PREFERENTE", ""),
            u.get("AGENCIA", "") or (p1 or {}).get("AGENCIA", ""),
            " | ".join(sorted({r.get("ARCHIVO", "") for r in c["polizas"]})), c["nota"]])
    anchos_b = {"FOLIO CRÉDITO": 11, "CLIENTE": 32, "TIPO PERSONA": 9, "RFC": 15, "EDAD": 6, "SEXO": 10,
                "TELÉFONO": 13, "CORREO": 22, "DIRECCIÓN": 48, "MUNICIPIO": 17, "ESTADO": 11, "CP": 7,
                "VEHÍCULO": 32, "MODELO MG": 8, "AÑO MODELO": 7, "SERIE": 20, "MOTOR": 18, "ASEGURADORA": 12,
                "PAQUETE": 14, "SUMA ASEGURADA DM": 14, "DEDUCIBLE DM": 9, "DEDUCIBLE RT": 9, "RC TERCEROS": 14,
                "PÓLIZA 1 (CONTADO)": 17, "P1 DESDE": 11, "P1 HASTA": 11, "PRIMA P1": 12,
                "PÓLIZA 2 (RESTO DEL PLAZO)": 17, "P2 DESDE": 11, "P2 HASTA": 11, "PRIMA P2 (TOTAL)": 13,
                "PRIMA ANUAL P2": 12, "INICIO DEL CRÉDITO": 11, "FIN DEL CRÉDITO": 11, "PLAZO (AÑOS)": 7, "ESTADO DEL CRÉDITO": 22,
                "OPORTUNIDADES EN TOTAL": 11, "BENEFICIARIO PREFERENTE": 18, "AGENCIA": 30, "ARCHIVO PDF": 40,
                "NOTA": 40}
    ultb = tabla(wbd, 1, cb, anchos_b, datos,
                 fechas=("P1 DESDE", "P1 HASTA", "P2 DESDE", "P2 HASTA", "INICIO DEL CRÉDITO", "FIN DEL CRÉDITO"),
                 texto=("FOLIO CRÉDITO", "RFC", "TELÉFONO", "CP", "AÑO MODELO", "SERIE", "MOTOR",
                        "PÓLIZA 1 (CONTADO)", "PÓLIZA 2 (RESTO DEL PLAZO)"),
                 centro=("TIPO PERSONA", "EDAD", "MODELO MG", "ASEGURADORA", "PLAZO (AÑOS)", "OPORTUNIDADES EN TOTAL"),
                 moneda=("SUMA ASEGURADA DM", "RC TERCEROS", "PRIMA P1", "PRIMA P2 (TOTAL)", "PRIMA ANUAL P2"),
                 pct=("DEDUCIBLE DM", "DEDUCIBLE RT"))
    wbd.freeze_panes = "C2"

    # Solo si hubo PDFs con problemas, una hoja más para revisarlos
    if errores:
        wr = wb.create_sheet("Por revisar")
        tabla(wr, 1, ["ARCHIVO", "MOTIVO"], {"ARCHIVO": 45, "MOTIVO": 90}, [list(e) for e in errores])

    wb.active = 0
    wb.save(salida)
    if recuperados:
        print(f"  Seguimiento recuperado del archivo anterior: {recuperados} clientes")
    return creditos


def main(rutas, salida, hilos=HILOS):
    print(f"Leyendo: {', '.join(rutas)}")
    zip_revisar = os.path.splitext(salida)[0] + "_POR_REVISAR.zip"
    filas, errores = extraer(rutas, hilos, zip_revisar)
    creditos = guardar_excel(filas, errores, salida)
    print(f"\nListo: {len(filas)} pólizas de {len(creditos)} créditos -> {salida}")
    if errores:
        print(f"\n  {len(errores)} archivo(s) NO se leyeron completos (hoja 'Por revisar'):")
        for archivo, motivo in errores:
            print(f"   - {archivo}: {motivo}")
        print(f"\n  Copia de esos PDFs, lista para enviar: {zip_revisar}")
    else:
        print("  Todos los archivos se leyeron completos.")
    return filas, errores


def subir_y_procesar():
    """Colab: botón para subir archivos, procesa y descarga el Excel."""
    from google.colab import files
    shutil.rmtree(CARPETA_SUBIDA, ignore_errors=True)  # no mezclar con archivos de una corrida anterior
    os.makedirs(CARPETA_SUBIDA, exist_ok=True)
    print("Sube los PDFs de las pólizas o los .zip (puedes elegir varios a la vez).")
    print("Si tienes el Excel anterior con tu seguimiento, súbelo también.\n")
    subidos = files.upload()
    if not subidos:
        print("No se subió ningún archivo.")
        return
    rutas = []
    for nombre, contenido in subidos.items():
        low = nombre.lower()
        if low.endswith(".xlsx"):  # Excel anterior: se usa para conservar el seguimiento
            with open(ARCHIVO_SALIDA, "wb") as f:
                f.write(contenido)
            print(f"  Excel anterior recibido: {nombre} (se conservará tu seguimiento)")
            continue
        if not (low.endswith(".pdf") or low.endswith(".zip")):
            print(f"  Omitido (no es PDF, ZIP ni XLSX): {nombre}")
            continue
        destino = os.path.join(CARPETA_SUBIDA, os.path.basename(nombre))
        with open(destino, "wb") as f:
            f.write(contenido)
        rutas.append(destino)
    del subidos  # libera memoria
    if not rutas:
        print("No hay PDFs ni ZIPs para procesar.")
        return
    main([CARPETA_SUBIDA], ARCHIVO_SALIDA)
    files.download(ARCHIVO_SALIDA)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("-f")]  # ignora args de Jupyter
    if args and not any(a.endswith(".json") for a in args):
        salida = "Polizas_extraidas.xlsx"
        if "-o" in args:
            k = args.index("-o")
            salida = args[k + 1]
            args = args[:k] + args[k + 2:]
        main(args, salida)
    else:
        try:
            import google.colab  # noqa: F401
            en_colab = True
        except ImportError:
            en_colab = False
        if en_colab:
            subir_y_procesar()
        else:
            print("Uso: python extraer_polizas.py carpeta_o_zip [...] -o Polizas_extraidas.xlsx")
