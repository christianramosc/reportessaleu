"""
CUADRE DE ARCHIVOS — App de Streamlit
======================================
Cuadra (reconcilia) dos archivos Excel o CSV por una columna llave y compara
las columnas que elijas, generando observaciones automáticas.

Modos:
  - Automático: detecta la fila de encabezados, sugiere la columna llave y
    qué columnas comparar (por nombre parecido y por valores que coinciden).
    Todo queda editable antes de ejecutar.
  - Manual: tú eliges la llave y armas la lista de columnas a comparar.

En la app de MG Colima se usa como página: paginas/cuadre_archivos.py llama a main().
Usa ui.py para conservar los archivos al cambiar de página y para el encabezado.

Requisitos: streamlit, pandas, openpyxl  (xlrd solo si usas archivos .xls)
"""

import datetime
import difflib
import io
import re
import unicodedata

import pandas as pd
import streamlit as st

import ui
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TIPOS = ["Texto flexible", "Texto exacto", "Número", "Fecha"]
MODOS = ["Igual", "Contiene"]

EST_OK = "Coincide"
EST_OBS = "Con observaciones"
PALABRAS_LLAVE = ("VIN", "SERIE", "FOLIO", "ID", "CLAVE", "RFC", "CUENTA",
                  "CREDITO", "CRÉDITO", "CONTRATO", "NUMERO", "NÚMERO", "NO.", "SOLICITUD")


# ---------------------------------------------------------------------------
# Normalización
# ---------------------------------------------------------------------------

def quitar_acentos(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def norm_llave(v, sin_ceros=False):
    if pd.isna(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip().upper()
    if sin_ceros:
        s = s.lstrip("0") or "0"
    return s


def norm_valor(v, tipo):
    if pd.isna(v):
        return None
    if tipo == "Texto exacto":
        s = str(v).strip()
        return s or None
    if tipo == "Número":
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return float(v)
        try:
            return float(str(v).replace(",", "").replace("$", "").strip())
        except ValueError:
            return None
    if tipo == "Fecha":
        try:
            f = pd.to_datetime(v, dayfirst=True)
            return None if pd.isna(f) else f.date()
        except (ValueError, TypeError, OverflowError):
            return None
    # Texto flexible: sin acentos, sin espacios/símbolos, sin ceros a la izquierda
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = quitar_acentos(str(v)).upper()
    s = re.sub(r"[^A-Z0-9]", "", s)
    s = re.sub(r"(?<![0-9])0+(?=[0-9])", "", s)
    return s or None


def coinciden(a, b, tipo, modo, tol):
    if tipo == "Número":
        return abs(a - b) <= (tol or 0)
    if modo == "Contiene" and isinstance(a, str) and isinstance(b, str):
        return a in b or b in a
    return a == b


def norm_nombre(c):
    return re.sub(r"[^a-z0-9]", "", quitar_acentos(str(c)).lower())


def inferir_tipo(serie):
    s = serie.dropna().head(200)
    if s.empty:
        return "Texto flexible"
    if all(isinstance(v, (pd.Timestamp, datetime.datetime, datetime.date)) for v in s):
        return "Fecha"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in s):
        return "Número"
    return "Texto flexible"


# ---------------------------------------------------------------------------
# Lectura de archivos
# ---------------------------------------------------------------------------

@st.cache_data(show_spinner=False)
def hojas_de(contenido, nombre):
    if nombre.lower().endswith(".csv"):
        return ["(CSV)"]
    return pd.ExcelFile(io.BytesIO(contenido)).sheet_names


@st.cache_data(show_spinner=False)
def leer_crudo(contenido, nombre, hoja):
    if nombre.lower().endswith(".csv"):
        try:
            return pd.read_csv(io.BytesIO(contenido), header=None, dtype=object)
        except UnicodeDecodeError:
            return pd.read_csv(io.BytesIO(contenido), header=None, dtype=object, encoding="latin-1")
    df = pd.read_excel(io.BytesIO(contenido), sheet_name=hoja, header=None)
    return df.dropna(axis=1, how="all").dropna(how="all")


def para_mostrar(df):
    """Convierte a texto para que Streamlit muestre columnas con tipos mezclados sin errores."""
    return df.astype(str).replace({"nan": "", "None": "", "NaT": ""})


def detectar_encabezado(df_raw, max_filas=30):
    """La fila con más celdas de texto dentro de las primeras filas."""
    mejor, mejor_score = 0, -1
    for pos in range(min(max_filas, len(df_raw))):
        fila = df_raw.iloc[pos]
        textos = sum(1 for v in fila if isinstance(v, str) and v.strip())
        if textos > mejor_score:
            mejor, mejor_score = pos, textos
    return mejor


def armar_tabla(df_raw, pos_enc):
    enc = df_raw.iloc[pos_enc].tolist()
    df = df_raw.iloc[pos_enc + 1:].copy()
    nombres, vistos = [], {}
    for i, h in enumerate(enc):
        n = str(h).strip() if pd.notna(h) and str(h).strip() else f"Columna_{i + 1}"
        if n in vistos:
            vistos[n] += 1
            n = f"{n}_{vistos[n]}"
        else:
            vistos[n] = 0
        nombres.append(n)
    df.columns = nombres
    return df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)


# ---------------------------------------------------------------------------
# Sugerencias automáticas
# ---------------------------------------------------------------------------

@st.cache_data(show_spinner=False)
def sugerir_llave(df_a, df_b, sin_ceros=False):
    """Busca el par de columnas (A, B) con valores casi únicos que más se cruzan."""
    sets_b = {}
    for cb in df_b.columns:
        if inferir_tipo(df_b[cb]) == "Fecha":
            continue
        sets_b[cb] = set(df_b[cb].dropna().map(lambda v: norm_llave(v, sin_ceros))) - {""}

    mejor, mejor_score = (None, None), 0.0
    for ca in df_a.columns:
        if inferir_tipo(df_a[ca]) == "Fecha":
            continue
        va = df_a[ca].dropna().map(lambda v: norm_llave(v, sin_ceros))
        va = va[va != ""]
        if va.empty:
            continue
        unicidad = va.nunique() / len(va)
        if unicidad < 0.8:
            continue
        set_a = set(va)
        largo = min(1.0, va.str.len().mean() / 6)  # castiga consecutivos tipo 1,2,3
        for cb, sb in sets_b.items():
            inter = len(set_a & sb)
            if not inter:
                continue
            score = inter / min(len(set_a), len(sb)) * unicidad * largo
            if any(p in str(ca).upper() or p in str(cb).upper() for p in PALABRAS_LLAVE):
                score += 0.15
            if score > mejor_score:
                mejor, mejor_score = (ca, cb), score
    return mejor


def _tasa_acuerdo(sa, sb):
    pares = [(x, y) for x, y in zip(sa, sb) if pd.notna(x) and pd.notna(y)]
    if len(pares) < 3:
        return 0.0, 0.0, "Texto flexible", len(pares)
    ta = inferir_tipo(pd.Series([p[0] for p in pares]))
    tb = inferir_tipo(pd.Series([p[1] for p in pares]))
    if "Fecha" in (ta, tb):
        tipo = "Fecha"
    elif ta == tb == "Número":
        tipo = "Número"
    else:
        tipo = "Texto flexible"
    igual = contiene = 0
    for x, y in pares:
        a, b = norm_valor(x, tipo), norm_valor(y, tipo)
        if a is None or b is None:
            continue
        if tipo == "Número":
            ok = abs(a - b) <= 0.01
            igual += ok
            contiene += ok
            continue
        igual += a == b
        if isinstance(a, str) and min(len(a), len(b)) >= 2 and (a in b or b in a):
            contiene += 1
    return igual / len(pares), contiene / len(pares), tipo, len(pares)


@st.cache_data(show_spinner=False)
def sugerir_comparaciones(df_a, df_b, key_a, key_b, sin_ceros=False, muestra=300):
    a = df_a.assign(_k=df_a[key_a].map(lambda v: norm_llave(v, sin_ceros)))
    b = df_b.assign(_k=df_b[key_b].map(lambda v: norm_llave(v, sin_ceros)))
    a = a[a["_k"] != ""].drop_duplicates("_k").set_index("_k")
    b = b[b["_k"] != ""].drop_duplicates("_k").set_index("_k")
    comunes = list(a.index.intersection(b.index))[:muestra]
    a_c, b_c = a.loc[comunes], b.loc[comunes]

    sugerencias, usadas_b = [], set()
    for ca in df_a.columns:
        if ca == key_a:
            continue
        mejor = None  # (score, cb, tipo, modo)
        for cb in df_b.columns:
            if cb == key_b or cb in usadas_b:
                continue
            sim = difflib.SequenceMatcher(None, norm_nombre(ca), norm_nombre(cb)).ratio()
            igual, cont, tipo, n_pares = _tasa_acuerdo(a_c[ca], b_c[cb])
            if igual >= 0.5:
                score, modo = igual + 1, "Igual"
            elif cont >= 0.7:
                score, modo = cont + 0.5, "Contiene"
            elif (sim >= 0.8 and n_pares < 3  # solo por nombre si no hay datos que lo contradigan
                  and inferir_tipo(df_a[ca]) == inferir_tipo(df_b[cb])):
                score, modo = sim, "Igual"
                tipo = inferir_tipo(df_a[ca])
            else:
                continue
            if mejor is None or score > mejor[0]:
                mejor = (score, cb, tipo, modo)
        if mejor:
            usadas_b.add(mejor[1])
            sugerencias.append({"Columna A": ca, "Columna B": mejor[1], "Tipo": mejor[2],
                                "Modo": mejor[3], "Tolerancia": 0.0})
    return sugerencias


# ---------------------------------------------------------------------------
# Motor del cuadre
# ---------------------------------------------------------------------------

def ejecutar_cuadre(df_a, df_b, key_a, key_b, nombre_a, nombre_b, comparaciones,
                    extras_a, extras_b, incluir_solo_b=True, sin_ceros=False):
    solo_a, solo_b = f"Solo en {nombre_a}", f"Solo en {nombre_b}"
    ka = df_a[key_a].map(lambda v: norm_llave(v, sin_ceros))
    kb = df_b[key_b].map(lambda v: norm_llave(v, sin_ceros))
    cnt_a = ka[ka != ""].value_counts()
    grupos_b = kb[kb != ""].groupby(kb[kb != ""]).groups

    def base(llave_original):
        return {"Llave": llave_original}

    filas, usadas_b = [], set()
    for i in df_a.index:
        k = ka[i]
        obs = []
        fila = base(df_a.at[i, key_a])
        fila_b = None

        if k == "":
            obs.append(f"Registro sin llave en {nombre_a}")
        else:
            if cnt_a.get(k, 0) > 1:
                obs.append(f"Llave duplicada en {nombre_a} ({cnt_a[k]} veces)")
            if k in grupos_b:
                idx_b = grupos_b[k]
                usadas_b.add(k)
                if len(idx_b) > 1:
                    obs.append(f"Llave duplicada en {nombre_b} ({len(idx_b)} veces)")
                fila_b = df_b.loc[idx_b[0]]

        for c in comparaciones:
            ca, cb = c["Columna A"], c["Columna B"]
            va = df_a.at[i, ca]
            vb = fila_b[cb] if fila_b is not None else None
            fila[f"{ca} ({nombre_a})"] = va
            fila[f"{cb} ({nombre_b})"] = vb
            if fila_b is None:
                continue
            na, nb = norm_valor(va, c["Tipo"]), norm_valor(vb, c["Tipo"])
            if na is None and nb is None:
                continue
            if na is None or nb is None:
                vacio = nombre_a if na is None else nombre_b
                obs.append(f"{ca}: vacío en {vacio}")
            elif not coinciden(na, nb, c["Tipo"], c["Modo"], c.get("Tolerancia", 0)):
                obs.append(f"{ca} no coincide ({nombre_a}: {va} / {nombre_b}: {vb})")

        for col in extras_a:
            fila[f"{col} ({nombre_a})"] = df_a.at[i, col]
        for col in extras_b:
            fila[f"{col} ({nombre_b})"] = fila_b[col] if fila_b is not None else None

        if fila_b is None:
            estatus = solo_a
            if k != "":
                obs.insert(0, f"No encontrado en {nombre_b}")
        else:
            estatus = EST_OBS if obs else EST_OK
        fila["Estatus"] = estatus
        fila["Observaciones"] = " | ".join(obs) if obs else "OK"
        filas.append(fila)

    if incluir_solo_b:
        for k, idx_b in grupos_b.items():
            if k in usadas_b:
                continue
            for j in idx_b:
                fb = df_b.loc[j]
                fila = base(fb[key_b])
                for c in comparaciones:
                    fila[f"{c['Columna A']} ({nombre_a})"] = None
                    fila[f"{c['Columna B']} ({nombre_b})"] = fb[c["Columna B"]]
                for col in extras_a:
                    fila[f"{col} ({nombre_a})"] = None
                for col in extras_b:
                    fila[f"{col} ({nombre_b})"] = fb[col]
                obs = f"Está en {nombre_b} pero no en {nombre_a}"
                if len(idx_b) > 1:
                    obs += f" | Llave duplicada en {nombre_b} ({len(idx_b)} veces)"
                fila["Estatus"] = solo_b
                fila["Observaciones"] = obs
                filas.append(fila)

    res = pd.DataFrame(filas)
    if res.empty:
        return res
    orden = ["Llave", "Estatus"] + [c for c in res.columns if c not in ("Llave", "Estatus", "Observaciones")] + ["Observaciones"]
    return res[orden]


# ---------------------------------------------------------------------------
# Exportar a Excel
# ---------------------------------------------------------------------------

def a_excel(res, resumen_cfg):
    buf = io.BytesIO()
    verde = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    naranja = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
    rojo = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    azul = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    blanco = Font(bold=True, color="FFFFFF")

    conteo = res["Estatus"].value_counts().rename_axis("Estatus").reset_index(name="Registros")

    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        conteo.to_excel(w, sheet_name="Resumen", index=False)
        cfg = pd.DataFrame(list(resumen_cfg.items()), columns=["Parámetro", "Valor"])
        cfg.to_excel(w, sheet_name="Resumen", index=False, startrow=len(conteo) + 3)
        res.to_excel(w, sheet_name="Cuadre", index=False)

        for nombre_hoja in ("Resumen", "Cuadre"):
            ws = w.sheets[nombre_hoja]
            for c in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(c)].width = 26
        ws_r = w.sheets["Resumen"]
        for fila_enc in (1, len(conteo) + 4):
            for c in (1, 2):
                ws_r.cell(row=fila_enc, column=c).fill = azul
                ws_r.cell(row=fila_enc, column=c).font = blanco

        ws = w.sheets["Cuadre"]
        n_cols = res.shape[1]
        for c in range(1, n_cols + 1):
            cel = ws.cell(row=1, column=c)
            cel.fill, cel.font = azul, blanco
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col_est = res.columns.get_loc("Estatus") + 1
        for r in range(2, len(res) + 2):
            est = ws.cell(row=r, column=col_est).value
            fill = verde if est == EST_OK else naranja if est == EST_OBS else rojo
            for c in range(1, n_cols + 1):
                ws.cell(row=r, column=c).fill = fill
        ws.column_dimensions[get_column_letter(n_cols)].width = 80
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Interfaz
# ---------------------------------------------------------------------------

def panel_archivo(etiqueta, clave):
    archivo = ui.archivo_persistente(f"Archivo {etiqueta}", f"cuadre_arch_{clave}",
                                     ["xlsx", "xlsm", "xls", "csv"])
    if not archivo:
        return None, None
    nombre_arch, contenido = archivo
    try:
        hojas = hojas_de(contenido, nombre_arch)
    except Exception as e:
        st.error(f"No pude abrir el archivo: {e}")
        return None, None
    hoja = st.selectbox("Hoja", hojas, key=f"hoja_{clave}") if len(hojas) > 1 else hojas[0]

    with st.spinner("Leyendo..."):
        crudo = leer_crudo(contenido, nombre_arch, hoja)
    if crudo.empty:
        st.warning("La hoja está vacía.")
        return None, None

    pos_auto = detectar_encabezado(crudo)
    fila_excel_auto = int(crudo.index[pos_auto]) + 1
    fila_enc = st.number_input("Fila de encabezados (como se ve en Excel)", min_value=1,
                               value=fila_excel_auto, step=1, key=f"enc_{clave}_{hoja}",
                               help="Detectada automáticamente; cámbiala si no es correcta.")
    posiciones = [p for p, idx in enumerate(crudo.index) if idx + 1 >= fila_enc]
    if not posiciones:
        st.error("Esa fila está fuera del rango de datos.")
        return None, None
    df = armar_tabla(crudo, posiciones[0])

    with st.expander("Filtrar filas (opcional)"):
        col_f = st.selectbox("Columna", ["(sin filtro)"] + list(df.columns), key=f"fcol_{clave}")
        if col_f != "(sin filtro)":
            valores = sorted(df[col_f].dropna().astype(str).unique().tolist())[:500]
            elegidos = st.multiselect("Conservar solo estos valores", valores, key=f"fval_{clave}")
            if elegidos:
                df = df[df[col_f].astype(str).isin(elegidos)].reset_index(drop=True)

    st.caption(f"{len(df):,} filas · {df.shape[1]} columnas")
    with st.expander("Vista previa"):
        st.dataframe(para_mostrar(df.head(20)), width="stretch")
    return df, f"{nombre_arch} / {hoja}"


def main():
    ui.hero("Cuadre de archivos",
            "Cruza dos archivos por una columna llave, compara columnas y genera observaciones",
            "Excel o CSV")

    with st.sidebar:
        st.header("Opciones")
        nombre_a = st.text_input("Nombre para el archivo A", "Archivo A")
        nombre_b = st.text_input("Nombre para el archivo B", "Archivo B")
        incluir_solo_b = st.checkbox(f"Incluir registros que solo están en B", value=True)
        sin_ceros = st.checkbox("Ignorar ceros a la izquierda en la llave", value=False,
                                help="Útil si en un archivo el folio es 00123 y en el otro 123.")

    c1, c2 = st.columns(2)
    with c1:
        st.subheader(f"📄 {nombre_a}")
        df_a, origen_a = panel_archivo("A", "a")
    with c2:
        st.subheader(f"📄 {nombre_b}")
        df_b, origen_b = panel_archivo("B", "b")

    if df_a is None or df_b is None:
        st.markdown("""<div class="vacio"><b>Sube los dos archivos para continuar</b><br>
        Se detecta la fila de encabezados, se sugiere la columna llave y las columnas a comparar.
        Todo se puede editar antes de ejecutar.</div>""", unsafe_allow_html=True)
        return
    if df_a.empty or df_b.empty:
        st.warning("Uno de los archivos quedó sin filas (revisa el filtro o la fila de encabezados).")
        return

    st.divider()
    modo = st.radio("¿Cómo quieres configurar el cuadre?", ["Automático", "Manual"], horizontal=True,
                    help="Automático propone la llave y las columnas; en ambos casos puedes editarlas.")
    cols_a, cols_b = list(df_a.columns), list(df_b.columns)

    # ---- Paso 1: llave
    st.subheader("1. Columna llave")
    sug_a, sug_b = (None, None)
    if modo == "Automático":
        with st.spinner("Buscando la columna llave..."):
            sug_a, sug_b = sugerir_llave(df_a, df_b, sin_ceros)
        if sug_a:
            st.success(f"Llave sugerida: **{sug_a}** ({nombre_a}) ↔ **{sug_b}** ({nombre_b})")
        else:
            st.warning("No encontré una llave clara; elígela manualmente.")
    k1, k2 = st.columns(2)
    key_a = k1.selectbox(f"Llave en {nombre_a}", cols_a,
                         index=cols_a.index(sug_a) if sug_a in cols_a else 0, key=f"ka_{modo}")
    key_b = k2.selectbox(f"Llave en {nombre_b}", cols_b,
                         index=cols_b.index(sug_b) if sug_b in cols_b else 0, key=f"kb_{modo}")

    # ---- Paso 2: columnas a comparar
    st.subheader("2. Columnas a comparar")
    if modo == "Automático":
        with st.spinner("Buscando columnas equivalentes..."):
            sugerencias = sugerir_comparaciones(df_a, df_b, key_a, key_b, sin_ceros)
        if not sugerencias:
            st.warning("No encontré columnas equivalentes; agrégalas en la tabla.")
        base = pd.DataFrame(sugerencias or [{"Columna A": None, "Columna B": None, "Tipo": "Texto flexible",
                                             "Modo": "Igual", "Tolerancia": 0.0}])
    else:
        base = pd.DataFrame([{"Columna A": None, "Columna B": None, "Tipo": "Texto flexible",
                              "Modo": "Igual", "Tolerancia": 0.0}])

    st.caption("Agrega o borra filas. **Texto flexible** ignora mayúsculas, acentos, espacios, guiones y ceros "
               "a la izquierda (MG-05 = MG5). **Contiene** acepta que un valor esté dentro del otro "
               "(MG5 ≈ MG5 - 2027). **Tolerancia** aplica solo a Número.")
    editor_key = f"cmp_{modo}_{origen_a}_{origen_b}_{key_a}_{key_b}_{len(df_a)}_{len(df_b)}"
    comparaciones_df = st.data_editor(
        base, num_rows="dynamic", width="stretch", key=editor_key,
        column_config={
            "Columna A": st.column_config.SelectboxColumn(f"Columna en {nombre_a}", options=cols_a),
            "Columna B": st.column_config.SelectboxColumn(f"Columna en {nombre_b}", options=cols_b),
            "Tipo": st.column_config.SelectboxColumn("Tipo", options=TIPOS, default="Texto flexible"),
            "Modo": st.column_config.SelectboxColumn("Modo", options=MODOS, default="Igual"),
            "Tolerancia": st.column_config.NumberColumn("Tolerancia", min_value=0.0, default=0.0, format="%.2f"),
        },
    )
    comparaciones = [
        {**r, "Tipo": r.get("Tipo") or "Texto flexible", "Modo": r.get("Modo") or "Igual",
         "Tolerancia": float(r.get("Tolerancia") or 0)}
        for r in comparaciones_df.to_dict("records")
        if r.get("Columna A") in cols_a and r.get("Columna B") in cols_b
    ]

    # ---- Paso 3: columnas extra
    with st.expander("3. Columnas adicionales para mostrar en el resultado (opcional)"):
        e1, e2 = st.columns(2)
        extras_a = e1.multiselect(f"De {nombre_a}", [c for c in cols_a if c != key_a])
        extras_b = e2.multiselect(f"De {nombre_b}", [c for c in cols_b if c != key_b])

    if st.button("▶️ Ejecutar cuadre", type="primary", width="stretch"):
        with st.spinner("Cuadrando..."):
            res = ejecutar_cuadre(df_a, df_b, key_a, key_b, nombre_a, nombre_b, comparaciones,
                                  extras_a, extras_b, incluir_solo_b, sin_ceros)
        cfg = {
            f"Origen {nombre_a}": origen_a, f"Origen {nombre_b}": origen_b,
            "Llave": f"{key_a} ↔ {key_b}",
            "Comparaciones": "; ".join(f"{c['Columna A']} ↔ {c['Columna B']} ({c['Tipo']}, {c['Modo']})"
                                       for c in comparaciones) or "Ninguna (solo cruce de llaves)",
            "Fecha de cuadre": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        }
        st.session_state["cuadre_res"] = (res, cfg, nombre_a, nombre_b, (origen_a, origen_b))

    if "cuadre_res" not in st.session_state:
        return
    res, cfg, na, nb, origenes = st.session_state["cuadre_res"]
    if origenes != (origen_a, origen_b):
        st.info("Cambiaron los archivos u hojas. Vuelve a ejecutar el cuadre para ver resultados.")
        return
    if res.empty:
        st.warning("El cuadre no produjo resultados.")
        return

    # ---- Resultados
    st.divider()
    st.subheader("Resultados")
    conteo = res["Estatus"].value_counts()
    m = st.columns(4)
    m[0].metric("✅ Coinciden", int(conteo.get(EST_OK, 0)))
    m[1].metric("⚠️ Con observaciones", int(conteo.get(EST_OBS, 0)))
    m[2].metric(f"❌ Solo en {na}", int(conteo.get(f"Solo en {na}", 0)))
    m[3].metric(f"❌ Solo en {nb}", int(conteo.get(f"Solo en {nb}", 0)))

    f1, f2 = st.columns([2, 1])
    estatus_sel = f1.multiselect("Filtrar por estatus", list(conteo.index), default=list(conteo.index))
    buscar = f2.text_input("Buscar texto")
    vista = res[res["Estatus"].isin(estatus_sel)]
    if buscar:
        mask = vista.astype(str).apply(lambda col: col.str.contains(buscar, case=False, regex=False)).any(axis=1)
        vista = vista[mask]
    st.dataframe(para_mostrar(vista), width="stretch", hide_index=True)

    with st.expander("Observaciones más frecuentes"):
        tipos_obs = (res.loc[res["Observaciones"] != "OK", "Observaciones"].str.split(r" \| ").explode()
                     .str.replace(r"\(.*\)", "", regex=True).str.strip().value_counts())
        st.dataframe(tipos_obs.rename_axis("Observación").reset_index(name="Veces"),
                     width="stretch", hide_index=True)

    st.download_button(
        "⬇️ Descargar Excel del cuadre", data=a_excel(res, cfg),
        file_name=f"Cuadre_{na}_vs_{nb}_{datetime.date.today():%Y%m%d}.xlsx".replace(" ", "_"),
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary", width="stretch",
    )


if __name__ == "__main__":
    st.set_page_config(page_title="Cuadre de archivos", page_icon="🔀", layout="wide")
    main()
